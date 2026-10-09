#!/usr/bin/env python3
"""Measure what a revised deck kept of the deck it revises, slide by slide.

    python3 evals/point-change/preservation.py source.pptx revised.pptx [--pairs 1:1,2:2,4:3] [--json out.json]

A point change must leave what it does not touch as it was. For each pair
`source slide : revised slide` (by default slide n against slide n) this says

    part       the slide's XML part is the source's, byte for byte
    parts      so is every part it draws on - its layout, master and theme,
               its charts and their workbooks, its pictures, its notes
    text       it prints the same paragraphs, in the same order
    geometry   the furthest any shape moved or resized, in px at 96 an inch:
               each source shape is set against the revised shape that says
               the same words (or, for a picture, a chart or a table, the
               next of its kind); `unmatched` counts the shapes with no
               counterpart

and sums them: the slides identical as bytes, the slides whose text held, and
the largest movement on any slide. A carried slide of a revision this skill
built is identical on all four; a slide recomposed from an inventory keeps
its words at best.
"""
from __future__ import annotations

import argparse
import io
import json
import posixpath
import sys
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.shapes.group import GroupShape

EMU_PER_PX = 9525
REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"
PML = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"


def parts_of(data: bytes) -> dict:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def rels(parts: dict, part: str) -> list:
    name = posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")
    if name not in parts:
        return []
    out = []
    for el in etree.fromstring(parts[name]):
        if el.get("TargetMode") == "External":
            continue
        target = el.get("Target", "")
        out.append((el.get("Type").rsplit("/", 1)[-1], target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(posixpath.dirname(part), target))))
    return out


def slide_parts(parts: dict) -> list:
    root = etree.fromstring(parts["ppt/presentation.xml"])
    by_id = {}
    for el in etree.fromstring(parts["ppt/_rels/presentation.xml.rels"]):
        by_id[el.get("Id")] = posixpath.normpath(posixpath.join("ppt", el.get("Target", "")))
    listed = root.find(f"{PML}sldIdLst")
    return [by_id[el.get(R_ID)] for el in (listed if listed is not None else [])]


def drawn_on(parts: dict, part: str) -> list:
    """The content of every part `part` draws on, in the order its relationships are followed, each with the kind of relationship that reached it."""
    seen, out, queue = {part}, [], [("slide", part)]
    while queue:
        kind, name = queue.pop(0)
        out.append((kind, parts.get(name)))
        for next_kind, target in sorted(rels(parts, name)):
            # A notes page points back at its slide, and a layout at every sibling through the master: follow each part once.
            if target not in seen:
                seen.add(target)
                queue.append((next_kind, target))
    return out


def leaves(shapes):
    for shape in shapes:
        if isinstance(shape, GroupShape):
            yield from leaves(shape.shapes)
        else:
            yield shape


def described(slide) -> list:
    """Each leaf shape as (kind, the words it says, its frame in px)."""
    out = []
    for shape in leaves(slide.shapes):
        kind = "chart" if getattr(shape, "has_chart", False) else "table" if getattr(shape, "has_table", False) else "picture" if shape.shape_type == 13 else "text" if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip() else "shape"
        words = " ".join(shape.text_frame.text.split()) if kind == "text" else ""
        frame = tuple(round((value or 0) / EMU_PER_PX, 1) for value in (shape.left, shape.top, shape.width, shape.height))
        out.append((kind, words, frame))
    return out


def paragraphs(slide) -> list:
    out = []
    for shape in leaves(slide.shapes):
        if getattr(shape, "has_table", False):
            out.extend(" ".join(cell.text.split()) for row in shape.table.rows for cell in row.cells)
        elif getattr(shape, "has_text_frame", False):
            out.extend(" ".join(p.text.split()) for p in shape.text_frame.paragraphs if p.text.strip())
    return out


def moved(before: list, after: list) -> tuple:
    """The furthest a shape of `before` moved or resized against its counterpart in `after`, and how many have none.
    A shape's counterpart says the same words; one whose words changed is set against the nearest shape of its kind left over."""
    distance = lambda a, b: max(abs(x - y) for x, y in zip(a, b))
    left, furthest, unmatched, reworded = list(after), 0.0, 0, []
    for kind, words, frame in before:
        match = next((item for item in left if item[0] == kind and (kind != "text" or item[1] == words)), None)
        if match is None:
            reworded.append((kind, frame))
            continue
        left.remove(match)
        furthest = max(furthest, distance(frame, match[2]))
    for kind, frame in reworded:
        rest = [item for item in left if item[0] == kind]
        if not rest:
            unmatched += 1
            continue
        match = min(rest, key=lambda item: distance(frame, item[2]))
        left.remove(match)
        furthest = max(furthest, distance(frame, match[2]))
    return round(furthest, 1), unmatched


def measure(source: Path, revised: Path, pairs=None) -> dict:
    a_bytes, b_bytes = source.read_bytes(), revised.read_bytes()
    a_parts, b_parts = parts_of(a_bytes), parts_of(b_bytes)
    a_slides, b_slides = slide_parts(a_parts), slide_parts(b_parts)
    a_deck, b_deck = Presentation(io.BytesIO(a_bytes)), Presentation(io.BytesIO(b_bytes))
    pairs = pairs or [(n, n) for n in range(1, min(len(a_slides), len(b_slides)) + 1)]
    slides = []
    for a, b in pairs:
        if not (1 <= a <= len(a_slides) and 1 <= b <= len(b_slides)):
            slides.append({"source": a, "revised": b, "missing": True})
            continue
        part = a_parts[a_slides[a - 1]] == b_parts[b_slides[b - 1]]
        before, after = described(a_deck.slides[a - 1]), described(b_deck.slides[b - 1])
        distance, unmatched = moved(before, after)
        slides.append({"source": a, "revised": b, "part": part, "parts": part and drawn_on(a_parts, a_slides[a - 1]) == drawn_on(b_parts, b_slides[b - 1]),
                       "text": paragraphs(a_deck.slides[a - 1]) == paragraphs(b_deck.slides[b - 1]), "geometry": distance, "unmatched": unmatched, "shapes": len(before)})
    held = [s for s in slides if not s.get("missing")]
    return {"source": source.name, "revised": revised.name, "sourceSlides": len(a_slides), "revisedSlides": len(b_slides), "compared": len(held),
            "identicalParts": sum(1 for s in held if s["parts"]), "sameText": sum(1 for s in held if s["text"]),
            "furthest": max([s["geometry"] for s in held], default=0), "unmatchedShapes": sum(s["unmatched"] for s in held), "slides": slides}


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("source", type=Path)
    parser.add_argument("revised", type=Path)
    parser.add_argument("--pairs", help="source:revised slide numbers, comma-separated (default: slide n against slide n)")
    parser.add_argument("--json", type=Path)
    args = parser.parse_args(argv)
    pairs = [tuple(int(n) for n in pair.split(":")) for pair in args.pairs.split(",")] if args.pairs else None
    report = measure(args.source, args.revised, pairs)
    if args.json:
        args.json.write_text(json.dumps(report, indent=1), encoding="utf-8")
    print(json.dumps({key: value for key, value in report.items() if key != "slides"}))
    return 0


if __name__ == "__main__":
    sys.exit(main())
