"""Every figure a chart prints from one set of numbers carries one precision.

The label code chose one decimal count for the chart and then printed the
rounded number with String(), which drops a trailing ".0": a series read
"32" beside "40.8" and "15" beside "14.9". The same drift sat in the value
axis, the growth column, radial rings, in-table bars and the native chart.
"""
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from node_probe import run_node

ROOT = Path(__file__).resolve().parents[2]
RUNTIME = ROOT / "skills" / "professional-slides" / "runtime"


class DrawnLabelPrecisionTests(unittest.TestCase):
    def test_a_series_prints_every_label_to_one_precision(self):
        result = run_node("""
import {formatValue} from './skills/professional-slides/runtime/value-format.mjs';
import {REGISTRY} from './skills/professional-slides/runtime/registry.mjs';
const values=[32,40.8,15,14.9];
const nodes=REGISTRY.get('chart.column').render({id:'c',frame:{x:60,y:60,width:900,height:480},
  props:{categories:['A','B','C','D'],series:[{name:'Share',values}]}}).nodes;
const range=REGISTRY.get('chart.range').render({id:'r',frame:{x:60,y:60,width:900,height:480},
  props:{categories:['A','B'],low:[12,14.5],high:[20,22]}}).nodes;
console.log(JSON.stringify({
  direct: values.map(v=>formatValue(v,{series:[{name:'Share',values}]})),
  drawn: nodes.filter(n=>n.role==='data-label').map(n=>n.text),
  range: range.filter(n=>n.role==='data-label').map(n=>n.text),
  whole: [120,40.5].map(v=>formatValue(v,{values:[120,40.5]})),
}));
""")
        self.assertEqual(result["direct"], ["32.0", "40.8", "15.0", "14.9"])
        self.assertEqual(result["drawn"], ["32.0", "40.8", "15.0", "14.9"])
        # A range's two ends are one chart: 12.0 and 14.5, 20.0 and 22.0.
        self.assertEqual(sorted(result["range"]), sorted(["12.0", "14.5", "20.0", "22.0"]))
        # Whole numbers once the chart reaches a hundred, for every label.
        self.assertEqual(result["whole"], ["120", "41"])

    def test_the_value_axis_prints_one_precision(self):
        result = run_node("""
import {axes} from './skills/professional-slides/runtime/charts.mjs';
const plot={x:100,y:60,width:600,height:300};
const text=(min,max)=>axes('a',plot,min,max,4).filter(n=>n.role==='axis-label').map(n=>n.text);
console.log(JSON.stringify({half:text(0,10),whole:text(0,40),signed:text(-5,5)}));
""")
        self.assertEqual(result["half"], ["0.0", "2.5", "5.0", "7.5", "10.0"])
        self.assertEqual(result["whole"], ["0", "10", "20", "30", "40"])
        self.assertEqual(result["signed"], ["-5.0", "-2.5", "0.0", "2.5", "5.0"])

    def test_a_growth_column_and_radial_rings_share_one_precision(self):
        result = run_node("""
import {REGISTRY} from './skills/professional-slides/runtime/registry.mjs';
import {radialBarsLayout} from './skills/professional-slides/runtime/figures.mjs';
const props={categories:['2020','2030'],series:[{name:'Fast',values:[100,400]},{name:'Slow',values:[100,110]}],directLabels:'end',seriesGrowth:{from:'2020',to:'2030'}};
const nodes=REGISTRY.get('chart.line').render({id:'l',frame:{x:0,y:0,width:900,height:420},props}).nodes;
const rings=radialBarsLayout({x:0,y:0,width:600,height:420},{items:[{label:'One',value:40},{label:'Two',value:32.5}]});
console.log(JSON.stringify({rates:nodes.filter(n=>n.data?.growth).map(n=>n.text),rings:(rings.items||rings.rows||[]).map(r=>(r.item||r).display)}));
""")
        # 14.9% a year beside 1.0%: the column takes the decimal throughout.
        self.assertEqual(result["rates"], ["+14.9%", "+1.0%"])
        self.assertEqual(result["rings"], ["40.0%", "32.5%"])

    def test_in_table_bars_share_one_precision_down_the_column(self):
        result = run_node("""
import {compileDeck,component} from './skills/professional-slides/runtime/core.mjs';
import {REGISTRY} from './skills/professional-slides/runtime/registry.mjs';
const scale={type:'bars',label:'Share',unit:'%',min:0,max:50,series:['Share']};
const spec={slides:[{id:'s',frame:{x:60,y:100,width:900,height:400},composition:component({id:'t',component:'table',size:{width:'fill',height:'fill'},
  props:{columns:['Case','Share'],rows:[['A',{type:'bars',values:[15],scale:'share'}],['B',{type:'bars',values:[14.9],scale:'share'}]],scales:{share:scale}}})}]};
const nodes=compileDeck(spec,REGISTRY).slides[0].nodes;
console.log(JSON.stringify(nodes.filter(n=>n.role==='table-cell-text'&&/^[0-9.]+$/.test(n.text)).map(n=>n.text)));
""")
        self.assertEqual(result, ["15.0", "14.9"])


class NativeLabelPrecisionTests(unittest.TestCase):
    def test_the_native_chart_prints_the_drawn_precision(self):
        try:
            import pptx  # noqa: F401
        except ImportError:
            self.skipTest("python-pptx is not installed")
        sys.path.insert(0, str(RUNTIME / "emit"))
        from pptx import Presentation
        scene = run_node("""
import { toDeckPlan } from './skills/professional-slides/runtime/compose.mjs';
import { planDeck } from './skills/professional-slides/runtime/planner.mjs';
const spec={schema:'professional-slides.deck/v3',id:'p',cover:{title:'x'},slides:[{title:'Share rose in both markets over the period',layout:'exhibit-full',
  exhibit:{type:'chart.column',heading:'Share by market',categories:['A','B','C','D'],series:[{name:'Share',values:[32,40.8,15,14.9]}]}}]};
console.log(JSON.stringify(planDeck(toDeckPlan(spec,'.')).deck));
""")
        labels = [n["text"] for n in scene["slides"][-1]["nodes"] if n["role"] == "data-label"]
        self.assertEqual(labels, ["32.0", "40.8", "15.0", "14.9"])
        with tempfile.TemporaryDirectory() as tmp:
            scene_path, pptx_path = Path(tmp) / "scene.json", Path(tmp) / "p.pptx"
            scene_path.write_text(json.dumps(scene))
            subprocess.run([sys.executable, str(RUNTIME / "emit" / "emit_pptx.py"), str(scene_path), str(pptx_path)],
                           check=True, capture_output=True)
            chart = next(s.chart for s in Presentation(pptx_path).slides[-1].shapes if s.has_chart)
            self.assertEqual(chart.plots[0].data_labels.number_format, "0.0")


if __name__ == "__main__":
    unittest.main()
