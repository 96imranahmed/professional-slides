"""How the build ends: refused, crashed, unrendered, or built - and how long it took.

The build threw a plain Error for every input it refused (a content or plan
gate, the variety contract, an evaluation too short), so the CLI printed a
stack and exited 1, as it does for a bug. A refusal now exits EXIT.refused
with its message alone; a crash keeps its stack and 1. A machine without
LibreOffice or poppler crashed at the render; it now finishes unrendered and
says what to install. And the stages that do not read each other's output run
side by side, each timed, beside the build's wall time.
"""
from __future__ import annotations

import contextlib
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from node_probe import HAS_RENDERER, NODE, requires_python_package, run_node

ROOT = Path(__file__).resolve().parents[2]
RUNTIME = ROOT / "skills" / "professional-slides" / "runtime"
BUILD = RUNTIME / "build-deck.mjs"
PYTHON = os.environ.get("RUNTIME_PYTHON") or sys.executable
sys.path.insert(0, str(RUNTIME / "emit"))


def spec(pages=3, **extra):
    return {"schema": "professional-slides.deck/v3", "id": "probe", "cover": {"title": "A probe deck"},
            "slides": [{"title": f"Page {i} states one finding in a sentence", "layout": "text",
                        "points": ["A point that says something about the finding and why it matters."]} for i in range(pages)], **extra}


class RefusalTests(unittest.TestCase):
    def setUp(self):
        if not NODE:
            self.skipTest("Node.js is not available")
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)

    def build(self, deck, *args, content=None, env=None):
        path = self.tmp / "probe.deck.json"
        path.write_text(deck if isinstance(deck, str) else json.dumps(deck), encoding="utf-8")
        if content is not None:
            (self.tmp / "probe.content.json").write_text(json.dumps(content), encoding="utf-8")
        return subprocess.run([NODE, str(BUILD), str(path), str(self.tmp / "out"), "--no-fetch", "--python", PYTHON, *args],
                              cwd=ROOT, capture_output=True, text=True, env=env)

    def test_a_refused_input_exits_refused_with_its_message_alone(self):
        run = self.build(spec(purpose="evaluation"), "--preflight")
        self.assertEqual(run.returncode, 2, run.stderr)
        self.assertIn("Refused (EVALUATION_TOO_SHORT)", run.stderr)
        self.assertNotIn("    at ", run.stderr)

    def test_a_rejected_content_plan_is_a_refusal_too(self):
        bad = {"schema": "professional-slides.content/v1", "id": "probe",
               "pages": [{"n": i, "claim": "Origins", "settles": {"kind": "qualitative", "what": "a premise"}, "adds": "", "highlight": ""}
                         for i in range(1, 11)]}
        run = self.build(spec(), "--preflight", content=bad)
        self.assertEqual(run.returncode, 2, run.stderr)
        self.assertIn("Refused (CONTENT_REJECTED): content gates rejected", run.stderr)
        self.assertNotIn("    at ", run.stderr)

    def test_a_crash_keeps_its_stack_and_exits_one(self):
        run = self.build("{ not json", "--preflight")
        self.assertEqual(run.returncode, 1)
        self.assertIn("SyntaxError", run.stderr)

    def test_the_refusal_codes_are_registered(self):
        # Read off the module's exported table, not its source; that a refusal
        # exits refused rather than crashing is the tests above.
        codes = run_node("""
import { REFUSAL_CODES } from './skills/professional-slides/runtime/build-deck.mjs';
console.log(JSON.stringify(Object.fromEntries(Object.entries(REFUSAL_CODES).map(([code, about]) => [code, String(about).trim().length > 0]))));
""")
        for code in ("CONTENT_REJECTED", "PLAN_REJECTED", "VARIETY_REJECTED", "STAGE_CONTRACT", "STAGE_MISSING", "TEXT_PLAN_CHANGED", "CONTENT_LOST"):
            self.assertTrue(codes.get(code), code)


@requires_python_package("pptx")
class BuildTimingTests(unittest.TestCase):
    def test_every_stage_is_timed_and_the_wall_time_recorded(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "probe.deck.json"
            path.write_text(json.dumps(spec()), encoding="utf-8")
            run = subprocess.run([NODE, str(BUILD), str(path), str(Path(tmp) / "out"), "--no-render", "--no-fetch", "--python", PYTHON],
                                 cwd=ROOT, capture_output=True, text=True)
            self.assertIn(run.returncode, (0, 2), run.stderr)
            timings = json.loads((Path(tmp) / "out" / "build-result.json").read_text())["timings"]
            for key in ("planMs", "preflightMs", "emitMs", "readbackMs", "wallMs"):
                self.assertIn(key, timings)
            # The scene's gates run beside the emitter, so the wall time is under the stages' sum.
            self.assertLess(timings["wallMs"], sum(v for k, v in timings.items() if k != "wallMs") + 1000)
            self.assertNotIn("totalMs", timings)


@requires_python_package("pptx")
class MissingRendererTests(unittest.TestCase):
    def test_render_says_what_is_missing_and_how_to_install_it(self):
        import render_pptx
        with mock.patch.object(render_pptx, "missing_tools", return_value={"pdftoppm": "brew install poppler"}):
            result = render_pptx.render(Path("deck.pptx"), Path("unused"))
        self.assertTrue(result["skipped"])
        self.assertEqual(result["missing"], ["pdftoppm"])
        self.assertIn("brew install poppler", result["message"])
        printed, said = io.StringIO(), io.StringIO()
        with mock.patch.object(render_pptx, "missing_tools", return_value={"soffice": "brew install --cask libreoffice"}), \
             contextlib.redirect_stdout(printed), contextlib.redirect_stderr(said):
            self.assertEqual(render_pptx.main(["deck.pptx", "unused"]), render_pptx.MISSING_EXIT)
        self.assertTrue(json.loads(printed.getvalue())["skipped"])
        self.assertIn("brew install --cask libreoffice", said.getvalue())

    @unittest.skipIf(shutil.which("pdftoppm", path="/usr/bin:/bin") is not None, "poppler is on the system path")
    def test_a_build_without_poppler_finishes_unrendered(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "probe.deck.json"
            path.write_text(json.dumps(spec()), encoding="utf-8")
            # Poppler off the path; python-pptx kept importable where a user install put it.
            base = subprocess.run([PYTHON, "-c", "import site; print(site.getuserbase())"], capture_output=True, text=True).stdout.strip()
            env = {"PATH": "/usr/bin:/bin", "HOME": tmp, "RUNTIME_NODE_MODULES": str(ROOT / "node_modules"),
                   **({"PYTHONUSERBASE": base} if base else {})}
            run = subprocess.run([NODE, str(BUILD), str(path), str(Path(tmp) / "out"), "--no-fetch", "--python", PYTHON],
                                 cwd=ROOT, capture_output=True, text=True, env=env)
            self.assertIn(run.returncode, (0, 2), run.stderr)
            self.assertNotIn("    at ", run.stderr)
            result = json.loads((Path(tmp) / "out" / "build-result.json").read_text())
            self.assertIn("pdftoppm", result["renderSkipped"])
            self.assertIn(result["status"], ("built-unrendered", "built-with-blockers"))
            self.assertIsNone(result.get("gates"))


class ProfileLockTests(unittest.TestCase):
    def test_a_held_profile_sends_the_next_render_to_a_private_one(self):
        import render_pptx
        with tempfile.TemporaryDirectory() as tmp, mock.patch.dict(os.environ, {"PROFESSIONAL_SLIDES_LO_PROFILE": str(Path(tmp) / "kept")}):
            with render_pptx.lo_profile(Path(tmp) / "one") as (first, shared):
                with render_pptx.lo_profile(Path(tmp) / "two") as (second, shared_too):
                    self.assertTrue(shared)
                    self.assertFalse(shared_too)
                    self.assertEqual(second, Path(tmp) / "two" / "profile")
                self.assertEqual(first, Path(tmp) / "kept" / "profile")
            with render_pptx.lo_profile(Path(tmp) / "three") as (again, shared_again):
                self.assertTrue(shared_again)  # released, so the next render keeps it
                self.assertEqual(again, first)


@unittest.skipUnless(HAS_RENDERER, "LibreOffice and poppler are not installed")
@requires_python_package("pptx", "PIL", "pypdf")
class RenderCanvasTests(unittest.TestCase):
    def test_pages_render_at_the_gates_canvas_and_the_sheets_come_separately(self):
        import render_pptx
        from PIL import Image
        from pptx import Presentation
        from pptx.util import Emu
        with tempfile.TemporaryDirectory() as tmp:
            deck = Presentation()
            deck.slide_width, deck.slide_height = Emu(12192000), Emu(6858000)
            for _ in range(2):
                deck.slides.add_slide(deck.slide_layouts[6])
            pptx = Path(tmp) / "canvas.pptx"
            deck.save(str(pptx))
            out = Path(tmp) / "rendered"
            result = render_pptx.render(pptx, out)
            self.assertEqual(len(result["renders"]), 2)
            for png in result["renders"]:
                with Image.open(png) as image:
                    self.assertEqual(image.size, (1280, 720))
            self.assertNotIn("montage", result)
            sheets = render_pptx.sheets(out)
            self.assertTrue(Path(sheets["montage"]).is_file())
            self.assertEqual(len(sheets["spreads"]), 1)
            with Image.open(sheets["montage"]) as montage:
                self.assertEqual(montage.size, (4 * 640 + 5 * 12, 360 + 2 * 12))


class OutcomeTests(unittest.TestCase):
    def test_only_blockers_fail_a_build_and_each_is_named(self):
        """PR #4 review: advisories failed a build, and a build held back named no blocker."""
        # Advisory-only builds are `built`; a build held back by text lost from
        # the rendered pages says so, rather than showing the advisory counts.
        result = run_node('''
import {buildOutcome} from './skills/professional-slides/runtime/build-deck.mjs';
const advisory = {code:'UNANNOTATED',severity:'advisory',slide:4};
const passed = {passed:true,findings:[advisory,{code:'DECK_CRAFT',severity:'advisory'}]};
const clean = buildOutcome({preflight:passed,gates:passed,readback:{accepted:true},textCoverage:{accepted:true,findings:[]}},{render:true});
const lost = buildOutcome({preflight:passed,gates:passed,readback:{accepted:true},textCoverage:{accepted:false,findings:[{id:'p13',code:'TEXT_EXPORT_LOST',text:'787',severity:'blocking'}]}},{render:true});
const gated = buildOutcome({preflight:{passed:false,findings:[{code:'TITLE_WORDS',severity:'blocker',slide:3},advisory]},gates:{passed:false,findings:[{code:'TITLE_WORDS',severity:'blocker',slide:3},advisory]},readback:{accepted:true}},{render:true});
const unrendered = buildOutcome({preflight:passed,readback:{accepted:true}},{render:false});
console.log(JSON.stringify({clean,lost,gated,unrendered}));
''')
        self.assertEqual(result['clean'], {'status': 'built', 'blockers': [], 'advisories': {'DECK_CRAFT': 1, 'UNANNOTATED': 1}})
        self.assertEqual(result['lost']['status'], 'built-with-blockers')
        self.assertEqual([(b['source'], b['code'], b['id']) for b in result['lost']['blockers']], [('rendered text', 'TEXT_EXPORT_LOST', 'p13')])
        self.assertEqual(result['gated']['status'], 'built-with-blockers')
        self.assertEqual([b['code'] for b in result['gated']['blockers']], ['TITLE_WORDS'])
        self.assertEqual(result['gated']['advisories'], {'UNANNOTATED': 1})
        self.assertEqual(result['unrendered']['status'], 'built-unrendered')


@requires_python_package("pptx")
class CoverageRefusalTests(unittest.TestCase):
    def test_coverage_blocks_build_and_delivery_and_python_option_is_used(self):
        """PR #4 review: a deck missing evidence for a criterion built and delivered, and --python lost to RUNTIME_PYTHON."""
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory);spec=root/'spec.json';out=root/'out'
            spec.write_text(json.dumps({'schema':'professional-slides.deck/v3','id':'coverage','cover':{'title':'Decision'},'criteria':['housing'],'slides':[]}))
            command=[NODE,str(RUNTIME/'build-deck.mjs'),str(spec),str(out),'--python',sys.executable]
            # Explicit interpreter must win over an unusable default.
            env={**os.environ,'RUNTIME_PYTHON':'/does-not-exist'}
            pre=subprocess.run(command+['--preflight'],capture_output=True,text=True,env=env)
            self.assertEqual(pre.returncode,2,pre.stderr)
            self.assertEqual(json.loads(pre.stdout)['status'],'preflight-findings')
            build=subprocess.run(command+['--no-render'],capture_output=True,text=True,env=env)
            self.assertEqual(build.returncode,2,build.stderr)
            self.assertEqual(json.loads(build.stdout)['status'],'built-with-blockers')
            self.assertIn('MISSING_EVIDENCE',json.loads(build.stdout)['blockers']['byCode'])
            report=json.loads((out/'preflight-gates.json').read_text())
            self.assertEqual(report['countsByCode']['MISSING_EVIDENCE'],1)
            review=root/'review.json';review.write_text(json.dumps({'accepted':True,'summary':'Accepted for the purpose of proving gates cannot be bypassed.','findings':[]}))
            delivered=subprocess.run([NODE,str(RUNTIME/'deliver-deck.mjs'),str(spec),str(out),'--skip-build','--review',str(review)],capture_output=True,text=True)
            self.assertEqual(delivered.returncode,2,delivered.stderr)
            self.assertFalse((out/'coverage-DELIVERED.pptx').exists())
            self.assertIn('MISSING_EVIDENCE',[x['code'] for x in json.loads(delivered.stdout)['blockers']])
            invalid=subprocess.run(command[:-2]+['--python'],capture_output=True,text=True)
            self.assertEqual(invalid.returncode,1)
            self.assertIn('--python requires an executable',invalid.stderr)


if __name__ == "__main__":
    unittest.main()
