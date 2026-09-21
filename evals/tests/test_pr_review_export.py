"""Read saved OOXML to verify paint order, alpha, masks and theme typography."""
import base64
import io
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn
from node_probe import ROOT

sys.path.insert(0,str(ROOT/'skills/professional-slides/runtime/emit'))
from emit_pptx import Emitter


def scene_fixture():
    frame={'x':100,'y':180,'width':800,'height':400}
    image=io.BytesIO();Image.new('RGB',(30,30),'blue').save(image,format='PNG')
    def rect(name,color,opacity=1):
        return {'id':name,'type':'rect','role':'section-surface','frame':frame,'style':{'fill':color,'opacity':opacity}}
    return {'typography':{'body':'Georgia','display':'Arial'},'tokens':{},'slides':[{'id':'one','tokens':{},'nodes':[
        rect('background','#EEEEEE'),
        {'id':'plot','type':'rect','role':'chart-mark','frame':frame,'data':{'componentInstance':'chart'},'style':{'fill':'#000000'}},
        rect('later-overlay','#FF0000',.72),
        {'id':'portrait','type':'image','frame':{'x':950,'y':180,'width':80,'height':80},'data':{'dataUri':'data:image/png;base64,'+base64.b64encode(image.getvalue()).decode(),'circular':True,'crop':{'left':.1,'right':.1}}},
    ],'componentInstances':[{'instanceId':'chart','component':'chart.column','frame':frame,'nativeChart':{
        'type':'column','frame':frame,'categories':['A','B'],'series':[{'name':'Sales','values':[1,2]}],'dataLabels':True,'legend':True,'valueFormat':{'decimals':1}}}]}]}


class ExportReviewTests(unittest.TestCase):
    def test_canvas_survives_export_and_readback_rejects_drift(self):
        from readback_pptx import readback
        from pptx.dml.color import RGBColor
        scene = {'tokens': {'color.canvas': {'kind': 'color', 'value': '#FFF9F0'}}, 'slides': [
            {'id': 'warm', 'tokens': {}, 'nodes': []},
            {'id': 'inverse', 'tokens': {'color.canvas': {'kind': 'color', 'value': '#18212B'}}, 'nodes': []}
        ]}
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'canvas.pptx'
            Emitter(scene).run(path)
            saved = Presentation(path)
            self.assertEqual(str(saved.slides[0].background.fill.fore_color.rgb), 'FFF9F0')
            self.assertEqual(str(saved.slides[1].background.fill.fore_color.rgb), '18212B')
            self.assertTrue(readback(scene, path)['accepted'])
            saved.slides[0].background.fill.fore_color.rgb = RGBColor(255, 255, 255)
            saved.save(path)
            report = readback(scene, path)
            self.assertFalse(report['accepted'])
            self.assertEqual(report['findings'][0]['code'], 'CANVAS_COLOR')

    @classmethod
    def setUpClass(cls):
        cls.tmp=tempfile.TemporaryDirectory();cls.path=Path(cls.tmp.name)/'review.pptx'
        Emitter(scene_fixture()).run(cls.path)
        cls.prs=Presentation(cls.path)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_native_chart_sits_between_earlier_and_later_surfaces(self):
        names=[s.name for s in self.prs.slides[0].shapes]
        self.assertLess(names.index('ps:background'),names.index('ps:chart:chart'))
        self.assertLess(names.index('ps:chart:chart'),names.index('ps:later-overlay'))

    def test_rectangle_fill_keeps_alpha(self):
        shape=next(s for s in self.prs.slides[0].shapes if s.name=='ps:later-overlay')
        alpha=shape._element.find('.//'+qn('a:alpha'))
        self.assertEqual(alpha.get('val'),'72000')

    def test_picture_keeps_oval_geometry_and_crop(self):
        pic=next(s for s in self.prs.slides[0].shapes if s.name=='ps:portrait')
        self.assertEqual(pic._element.spPr.find(qn('a:prstGeom')).get('prst'),'ellipse')
        self.assertAlmostEqual(pic.crop_left,.1)
        self.assertAlmostEqual(pic.crop_right,.1)

    def test_theme_keeps_distinct_display_and_body_fonts(self):
        with zipfile.ZipFile(self.path) as archive:
            root=etree.fromstring(archive.read('ppt/theme/theme1.xml'))
        ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
        self.assertEqual(root.find('.//a:majorFont/a:latin',ns).get('typeface'),'Arial')
        self.assertEqual(root.find('.//a:minorFont/a:latin',ns).get('typeface'),'Georgia')

    def test_native_chart_inherits_resolved_body_font(self):
        chart=next(s.chart for s in self.prs.slides[0].shapes if s.has_chart)
        self.assertEqual(chart.font.name,'Georgia')
        self.assertEqual(chart.plots[0].data_labels.number_format,'0.0')

class ZeroStackLabelTests(unittest.TestCase):
    def test_native_line_values_stay_above_markers_after_save(self):
        from pptx.enum.chart import XL_LABEL_POSITION
        scene = scene_fixture()
        spec = scene['slides'][0]['componentInstances'][0]['nativeChart']
        spec.update(type='line', categories=['A', 'B', 'C'],
                    series=[{'name': 'Cash flow', 'values': [1.62, 1.94, 2.37]}])
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'line.pptx'
            Emitter(scene).run(path)
            chart = next(s.chart for s in Presentation(path).slides[0].shapes if s.has_chart)
            self.assertEqual(chart.plots[0].data_labels.position, XL_LABEL_POSITION.ABOVE)
            self.assertEqual(tuple(chart.series[0].values), (1.62, 1.94, 2.37))

    def test_zero_stack_point_keeps_data_but_suppresses_native_label(self):
        scene=scene_fixture()
        spec=scene['slides'][0]['componentInstances'][0]['nativeChart']
        spec.update(type='stacked-bar',categories=['A','B'],series=[{'name':'First','values':[0,2]},{'name':'Second','values':[4,3]}])
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'zero.pptx';Emitter(scene).run(path)
            prs=Presentation(path)
            chart=next(s.chart for s in prs.slides[0].shapes if s.has_chart)
            self.assertEqual(chart.series[0].values[0],0)
            labels=chart.series[0]._element.findall('.//'+qn('c:dLbl'))
            self.assertTrue(any(l.find(qn('c:idx')).get('val')=='0' and l.find(qn('c:delete')).get('val')=='1' for l in labels))
