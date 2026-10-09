"""The saved file is a function of its scene, and its labels read on their marks.

The review binds itself to the .pptx's hash. Two builds of one scene wrote
byte-identical scene.json and different files: python-pptx stamps every zip
entry with the clock and writes each native chart's workbook through
XlsxWriter, which stamps its creation date, so every review was void for the
next build of the same deck.

And a stacked segment's label took white wherever the fill's luminance fell
under 0.45 - white at 2.3:1 on a green where ink reads at 9:1 - while the pie
beside it, and the scene, chose whichever of the two reads better.
"""
from __future__ import annotations

import io
import json
import os
import subprocess
import sys
import tempfile
import time
import unittest
import zipfile
from pathlib import Path

from node_probe import RUNTIME_PYTHON, requires_python_package, run_node

ROOT = Path(__file__).resolve().parents[2]
EMIT = ROOT / "skills" / "professional-slides" / "runtime" / "emit"
PYTHON = RUNTIME_PYTHON or sys.executable
sys.path.insert(0, str(EMIT))


def chart_scene():
    """One slide with a native stacked chart and a freeform, the parts that stamp dates and cost time."""
    return json.loads(run_node('''
import { composeAll } from './skills/professional-slides/runtime/compose-all.mjs';
const spec = { schema: 'professional-slides.deck/v3', id: 'det', cover: { title: 'Determinism' }, slides: [
  { id: 's1', title: 'Segments stack to the whole in every year', exhibit: { type: 'chart.stacked-column', categories: ['2022', '2023', '2024'],
    series: [{ name: 'A', values: [3, 4, 5] }, { name: 'B', values: [2, 2, 3] }] } },
  { id: 's2', title: 'A second chart keeps the workbook count honest', exhibit: { type: 'chart.column', categories: ['A', 'B', 'C', 'D'],
    series: [{ name: 'x', values: [4, 3, 2, 1] }] } }] };
console.log(JSON.stringify(JSON.stringify(composeAll(spec, process.cwd()).deck)));
'''))


@requires_python_package("pptx")
class DeterministicEmitTests(unittest.TestCase):
    def test_two_builds_of_one_scene_write_the_same_bytes(self):
        from emit_pptx import Emitter
        scene = chart_scene()
        with tempfile.TemporaryDirectory() as tmp:
            first, second = Path(tmp) / "a.pptx", Path(tmp) / "b.pptx"
            Emitter(json.loads(json.dumps(scene))).run(first)
            # A zip records time to two seconds; wait past a tick so the clock
            # would show if anything still read it.
            time.sleep(2.1)
            Emitter(json.loads(json.dumps(scene))).run(second)
            self.assertEqual(first.read_bytes(), second.read_bytes())
            with zipfile.ZipFile(first) as package:
                workbooks = [n for n in package.namelist() if n.startswith("ppt/embeddings/") and n.endswith(".xlsx")]
                self.assertTrue(workbooks, "the scene carries native charts")
                self.assertEqual({info.date_time for info in package.infolist()}, {(1980, 1, 1, 0, 0, 0)})
                self.assertIn(b"1980-01-01T00:00:00Z", package.read("docProps/core.xml"))
                with zipfile.ZipFile(io.BytesIO(package.read(workbooks[0]))) as book:
                    self.assertIn(b"1980-01-01T00:00:00Z", book.read("docProps/core.xml"))
                    self.assertEqual({info.date_time for info in book.infolist()}, {(1980, 1, 1, 0, 0, 0)})

    def test_the_file_still_opens_as_a_deck_with_its_charts(self):
        from emit_pptx import Emitter
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "a.pptx"
            Emitter(chart_scene()).run(path)
            deck = Presentation(str(path))
            charts = [shape for slide in deck.slides for shape in slide.shapes if getattr(shape, "has_chart", False)]
            self.assertEqual(len(charts), 2)
            self.assertEqual(len(list(charts[0].chart.plots[0].series)), 2)


@requires_python_package("pptx")
class SourceDateEpochTests(unittest.TestCase):
    """SOURCE_DATE_EPOCH=0, the reproducible-builds default, is 1970: a zip
    entry cannot record a date before 1980 (nor after 2107), so the archive
    raised and no deck was written. The entries take the nearest date a zip
    holds; the package's and the workbooks' XML dates keep the epoch itself."""

    @classmethod
    def setUpClass(cls):
        cls.tmp_dir = tempfile.TemporaryDirectory()
        cls.scene = Path(cls.tmp_dir.name) / "scene.json"
        cls.scene.write_text(json.dumps(chart_scene()))

    @classmethod
    def tearDownClass(cls):
        cls.tmp_dir.cleanup()

    def emit(self, epoch: str, name: str) -> Path:
        # The epoch is read when the emitter loads, so each build is its own process.
        out = Path(self.tmp_dir.name) / name
        run = subprocess.run([PYTHON, str(EMIT / "emit_pptx.py"), str(self.scene), str(out)], capture_output=True, text=True,
                             env={**os.environ, "SOURCE_DATE_EPOCH": epoch}, timeout=240)
        self.assertEqual(run.returncode, 0, run.stderr)
        return out

    def assert_stamped(self, epoch: str, xml_date: bytes, zip_date: tuple):
        first, second = self.emit(epoch, f"{epoch}-a.pptx"), self.emit(epoch, f"{epoch}-b.pptx")
        self.assertEqual(first.read_bytes(), second.read_bytes())
        with zipfile.ZipFile(first) as package:
            self.assertEqual({info.date_time for info in package.infolist()}, {zip_date})
            self.assertIn(xml_date, package.read("docProps/core.xml"))
            workbooks = [n for n in package.namelist() if n.startswith("ppt/embeddings/") and n.endswith(".xlsx")]
            self.assertTrue(workbooks, "the scene carries native charts")
            for workbook in workbooks:
                with zipfile.ZipFile(io.BytesIO(package.read(workbook))) as book:
                    self.assertIn(xml_date, book.read("docProps/core.xml"))
                    self.assertEqual({info.date_time for info in book.infolist()}, {zip_date})

    def test_an_epoch_before_1980_stamps_the_entries_at_1980(self):
        self.assert_stamped("0", b"1970-01-01T00:00:00Z", (1980, 1, 1, 0, 0, 0))

    def test_an_epoch_after_2107_stamps_the_entries_at_2107(self):
        self.assert_stamped(str(2 ** 33), b"2242-03-16T12:56:32Z", (2107, 12, 31, 23, 59, 58))

    def test_an_epoch_inside_the_range_stamps_both_alike(self):
        self.assert_stamped("1700000000", b"2023-11-14T22:13:20Z", (2023, 11, 14, 22, 13, 20))

    def test_a_revision_assembled_at_an_epoch_before_1980_stamps_its_new_entries_at_1980(self):
        # The assembler writes the parts it rewrites - the slide order, the content types - at the epoch too, held to the same range.
        source = self.emit("0", "assemble-source.pptx")
        plan = Path(self.tmp_dir.name) / "assemble-plan.json"
        plan.write_text(json.dumps({"source": source.name, "sha256": None, "composed": None, "slides": [{"carry": 2}, {"carry": 1}]}))
        out = Path(self.tmp_dir.name) / "assembled.pptx"
        run = subprocess.run([PYTHON, str(EMIT / "assemble_pptx.py"), str(plan), str(out)], capture_output=True, text=True,
                             env={**os.environ, "SOURCE_DATE_EPOCH": "0"}, timeout=240)
        self.assertEqual(run.returncode, 0, run.stderr)
        with zipfile.ZipFile(out) as package:
            self.assertEqual({info.date_time for info in package.infolist()}, {(1980, 1, 1, 0, 0, 0)})


@requires_python_package("pptx")
class OnFillParityTests(unittest.TestCase):
    def test_the_emitter_sets_labels_as_the_scene_does_on_every_palette_colour(self):
        from emit_pptx import on_fill
        result = run_node('''
import { onFill, tokenValue, TOKENS, THEME_SLOT_TOKENS } from './skills/professional-slides/runtime/core.mjs';
import { PALETTES, resolvePalette } from './skills/professional-slides/runtime/palettes.mjs';
import { DESIGN_SYSTEMS } from './skills/professional-slides/runtime/design-systems.mjs';
import { withDesignTokens } from './skills/professional-slides/runtime/design-context.mjs';
const hex = (v) => typeof v === 'string' && /^#[0-9A-Fa-f]{6}$/.test(v);
const cases = [];
const collect = (tokens, colors, source) => withDesignTokens(tokens, () => {
  const on = tokenValue('color.onPrimary'), ink = tokenValue('color.ink');
  for (const [key, value] of Object.entries(colors)) if (key.startsWith('color.') && hex(value))
    cases.push({ source, key, fill: value, on, ink, chosen: tokenValue(onFill(value)) });
});
for (const id of Object.keys(PALETTES)) { const tokens = resolvePalette(id, TOKENS, THEME_SLOT_TOKENS).tokens; collect(tokens, Object.fromEntries(Object.entries(tokens).map(([k, v]) => [k, v.value])), id); }
const walk = (value, out) => { if (value && typeof value === 'object') for (const [k, v] of Object.entries(value)) { if (k.startsWith('color.') && hex(v)) out[k] = v; else walk(v, out); } return out; };
for (const [name, system] of Object.entries(DESIGN_SYSTEMS)) {
  const colors = walk(system, {});
  const tokens = { ...TOKENS, ...Object.fromEntries(Object.entries(colors).filter(([k]) => TOKENS[k]).map(([k, v]) => [k, { ...TOKENS[k], value: v }])) };
  collect(tokens, colors, name);
}
console.log(JSON.stringify(cases));
''')
        self.assertGreater(len(result), 60)
        mismatched = [c for c in result if on_fill(c["fill"], c["on"], c["ink"]).upper() != c["chosen"].upper()]
        self.assertEqual(mismatched, [])

    def test_a_mid_tone_fill_takes_the_ink(self):
        from color import contrast, on_fill
        # The accent green that took white at 2.3:1 under the luminance cut.
        self.assertEqual(on_fill("#86BC25", "#FFFFFF", "#000000"), "#000000")
        self.assertGreater(contrast("#86BC25", "#000000"), 9)
        self.assertLess(contrast("#86BC25", "#FFFFFF"), 2.5)
        self.assertEqual(on_fill("#051C2C", "#FFFFFF", "#051C2C"), "#FFFFFF")

    def test_one_rule_serves_both_label_sites(self):
        import color
        import emit_pptx
        from pptx import Presentation
        self.assertIs(emit_pptx.on_fill, color.on_fill)  # one definition (color.py)
        # Both sites, read off the saved file: a stacked segment's label and a
        # pie slice's share on the mid-tone green take the ink, as the scene does,
        # where the old luminance cut gave both white at 2.3:1.
        frame = {"x": 60, "y": 160, "width": 1000, "height": 460}
        tokens = {"color.chartSeries1": {"kind": "color", "value": "#86BC25"}, "color.chartSeries2": {"kind": "color", "value": "#86BC25"},
                  "color.onPrimary": {"kind": "color", "value": "#FFFFFF"}, "color.ink": {"kind": "color", "value": "#000000"}}
        specs = {"stacked-bar": {"type": "stacked-bar", "frame": frame, "categories": ["X", "Y"], "dataLabels": True,
                                 "series": [{"name": "A", "values": [3, 5]}, {"name": "B", "values": [2, 4]}]},
                 "pie": {"type": "pie", "frame": frame, "categories": ["X", "Y"], "dataLabels": True, "series": [{"name": "Share", "values": [60, 40]}]}}
        scene = {"typography": {"body": "Arial", "display": "Arial"}, "tokens": tokens, "slides": [
            {"id": kind, "tokens": {}, "componentInstances": [{"instanceId": kind, "component": "chart", "frame": frame, "nativeChart": spec}],
             "nodes": [{"id": f"{kind}-plot", "type": "rect", "role": "chart-mark", "frame": frame, "data": {"componentInstance": kind}, "style": {"fill": "#000000"}}]}
            for kind, spec in specs.items()]}
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "labels.pptx"
            emit_pptx.Emitter(scene).run(path)
            stacked, pie = [next(shape.chart for shape in slide.shapes if shape.has_chart) for slide in Presentation(path).slides]
            segment = str(stacked.series[0].data_labels.font.color.rgb)
            share = str(pie.series[0].points[0].data_label.font.color.rgb)
        self.assertEqual(segment, "000000")
        self.assertEqual(share, "000000")


@requires_python_package("pptx")
class FreeformTests(unittest.TestCase):
    def test_a_cached_offset_draws_the_same_path(self):
        from emit_pptx import CachedFreeform
        from pptx import Presentation
        from pptx.shapes.freeform import FreeformBuilder
        from pptx.util import Emu
        points = [(Emu(1000 + (i * 37) % 900), Emu(2000 + (i * 53) % 700)) for i in range(400)]

        def draw(builder_class):
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            builder = builder_class.new(slide.shapes, points[0][0], points[0][1], 1.0, 1.0)
            builder.add_line_segments(points[1:], close=True)
            shape = builder.convert_to_shape(0, 0)
            return (shape.left, shape.top, shape.width, shape.height), shape._element.xml.split("<a:custGeom>")[1]

        self.assertEqual(draw(CachedFreeform), draw(FreeformBuilder))


if __name__ == "__main__":
    unittest.main()
