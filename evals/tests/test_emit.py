"""The PPTX emitter (runtime/emit/emit_pptx.py), read back from the saved OOXML.

Paint order, alpha, picture masks and theme typography; native chart labels,
axes and point colours as PowerPoint will show them. Each test emits a scene
with the real emitter and opens the file it wrote.
"""
import base64
import io
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from node_probe import ROOT, requires_python_package, run_node

# The emitter's dependencies are optional at test time (see requirements.txt).
# Imported defensively so the module still loads and its classes skip with a
# message naming what is missing, rather than failing collection.
sys.path.insert(0, str(ROOT / 'skills/professional-slides/runtime/emit'))
try:
    from lxml import etree
    from PIL import Image
    from pptx import Presentation
    from pptx.oxml.ns import qn
    from emit_pptx import Emitter
except ImportError:  # pragma: no cover - exercised only without the packages
    etree = Image = Presentation = qn = Emitter = None

needs_emitter = requires_python_package('pptx', 'lxml', 'PIL')
needs_pptx = requires_python_package('pptx')


def scene_fixture():
    frame={'x':100,'y':180,'width':800,'height':400}
    image=io.BytesIO();Image.new('RGB',(30,30),'blue').save(image,format='PNG')
    def rect(name,color,opacity=1):
        return {'id':name,'type':'rect','role':'section-surface','frame':frame,'style':{'fill':color,'opacity':opacity}}
    return {'typography':{'body':'Georgia','display':'Arial'},'tokens':{},'slides':[{'id':'one','tokens':{},'nodes':[
        rect('background','#EEEEEE'),
        {'id':'plot','type':'rect','role':'chart-mark','frame':frame,'data':{'componentInstance':'chart'},'style':{'fill':'#000000'}},
        rect('later-overlay','#FF0000',.72),
        {'id':'portrait','type':'image','frame':{'x':950,'y':180,'width':80,'height':80},'data':{'dataUri':'data:image/png;base64,'+base64.b64encode(image.getvalue()).decode(),'circular':True,'crop':{'left':.1,'right':.1}}},
    ],'componentInstances':[{'instanceId':'chart','component':'chart.column','frame':frame,'nativeChart':{
        'type':'column','frame':frame,'categories':['A','B'],'series':[{'name':'Sales','values':[1,2]}],'dataLabels':True,'legend':True,'valueFormat':{'decimals':1}}}]}]}


@needs_emitter
class OoxmlFidelityTests(unittest.TestCase):
    """What the saved file paints, in what order, and in which faces."""

    @classmethod
    def setUpClass(cls):
        cls.tmp=tempfile.TemporaryDirectory();cls.path=Path(cls.tmp.name)/'review.pptx'
        Emitter(scene_fixture()).run(cls.path)
        cls.prs=Presentation(cls.path)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_native_chart_sits_between_earlier_and_later_surfaces(self):
        """PR #4 review: a native chart was painted above surfaces the scene drew after it."""
        names=[s.name for s in self.prs.slides[0].shapes]
        self.assertLess(names.index('ps:background'),names.index('ps:chart:chart'))
        self.assertLess(names.index('ps:chart:chart'),names.index('ps:later-overlay'))

    def test_rectangle_fill_keeps_alpha(self):
        """PR #4 review: a translucent overlay was saved opaque."""
        shape=next(s for s in self.prs.slides[0].shapes if s.name=='ps:later-overlay')
        alpha=shape._element.find('.//'+qn('a:alpha'))
        self.assertEqual(alpha.get('val'),'72000')

    def test_picture_keeps_oval_geometry_and_crop(self):
        """PR #4 review: a circular portrait lost its oval mask and its crop."""
        pic=next(s for s in self.prs.slides[0].shapes if s.name=='ps:portrait')
        self.assertEqual(pic._element.spPr.find(qn('a:prstGeom')).get('prst'),'ellipse')
        self.assertAlmostEqual(pic.crop_left,.1)
        self.assertAlmostEqual(pic.crop_right,.1)

    def test_theme_keeps_distinct_display_and_body_fonts(self):
        """PR #4 review: the theme collapsed the display and body faces into one."""
        with zipfile.ZipFile(self.path) as archive:
            root=etree.fromstring(archive.read('ppt/theme/theme1.xml'))
        ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
        self.assertEqual(root.find('.//a:majorFont/a:latin',ns).get('typeface'),'Arial')
        self.assertEqual(root.find('.//a:minorFont/a:latin',ns).get('typeface'),'Georgia')

    def test_native_chart_inherits_resolved_body_font(self):
        """PR #4 review: a native chart ignored the resolved body face and its number format."""
        chart=next(s.chart for s in self.prs.slides[0].shapes if s.has_chart)
        self.assertEqual(chart.font.name,'Georgia')
        self.assertEqual(chart.plots[0].data_labels.number_format,'0.0')


@needs_emitter
class NativeChartLabelTests(unittest.TestCase):
    """Native data labels say what the scene's labels said, where they said it."""

    def test_native_line_endpoint_overrides_preserve_values_and_point_policy(self):
        """PR #4 review: a native line's end label override dropped its value and the series' point-label policy."""
        for point_labels in (True, False):
            scene = scene_fixture()
            scene['tokens']['color.chartSeries1'] = {'kind': 'color', 'value': '#06202E'}
            spec = scene['slides'][0]['componentInstances'][0]['nativeChart']
            spec.update(type='line', categories=['0', '12', '24', '36'],
                        series=[{'name': 'Monthly support', 'values': [0, 6, 12, 18]}],
                        endLabels=True, pointDataLabels=point_labels)
            with tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp) / 'line.pptx'
                Emitter(scene).run(path)
                chart = next(s.chart for s in Presentation(path).slides[0].shapes if s.has_chart)
                labels = chart.series[0]._element.find(qn('c:dLbls'))
                self.assertEqual(labels.find(qn('c:showVal')).get('val'), '1' if point_labels else '0')
                # The first point, under a rising line, carries its own
                # placement override too; the endpoint is the one at idx 3.
                endpoint = next(l for l in labels.findall(qn('c:dLbl')) if l.find(qn('c:idx')).get('val') == '3')
                self.assertEqual(endpoint.find(qn('c:idx')).get('val'), '3')
                self.assertEqual(''.join(t.text for t in endpoint.findall('.//' + qn('a:t'))), 'Monthly support 18.0')

    def test_native_line_values_stay_above_markers_after_save(self):
        """PR #4 review: native line values sat on their markers after save."""
        from pptx.enum.chart import XL_LABEL_POSITION
        scene = scene_fixture()
        spec = scene['slides'][0]['componentInstances'][0]['nativeChart']
        spec.update(type='line', categories=['A', 'B', 'C'],
                    series=[{'name': 'Cash flow', 'values': [1.62, 1.94, 2.37]}])
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'line.pptx'
            Emitter(scene).run(path)
            chart = next(s.chart for s in Presentation(path).slides[0].shapes if s.has_chart)
            self.assertEqual(chart.plots[0].data_labels.position, XL_LABEL_POSITION.ABOVE)
            self.assertEqual(tuple(chart.series[0].values), (1.62, 1.94, 2.37))

    def test_signed_native_line_categories_stay_below_plot_after_save(self):
        """PR #4 review: a signed native line's category labels rode the zero line through the plot."""
        from pptx.enum.chart import XL_TICK_LABEL_POSITION
        scene = scene_fixture()
        spec = scene['slides'][0]['componentInstances'][0]['nativeChart']
        spec.update(type='line', categories=['0', '1', '2', '3', '4', '5'],
                    series=[{'name': 'Contribution', 'values': [-250, -170, -90, -10, 70, 150]}])
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'signed-line.pptx'
            Emitter(scene).run(path)
            chart = next(s.chart for s in Presentation(path).slides[0].shapes if s.has_chart)
            self.assertEqual(chart.category_axis.tick_label_position, XL_TICK_LABEL_POSITION.LOW)
            self.assertEqual(tuple(chart.series[0].values), (-250, -170, -90, -10, 70, 150))

    def test_zero_stack_point_keeps_data_but_suppresses_native_label(self):
        """PR #4 review: a zero segment in a native stack printed a "0" label."""
        scene=scene_fixture()
        spec=scene['slides'][0]['componentInstances'][0]['nativeChart']
        spec.update(type='stacked-bar',categories=['A','B'],series=[{'name':'First','values':[0,2]},{'name':'Second','values':[4,3]}])
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'zero.pptx';Emitter(scene).run(path)
            prs=Presentation(path)
            chart=next(s.chart for s in prs.slides[0].shapes if s.has_chart)
            self.assertEqual(chart.series[0].values[0],0)
            labels=chart.series[0]._element.findall('.//'+qn('c:dLbl'))
            self.assertTrue(any(l.find(qn('c:idx')).get('val')=='0' and l.find(qn('c:delete')).get('val')=='1' for l in labels))


@needs_pptx
class NativeChartExportTests(unittest.TestCase):
    """A chart compiled to go native keeps its axis, colours and shares in the saved file."""

    def test_native_chart_axis_and_point_colors_survive_saved_pptx(self):
        """PR #4 follow-up: native charts lost the value-axis choice and per-point colours on save."""
        with tempfile.TemporaryDirectory() as tmp:
            scene = run_node('''
import {compileDeck} from './skills/professional-slides/runtime/core.mjs';
import {REGISTRY} from './skills/professional-slides/runtime/registry.mjs';
const cases=[];
for (const type of ['column','bar']) for (const labels of [true,false]) for (const axis of [true,false])
 for (const extra of [{},{colorIndices:[2]},{highlights:[{category:'B',style:'bar'}]}])
  cases.push({component:'chart.'+type,props:{categories:['A','B'],series:[{name:'Revenue',values:[20,40]}],dataLabels:labels,showValueAxis:axis,...extra}});
const slides=cases.map((c,i)=>({id:'s'+i,composition:{nodeType:'component',id:'plot',...c,frame:{x:60,y:160,width:1000,height:460}}}));
console.log(JSON.stringify(compileDeck({id:'charts',slides},REGISTRY)));
''')
            file = Path(tmp) / 'charts.pptx'; Emitter(scene).run(file); prs = Presentation(file)
            for saved, source in zip(prs.slides, scene['slides']):
                chart = next(shape.chart for shape in saved.shapes if shape.has_chart)
                native = source['componentInstances'][0]['nativeChart']
                self.assertEqual(chart.value_axis.visible, native['showValueAxis'])
                self.assertEqual(chart.value_axis.visible, any(n['role'] == 'axis-label' for n in source['nodes']))
                marks = [n for n in source['nodes'] if n['role'] == 'chart-mark']
                series = chart.series[0]
                for point, mark in zip(series.points, marks):
                    color = point.format.fill.fore_color.rgb if point.format.fill.type else series.format.fill.fore_color.rgb
                    self.assertEqual(str(color), mark['style']['fill']['value'].lstrip('#').upper())

    @needs_pptx
    def test_native_pie_prints_the_scenes_shares_inside(self):
        """Emirates deck: a native pie printed raw values outside the rim instead of the scene's shares inside."""
        scene = run_node("""
import {compileDeck} from './skills/professional-slides/runtime/core.mjs';
import {REGISTRY} from './skills/professional-slides/runtime/registry.mjs';
const slides=[{id:'p',composition:{nodeType:'component',id:'pie',component:'chart.pie',props:{labels:['International','Domestic'],values:[75.8,65.1]},frame:{x:72,y:200,width:673,height:370}}}];
console.log(JSON.stringify(compileDeck({id:'pie',slides},REGISTRY)));
""")
        self.assertIsNotNone(scene['slides'][0]['componentInstances'][0].get('nativeChart'))
        with tempfile.TemporaryDirectory() as tmp:
            file = Path(tmp) / 'pie.pptx'
            Emitter(scene).run(file)
            chart = next(s.chart for s in Presentation(file).slides[0].shapes if s.has_chart)
            xml = chart._chartSpace.xml
        # Percentages, formatted as the scene prints them, centred on the slice
        # - never the raw value outside the rim.
        self.assertIn('<c:showPercent val="1"/>', xml)
        self.assertNotIn('<c:showVal val="1"/>', xml)
        self.assertIn('formatCode="0%"', xml)
        self.assertNotIn('<c:dLblPos val="outEnd"/>', xml)


if __name__ == "__main__":
    unittest.main()
