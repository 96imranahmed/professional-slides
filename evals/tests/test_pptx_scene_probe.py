"""The PPTX scene probe (evals/scripts/pptx_scene_probe.py) reads a saved deck back as a scene.

It reads the semantic role and owner the emitter writes into each shape, and
both shape-name formats - the current `ps:<slide>-<owner>-<role>` and the older
`ps:<slide>-<owner>:<role>` - so decks saved before the change still probe.
"""
import importlib.util
import json
import tempfile
import unittest
from pathlib import Path

from node_probe import ROOT, requires_python_package

needs_pptx = requires_python_package('pptx')
try:
    from pptx import Presentation
    from pptx.oxml.ns import qn
    from pptx.util import Inches, Pt
    _spec = importlib.util.spec_from_file_location('scene_probe', ROOT / 'evals/scripts/pptx_scene_probe.py')
    probe = importlib.util.module_from_spec(_spec)
    _spec.loader.exec_module(probe)
except ImportError:  # pragma: no cover - exercised only without python-pptx
    Presentation = qn = Inches = Pt = probe = None


@needs_pptx
class SceneProbeTests(unittest.TestCase):
    def test_probe_reads_semantic_metadata_and_both_shape_name_formats(self):
        """PR #4 follow-up: the probe read only one shape-name format and ignored the semantic metadata."""
        with tempfile.TemporaryDirectory() as tmp:
            prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
            fixtures = [
                ('ps:s01-chrome-title', 'Growth funds expansion', 'action-title', None),
                ('ps:s01-body-text', 'Evidence supports the decision', 'paragraph', None),
                ('ps:s01-grid-cell-text-0-0', 'North', 'table-cell-text', None),
                ('ps:s01-bars-value-label-revenue-2026', '42', 'data-label', None),
                ('ps:s01-chrome:source', 'Source: audited results', 'source-text', None),
                ('ps:opaque-node', 'Footnote', 'footnote-text', {'role': 'footnote-text', 'owner': 's01:chrome'}),
                ('ps:opaque-value', '85', 'table-cell-text', {'role': 'table-cell-text', 'owner': 's01:table'}),
            ]
            for i, (name, text, role, semantic) in enumerate(fixtures):
                shape = slide.shapes.add_textbox(Inches(1), Inches(.3 + i * .7), Inches(6), Inches(.5))
                shape.name = name; shape.text = text; shape.text_frame.paragraphs[0].font.size = Pt(12)
                if semantic:
                    shape._element.find('.//' + qn('p:cNvPr')).set('descr', json.dumps(semantic))
            file = Path(tmp) / 'legacy.pptx'; prs.save(file)
            scene = probe.probe(file)['slides'][0]
            self.assertEqual([(n['text'], n['role']) for n in scene['nodes']], [(text, role) for _, text, role, _ in fixtures])
            self.assertEqual({c['component'] for c in scene['componentInstances']}, {'slide-chrome', 'paragraph', 'table', 'chart.column'})


if __name__ == "__main__":
    unittest.main()
