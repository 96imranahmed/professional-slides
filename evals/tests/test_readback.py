"""Emit the fixture deck and read it back with an independent reader.

The old acceptance manifest was generated from the scene it then validated. This
does the opposite: python-pptx opens the saved file and compares what PowerPoint
would see - paragraph counts, wrap mode, placeholders, native chart series -
against the scene the emitter was given.

The properties asserted here are the ones item 1 and item 2 of the revamp exist
to produce: no `wrap="none"`, a real title placeholder on every slide, and
native charts with embedded data.
"""

from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
EMIT = ROOT / "skills" / "professional-slides" / "runtime" / "emit"
GATES = ROOT / "skills" / "professional-slides" / "runtime" / "gates"

sys.path.insert(0, str(GATES))
import page_gates  # noqa: E402
from node_probe import example_scene  # noqa: E402

try:
    import pptx  # noqa: F401
except ImportError:  # pragma: no cover - environment without python-pptx
    pptx = None


def requires_pptx(test):
    return unittest.skipIf(pptx is None, "python-pptx is not installed")(test)


class ReadbackTests(unittest.TestCase):
    built = None

    @classmethod
    def setUpClass(cls):
        if pptx is None:
            return
        cls.tmp = tempfile.TemporaryDirectory()
        root = Path(cls.tmp.name)
        cls.scene_path = root / "scene.json"
        cls.scene = example_scene("nyc-or-sf")
        cls.scene_path.write_text(json.dumps(cls.scene))
        cls.pptx_path = root / "deck.pptx"
        emit = subprocess.run(
            [sys.executable, str(EMIT / "emit_pptx.py"), str(cls.scene_path), str(cls.pptx_path)],
            capture_output=True, text=True, cwd=str(ROOT))
        if emit.returncode:
            raise AssertionError(f"emit_pptx exited {emit.returncode}\n{emit.stderr}")
        cls.emit_stats = json.loads(emit.stdout)
        read = subprocess.run(
            [sys.executable, str(EMIT / "readback_pptx.py"), str(cls.scene_path), str(cls.pptx_path)],
            capture_output=True, text=True, cwd=str(ROOT))
        cls.readback_code = read.returncode
        cls.report = json.loads(read.stdout)

    @classmethod
    def tearDownClass(cls):
        if pptx is not None:
            cls.tmp.cleanup()

    @requires_pptx
    def test_the_saved_file_is_accepted_by_an_independent_reader(self):
        self.assertTrue(self.report["accepted"], json.dumps(self.report["findings"][:10], indent=1))
        self.assertEqual(self.readback_code, 0)
        self.assertEqual(self.report["findings"], [])

    @requires_pptx
    def test_no_text_shape_freezes_its_line_breaks(self):
        self.assertEqual(self.report["stats"]["wrap_none"], 0)
        self.assertGreater(self.report["stats"]["text_checked"], 200)

    @requires_pptx
    def test_every_slide_has_a_real_title_placeholder(self):
        self.assertEqual(self.report["stats"]["title_placeholders"], len(self.scene["slides"]))

    @requires_pptx
    def test_charts_are_native_objects_with_their_own_series(self):
        self.assertGreater(self.report["stats"]["charts_checked"], 0)
        self.assertEqual(self.emit_stats["native_charts"], self.report["stats"]["charts_checked"])
        self.assertGreater(self.emit_stats["grouped"], 0)

    @requires_pptx
    def test_powerpoints_paragraph_count_matches_the_scene(self):
        """The comparison the old validator could not make: file against plan."""
        from pptx import Presentation

        presentation = Presentation(str(self.pptx_path))
        by_name = {}
        for index, slide in enumerate(presentation.slides):
            def collect(shapes):
                for shape in shapes:
                    if shape.shape_type == 6:
                        collect(shape.shapes)
                    else:
                        by_name[(index, shape.name)] = shape
            collect(slide.shapes)

        compared = 0
        for index, slide in enumerate(self.scene["slides"]):
            for node in slide.get("nodes", []):
                if node.get("type") != "text":
                    continue
                source = ((node.get("data") or {}).get("textLayout") or {}).get("source")
                shape = by_name.get((index, f"ps:{node['id']}"))
                if source is None or shape is None or not shape.has_text_frame:
                    continue
                self.assertEqual(
                    len(shape.text_frame.paragraphs), len(str(source).split("\n")),
                    f"slide {index + 1} shape {shape.name}")
                compared += 1
        self.assertGreater(compared, 200, "the readback compared too few shapes")


if __name__ == "__main__":
    unittest.main()
