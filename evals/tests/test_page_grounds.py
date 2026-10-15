"""Pages set on a dark or a primary ground: `tone` on the deck or the page.

A calibration audit found analytical pages - charts, tables, KPI pages, text -
on a navy or black ground, or on a brand colour, in about one strong deck page
in ten, some decks dark throughout. The runtime drew a dark ground on its
structural pages and its panels only, no palette had a dark canvas, and a
palette written with one drew its chart navy on navy while every gate passed
it. `tone` now sets a deck, or one page, on the house's own dark or on its
primary, every token resolved for that ground from the deck's palette
(palettes.mjs groundTokens); the file keeps the page's colours (readback) and
the gates read type and marks on the ground they sit on.
"""
from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

from node_probe import RUNTIME, run_node

sys.path.insert(0, str(RUNTIME / "gates"))
sys.path.insert(0, str(RUNTIME / "emit"))
import ink  # noqa: E402
import page_gates  # noqa: E402
import scene_ink  # noqa: E402

try:
    from pptx import Presentation
    from emit_pptx import Emitter
    from readback_pptx import readback
except ImportError:  # pragma: no cover - environment without python-pptx
    Presentation = Emitter = readback = None


PRELUDE = """
import { composeAll } from './skills/professional-slides/runtime/compose-all.mjs';
import { resolvePalette, groundTokens, darkGroundOf, contrastRatio } from './skills/professional-slides/runtime/palettes.mjs';
import { relativeLuminance } from './skills/professional-slides/runtime/color.mjs';
import { applyDesign } from './skills/professional-slides/runtime/design-systems.mjs';
import { TOKENS, THEME_SLOT_TOKENS } from './skills/professional-slides/runtime/core.mjs';
import { deckKeyProblems } from './skills/professional-slides/runtime/deck-keys.mjs';
const BASE = './skills/professional-slides/examples';
const chart = { type: 'chart.column', heading: 'Passengers', unit: 'million', categories: ['2022', '2023', '2024', '2025'], series: [{ name: 'Passengers', values: [71, 79, 84, 85] }], highlights: [{ category: '2025' }] };
const table = { type: 'table', columns: ['Line', 'Journeys, million', 'Growth'], rows: [['Dales', '9.1', '19%'], ['Valley', '11.2', '17%'], ['Coast', '5.5', '6%'], { style: 'total', cells: ['Network', '25.8', '15%'] }] };
const page = (id, exhibit, more = {}) => ({ id, title: 'Passengers reached a record 85 million in 2025', layout: 'exhibit-full', exhibit, source: 'Source: illustrative data', ...more });
const deck = (slides, more = {}) => ({ schema: 'professional-slides.deck/v3', id: 't', design: 'consulting', ...more, slides });
const scene = (spec) => composeAll(structuredClone(spec), BASE).deck;
const refusal = (spec) => { try { composeAll(structuredClone(spec), BASE); return null; } catch (error) { return (error.pageErrors ?? [error.message])[0]; } };
const light = (design) => resolvePalette(applyDesign({ design, id: 'x' }).palette, TOKENS, THEME_SLOT_TOKENS).tokens;
const canvasOf = (slide) => slide.tokens['color.canvas'].value;
"""


class GroundTokenTests(unittest.TestCase):
    def test_every_colour_on_a_derived_ground_reads_on_it_in_every_design(self):
        result = run_node(PRELUDE + """
const out = {};
for (const design of ['consulting', 'editorial', 'journal', 'keynote']) for (const ground of ['dark', 'primary']) {
  const t = groundTokens(light(design), ground, THEME_SLOT_TOKENS), v = (k) => t[k].value, g = v('color.canvas');
  const on = (k) => contrastRatio(v(k), g);
  const marks = ['color.componentPrimary', 'color.chartComparator', ...[1, 2, 3, 4, 5, 6].map((i) => `color.chartSeries${i}`)];
  out[`${design}/${ground}`] = { ink: on('color.ink'), room: on('color.ink'),
    text: Math.min(on('color.textSecondary'), on('color.chartUnit'), on('color.accent'), ...['positive', 'caution', 'negative', 'info'].map((s) => on(`color.${s}`))),
    marks: Math.min(...marks.map(on)), onBand: contrastRatio(v('color.ink'), v('color.band')), band: contrastRatio(v('color.band'), g),
    lines: ['color.rule', 'color.chartGrid'].every((k) => relativeLuminance(v(k)) > relativeLuminance(g)),
    reversed: v('color.onPrimary') === g,
    unchanged: Object.entries(t).filter(([k, d]) => d.kind === 'color' && !k.startsWith('color.heat.') && d.value === light(design)[k].value && contrastRatio(d.value, g) < 3).map(([k]) => k) };
}
console.log(JSON.stringify(out));
""")
        for key, got in result.items():
            with self.subTest(ground=key):
                self.assertGreaterEqual(got["ink"], 4.5)
                self.assertGreaterEqual(got["text"], 4.5, "secondary type, units, the accent and the status colours set text")
                # 3:1 where the house leaves room above it, and never under the palette's own 2:1 floor.
                self.assertGreaterEqual(got["marks"], min(3, got["room"]) - 0.01)
                self.assertGreaterEqual(got["marks"], 2)
                self.assertGreaterEqual(got["onBand"], 4.5, "the type on a header band or a dark panel")
                self.assertGreaterEqual(got["band"], 1.4, "the band is a step off the ground")
                self.assertTrue(got["lines"], "rules and gridlines are light lines on a dark ground")
                self.assertTrue(got["reversed"], "type on a mark or the primary takes the ground's colour")
                self.assertEqual(got["unchanged"], [], "no colour left in its light-page value where it does not read")
        self.assertGreaterEqual(result["consulting/dark"]["ink"], 13)

    def test_the_dark_ground_is_the_house_own_dark(self):
        result = run_node(PRELUDE + """
const of = (palette) => darkGroundOf(resolvePalette(palette, TOKENS, THEME_SLOT_TOKENS).tokens);
console.log(JSON.stringify({ midnight: of('midnight'), crimson: of('crimson'), graphite: of('graphite'),
  navyHouse: of({ base: 'midnight', colors: { 'color.componentPrimary': '#00205B', 'color.ink': '#1A1A1A' } }),
  inkHouse: of({ base: 'evergreen', colors: { 'color.ink': '#0B2A4A' } }),
  editorial: darkGroundOf(light('editorial')), journal: darkGroundOf(light('journal')), keynote: darkGroundOf(light('keynote')) }));
""")
        # A navy house is set on its navy, not on a fixed black.
        self.assertEqual(result["midnight"], "#051C2C")
        self.assertEqual(result["navyHouse"], "#00205B")
        # A brand colour that is not a dark of its own (a red, a green) gives way to the house's ink.
        self.assertEqual(result["crimson"], "#252525")
        self.assertEqual(result["inkHouse"], "#0B2A4A")
        self.assertEqual(result["graphite"], "#000000")
        self.assertEqual(result["editorial"], "#221E1A")
        self.assertNotEqual(result["keynote"], "#000000")

    def test_what_cannot_be_drawn_on_a_ground_is_refused(self):
        result = run_node(PRELUDE + """
let unnamed = null;
try { groundTokens({ ...light('consulting'), 'color.extra': { kind: 'color', cssVar: '--x', value: '#123456', themeSlot: null } }, 'dark'); } catch (error) { unnamed = error.message; }
console.log(JSON.stringify({
  paleBrand: refusal(deck([page('p1', chart, { tone: 'primary' })], { palette: { base: 'midnight', colors: { 'color.componentPrimary': '#7FB2E5' } } })),
  word: refusal(deck([page('p1', chart, { tone: 'navy' })])),
  divider: refusal(deck([{ id: 's1', kind: 'section', title: 'The demand', tone: 'dark' }, page('p1', chart)])),
  takeaways: refusal(deck([{ id: 'k1', kind: 'takeaways', items: ['One finding', 'Another finding'], tone: 'primary' }])),
  deckPrimary: refusal(deck([page('p1', chart)], { tone: 'primary' })),
  deckKey: deckKeyProblems({ tone: 'primary' }), unnamed }));
""")
        self.assertIn("carries pale type at 2.2:1", result["paleBrand"])
        self.assertIn('`tone: "dark"`', result["paleBrand"])
        self.assertIn('`tone` on a content page is "light", "dark", "primary"', result["word"])
        self.assertIn("a section page takes no `tone`", result["divider"])
        self.assertIn('`tone` on a takeaways page is "dark", "light"', result["takeaways"])
        self.assertIn("Unknown deck tone", result["deckPrimary"])
        self.assertIn('a single page takes "primary"', result["deckKey"][0])
        self.assertIn("No ground role for color.extra", result["unnamed"])


class GroundPageTests(unittest.TestCase):
    def test_a_dark_page_in_a_light_deck_leaves_the_rest_untouched(self):
        result = run_node(PRELUDE + """
const plain = scene(deck([page('p1', chart), page('p2', chart), page('p3', table)]));
const toned = scene(deck([page('p1', chart), page('p2', chart, { tone: 'dark' }), page('p3', table)]));
const strip = (slide) => JSON.stringify(slide.nodes);
const dark = toned.slides[1], node = (role) => dark.nodes.find((n) => n.role === role);
const marks = dark.nodes.filter((n) => n.role === 'chart-mark');
console.log(JSON.stringify({ same: [0, 2].map((i) => strip(plain.slides[i]) === strip(toned.slides[i])),
  plainGround: plain.slides.map((s) => s.ground ?? null), ground: dark.ground, canvas: canvasOf(dark), lightCanvas: canvasOf(toned.slides[0]),
  title: node('action-title').style.color.value, source: contrastRatio(node('source-text').style.color.value, canvasOf(dark)),
  marks: Math.min(...marks.map((n) => contrastRatio(n.style.fill.value, canvasOf(dark)))),
  deckTokens: toned.tokens['color.canvas'].value }));
""")
        self.assertEqual(result["same"], [True, True], "the pages around it compile exactly as before")
        self.assertEqual(result["plainGround"], [None, None, None])
        self.assertEqual(result["ground"], "dark")
        self.assertEqual(result["canvas"], "#051C2C")
        self.assertEqual(result["lightCanvas"], "#FFFFFF")
        self.assertEqual(result["deckTokens"], "#FFFFFF")
        self.assertEqual(result["title"], "#FFFFFF")
        self.assertGreaterEqual(result["source"], 4.5)
        self.assertGreaterEqual(result["marks"], 3)

    def test_a_dark_deck_keeps_the_pages_that_paint_their_own_ground(self):
        result = run_node(PRELUDE + """
const slides = [{ id: 's1', kind: 'section', title: 'The demand' }, page('p1', chart), { id: 'm1', kind: 'statement', text: 'The peak is full.' },
  { id: 'm2', kind: 'statement', text: 'The off-peak is not.', tone: 'dark' }, page('p2', table, { tone: 'light' }), { id: 'k1', kind: 'takeaways', items: ['One finding', 'Another finding'] }];
const grounds = (design) => Object.fromEntries(scene(deck(slides, { design, tone: 'dark', cover: { title: 'Northvale' } })).slides.map((s) => [s.id, canvasOf(s)]));
console.log(JSON.stringify({ consulting: grounds('consulting'), editorial: grounds('editorial') }));
""")
        c = result["consulting"]
        # The navy cover, divider panel and dark statement and takeaways draw their own navy on the light page's tokens.
        self.assertEqual([c["cover"], c["s1"], c["m2"], c["k1"]], ["#FFFFFF"] * 4)
        self.assertEqual([c["p1"], c["m1"]], ["#051C2C"] * 2)
        self.assertEqual(c["p2"], "#FFFFFF", "a page's own tone wins over the deck's")
        e = result["editorial"]
        # The editorial cover and divider are set on the page, so on a dark deck they are dark.
        self.assertEqual([e["cover"], e["s1"], e["p1"]], ["#221E1A"] * 3)

    def test_a_surface_type_is_reversed_out_of_is_the_band_on_a_dark_page(self):
        result = run_node(PRELUDE + """
const out = scene(deck([page('p1', table), page('p2', table, { tone: 'dark' })]));
const header = (slide) => slide.nodes.find((n) => n.role === 'table-header-cell').style.fill;
const total = (slide) => slide.nodes.find((n) => n.role === 'table-row-band').style.fill;
const text = (slide) => slide.nodes.find((n) => n.role === 'table-header-text').style.color.value;
const [lit, dark] = out.slides;
console.log(JSON.stringify({ lit: [header(lit).tokenId, total(lit).tokenId, text(lit)], dark: [header(dark).tokenId, total(dark).tokenId, text(dark)],
  band: dark.tokens['color.band'].value, step: contrastRatio(header(dark).value, canvasOf(dark)), onBand: contrastRatio(text(dark), header(dark).value) }));
""")
        self.assertEqual(result["lit"], ["color.ink", "color.componentPrimary", "#FFFFFF"], "a light page draws its band in the ink, as before")
        self.assertEqual(result["dark"][:2], ["color.band", "color.band"])
        self.assertEqual(result["dark"][2], "#FFFFFF")
        self.assertGreaterEqual(result["step"], 1.4)
        self.assertGreaterEqual(result["onBand"], 4.5)


def ground_scene(spec_js: str) -> dict:
    return run_node(PRELUDE + f"console.log(JSON.stringify(scene({spec_js})));")


class GroundGateTests(unittest.TestCase):
    def test_type_and_marks_lost_on_their_ground_are_refused_and_a_toned_page_is_not(self):
        # A palette that sets a dark canvas under a light ink and primary passes the palette's own check, and still draws its
        # series and surfaces for a light page: the audit's navy-on-navy chart.
        hand = ground_scene("deck([page('p1', chart), page('p2', table)], { palette: { base: 'midnight', colors: { 'color.canvas': '#0B2341', 'color.ink': '#FFFFFF', 'color.componentPrimary': '#9EB7D8' } } })")
        toned = ground_scene("deck([page('p1', chart), page('p2', table)], { tone: 'dark' })")
        primary = ground_scene("deck([page('p1', chart, { tone: 'primary' }), page('p2', table, { tone: 'primary' })], { design: 'journal' })")
        lost = page_gates.run_gates(hand, gates={"GROUND_CONTRAST"})
        self.assertEqual({f["slide"] for f in lost["findings"]}, {1, 2})
        self.assertTrue(all(f["severity"] == "blocker" for f in lost["findings"]))
        self.assertIn("`tone`", lost["findings"][0]["repair"])
        for clean in (toned, primary):
            self.assertEqual(page_gates.run_gates(clean, gates={"GROUND_CONTRAST"})["findings"], [])

    def test_a_chart_whose_every_mark_is_lost_is_refused(self):
        slide = {"tokens": {"color.canvas": {"value": "#051C2C"}}, "nodes": [
            {"id": "m1", "type": "rect", "role": "chart-mark", "frame": {"x": 100, "y": 100, "width": 40, "height": 200}, "style": {"fill": {"value": "#0B2341"}}, "data": {"componentInstance": "c"}},
            {"id": "m2", "type": "rect", "role": "chart-mark", "frame": {"x": 160, "y": 150, "width": 40, "height": 150}, "style": {"fill": {"value": "#06233B"}}, "data": {"componentInstance": "c"}}]}
        lost = page_gates.ground_contrast(slide)
        self.assertEqual([item["kind"] for item in lost], ["marks"])
        # One mark that reads is enough: on a light page a comparator's grey sits at 1.4:1 beside the mark that carries the point.
        slide["nodes"][1]["style"]["fill"]["value"] = "#FFFFFF"
        self.assertEqual(page_gates.ground_contrast(slide), [])

    def test_the_render_reads_a_dark_page_as_marks_lighter_than_it(self):
        import numpy as np
        grey = np.full((ink.CANVAS_H, ink.CANVAS_W), 30, dtype=np.uint8)
        grey[300:320, 100:600] = 250  # a line of pale type on navy
        self.assertEqual(int(ink.ink_matrix(grey).sum()), 0, "read as a light page, a dark page has no ink at all")
        self.assertEqual(int(ink.ink_matrix(grey, dark=True).sum()), 20 * 500)
        rows = ink.ink_rows(grey, ink.SURFACE_LUMINANCE, dark=True)
        self.assertEqual(sum(1 for r in rows if r), 20)
        dark = {"tokens": {"color.canvas": {"value": "#051C2C"}}, "nodes": []}
        navy_panel = {"tokens": {"color.canvas": {"value": "#FFFFFF"}}, "nodes": [
            {"type": "rect", "frame": {"x": 0, "y": 0, "width": 1280, "height": 720}, "style": {"fill": {"value": "#051C2C"}}}]}
        self.assertTrue(scene_ink.dark_ground(dark))
        # A light page that paints a navy panel over itself (a divider) is read as it always was.
        self.assertFalse(scene_ink.dark_ground(navy_panel))

    def test_type_the_render_shows_lost_on_its_ground_is_refused(self):
        import numpy as np
        slide = {"tokens": {"color.canvas": {"value": "#051C2C"}}, "nodes": [
            {"id": "t1", "type": "text", "role": "paragraph", "text": "Passengers rose to 85 million", "frame": {"x": 100, "y": 200, "width": 400, "height": 20},
             "style": {"color": {"value": "#FFFFFF"}, "fontSize": {"value": 12}}}]}
        navy = np.full((ink.CANVAS_H, ink.CANVAS_W), 26, dtype=np.uint8)
        lost = page_gates.render_contrast(slide, navy)
        self.assertEqual([item["node"] for item in lost], ["t1"])
        navy[205:215, 110:300] = 240
        self.assertEqual(page_gates.render_contrast(slide, navy), [])
        findings = []
        page_gates.gate_render_contrast(1, slide, np.full((ink.CANVAS_H, ink.CANVAS_W), 26, dtype=np.uint8), findings)
        self.assertEqual(findings[0]["code"], "RENDER_CONTRAST")


@unittest.skipIf(Emitter is None, "python-pptx is not installed")
class GroundFileTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.scene = ground_scene("deck([page('p1', chart), page('p2', chart, { tone: 'dark', highlight: 'record' }), page('p3', table, { tone: 'dark' })])")
        cls.tmp = tempfile.TemporaryDirectory()
        cls.scene_path = Path(cls.tmp.name) / "scene.json"
        cls.scene_path.write_text(json.dumps(cls.scene))

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def emit(self, name, emitter=Emitter):
        out = Path(self.tmp.name) / name
        emitter(json.loads(self.scene_path.read_text())).run(out)
        return out

    def test_the_file_keeps_the_dark_page_colours(self):
        path = self.emit("deck.pptx")
        report = readback(self.scene, path)
        self.assertTrue(report["accepted"], report["findings"])
        prs = Presentation(str(path))
        dark = prs.slides[1]
        self.assertEqual(str(dark.background.fill.fore_color.rgb), "051C2C")
        chart = next(shape for shape in dark.shapes if getattr(shape, "has_chart", False)).chart
        self.assertEqual(str(chart.font.color.rgb), "FFFFFF")
        light_chart = next(shape for shape in prs.slides[0].shapes if getattr(shape, "has_chart", False)).chart
        self.assertIsNone(light_chart.font.color.type, "a light page's chart keeps the theme's text, as before")

    def test_a_file_drawn_in_the_deck_light_colours_is_refused(self):
        class DeckColours(Emitter):
            def use_colors(self, colors):
                super().use_colors(getattr(self, "deck_colors", colors))
        report = readback(self.scene, self.emit("deck-colours.pptx", DeckColours))
        codes = {finding["code"] for finding in report["findings"]}
        self.assertIn("CHART_TEXT_COLOR", codes)
        tampered = json.loads(json.dumps(self.scene))
        title = next(n for n in tampered["slides"][1]["nodes"] if n["role"] == "action-title")
        title["style"]["color"]["value"] = "#051C2C"
        rect = next(n for n in tampered["slides"][2]["nodes"] if n["role"] == "table-header-cell")
        rect["style"]["fill"]["value"] = "#123456"
        codes = {finding["code"] for finding in readback(tampered, self.emit("again.pptx"))["findings"]}
        self.assertTrue({"TEXT_COLOR", "SHAPE_COLOR"} <= codes)


if __name__ == "__main__":
    unittest.main()
