"""PR #4 follow-up regressions at the package, CLI and saved-PPTX boundaries."""
import importlib.util
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from node_probe import NODE, ROOT, RUNTIME, run_node
from test_plugin_distribution import packager
from test_pr_review_export import Emitter


def load_module(name, file):
    spec = importlib.util.spec_from_file_location(name, file)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


probe = load_module('followup_probe', ROOT / 'evals/scripts/pptx_scene_probe.py')
importer = load_module('followup_importer', RUNTIME / 'import-template.py')


class FollowupReviewTests(unittest.TestCase):
    def cli(self, *args):
        return subprocess.run([NODE, *map(str, args)], capture_output=True, text=True,
                              env={**os.environ, 'RUNTIME_PYTHON': sys.executable}, cwd=ROOT, timeout=120)

    def test_package_compiles_examples_and_rejects_plugin_local_outputs(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp).resolve()
            package = root / 'plugin'
            packager.package(ROOT, package)
            skill = package / 'skills/professional-slides'
            image = skill / 'examples/assets/hills.jpg'
            self.assertEqual(image.read_bytes(), (ROOT / 'skills/professional-slides/examples/assets/hills.jpg').read_bytes())
            manifest = json.loads((package / 'package-manifest.json').read_text())
            self.assertIn('skills/professional-slides/examples/assets/hills.jpg', manifest['files'])
            for example in ['slideworks', 'nyc-or-sf']:
                result = self.cli(package / 'evals/scripts/compile_scene.mjs',
                                  skill / f'examples/{example}.deck.json', root / f'{example}.scene.json')
                self.assertEqual(result.returncode, 0, result.stderr)
                self.assertGreater(len(json.loads((root / f'{example}.scene.json').read_text())['slides']), 1)
            spec = root / 'spec.json'
            spec.write_text(json.dumps({'schema': 'professional-slides.deck/v3', 'id': 'safe', 'cover': {'title': 'Decision'}, 'slides': []}))
            alias = root / 'alias'
            alias.symlink_to(package, target_is_directory=True)
            dangling = root / 'dangling'
            dangling.symlink_to(package / 'new-output', target_is_directory=True)
            for output in [package, package / 'new-output', alias / 'nested/output', dangling / 'nested']:
                for command in ['build-deck.mjs', 'deliver-deck.mjs']:
                    result = self.cli(skill / f'runtime/{command}', spec, output, '--preflight')
                    self.assertEqual(result.returncode, 1, result.stderr)
                    self.assertIn('Plugin files are read-only', result.stderr)
            self.assertFalse((package / 'new-output').exists())
            self.assertFalse((package / 'nested').exists())
            self.assertFalse((package / 'scene.json').exists())
            # A writable installation must retain existing reports and deliveries.
            (package / 'scene.json').write_text('keep scene')
            (package / 'safe-DELIVERED.pptx').write_text('keep delivery')
            for command in ['build-deck.mjs', 'deliver-deck.mjs']:
                self.cli(skill / f'runtime/{command}', spec, alias)
            self.assertEqual((package / 'scene.json').read_text(), 'keep scene')
            self.assertEqual((package / 'safe-DELIVERED.pptx').read_text(), 'keep delivery')
            result = self.cli(skill / 'runtime/build-deck.mjs', spec, root / 'external', '--preflight')
            self.assertEqual(result.returncode, 0, result.stderr)

    def test_developer_output_exception_and_symlink_escape(self):
        run_node('''
import assert from 'node:assert/strict';
import fs from 'node:fs/promises';
import os from 'node:os';
import path from 'node:path';
import {assertOutputDirectory} from './skills/professional-slides/runtime/output-path.mjs';
const root=await fs.mkdtemp(path.join(os.tmpdir(),'output-guard-'));
try {
 await fs.mkdir(path.join(root,'.git'));
 await fs.mkdir(path.join(root,'output'));
 assert.equal(await assertOutputDirectory(path.join(root,'output/task'),root),path.join(await fs.realpath(root),'output/task'));
 await fs.symlink(root,path.join(root,'output/escape'));
 await assert.rejects(assertOutputDirectory(path.join(root,'output/escape/skills/new'),root),/read-only/);
 await assert.rejects(assertOutputDirectory(path.join(root,'skills/new'),root),/read-only/);
} finally {await fs.rm(root,{recursive:true,force:true});}
console.log(JSON.stringify({ok:true}));
''')

    def test_unrendered_build_cannot_be_delivered_even_with_old_built_status(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); out = root / 'out'; spec = root / 'spec.json'
            spec.write_text(json.dumps({'schema': 'professional-slides.deck/v3', 'id': 'safe', 'cover': {'title': 'Decision'}, 'slides': []}))
            result = self.cli(RUNTIME / 'build-deck.mjs', spec, out, '--no-render')
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertEqual(json.loads(result.stdout)['status'], 'built-unrendered')
            review = root / 'accepted.json'
            review.write_text(json.dumps({'accepted': True, 'summary': 'Accepted to verify that rendered gates cannot be bypassed.', 'findings': []}))
            for status in ['built-unrendered', 'built']:
                build = json.loads((out / 'build-result.json').read_text()); build['status'] = status
                (out / 'build-result.json').write_text(json.dumps(build))
                delivered = out / 'safe-DELIVERED.pptx'; delivered.write_text('stale')
                result = self.cli(RUNTIME / 'deliver-deck.mjs', spec, out, '--skip-build', '--review', review)
                self.assertEqual(result.returncode, 2, result.stderr)
                self.assertIn('MISSING_RENDERED_GATES', [f['code'] for f in json.loads(result.stdout)['blockers']])
                self.assertFalse(delivered.exists())

    def test_removed_dependency_route_has_explicit_usage_error(self):
        result = self.cli(ROOT / 'evals/scripts/run_tests.mjs', '--dependencies')
        self.assertNotEqual(result.returncode, 0)
        self.assertIn('Usage: run_tests.mjs [--release]', result.stderr)
        self.assertNotIn('MODULE_NOT_FOUND', result.stderr)

    def test_empty_footer_overrides_house_and_raw_planner_compilation_still_works(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            house = root / 'house.json'
            house.write_text(json.dumps({'schema': 'professional-slides.house/v1', 'footer': 'House name'}))
            run_node(f'''
import assert from 'node:assert/strict';
import {{applyTemplate,toDeckPlan}} from './skills/professional-slides/runtime/compose.mjs';
const spec={{schema:'professional-slides.deck/v3',template:'house.json',slides:[{{title:'Growth funds expansion',points:['Evidence supports the decision']}}]}};
const base={json.dumps(str(root))};
assert.equal(applyTemplate(spec,base).footer,'House name');
assert.equal(applyTemplate({{...spec,footer:''}},base).footer,'');
assert.equal(applyTemplate({{...spec,footer:'Custom'}},base).footer,'Custom');
console.log(JSON.stringify({{ok:true}}));
''')
            plan = {'id': 'raw', 'slides': [{'id': 'one', 'title': 'Growth funds expansion', 'items': [{'id': 'body', 'component': 'paragraph', 'props': {'text': 'Evidence supports expansion.'}}]}]}
            for payload in [plan, {'deckPlan': plan}]:
                file = root / 'spec.json'; file.write_text(json.dumps(payload))
                result = self.cli(ROOT / 'evals/scripts/compile_scene.mjs', file, root / 'scene.json')
                self.assertEqual(result.returncode, 0, result.stderr)

    def test_template_coordinates_match_across_aspect_ratios(self):
        with tempfile.TemporaryDirectory() as tmp:
            profiles = []
            for width, height in [(10, 7.5), (16, 9), (9, 16)]:
                prs = Presentation()
                slide = prs.slides.add_slide(prs.slide_layouts[1])
                slide.shapes.title.text = 'Growth funds expansion'
                slide.placeholders[1].text = 'Evidence supports expansion across three regions.'
                sx, sy = Inches(width) / prs.slide_width, Inches(height) / prs.slide_height
                surfaces = [*prs.slide_masters, *prs.slide_layouts, *prs.slides]
                # Snapshot inherited placeholder geometry before changing its
                # master; otherwise later placeholders are scaled twice.
                frames = [(shape, shape.left, shape.top, shape.width, shape.height)
                          for surface in surfaces for shape in surface.shapes]
                for shape, left, top, shape_width, shape_height in frames:
                    shape.left = round(left * sx); shape.width = round(shape_width * sx)
                    shape.top = round(top * sy); shape.height = round(shape_height * sy)
                prs.slide_width, prs.slide_height = Inches(width), Inches(height)
                file = Path(tmp) / f'{width}-{height}.pptx'; prs.save(file)
                profiles.append(importer.analyse(file, 'mckinsey'))
            for profile in profiles[1:]:
                self.assertEqual(profile['chrome'], profiles[0]['chrome'])
                self.assertEqual(profile['stats']['medianBodyCoverage'], profiles[0]['stats']['medianBodyCoverage'])
            self.assertLess(profiles[0]['chrome']['footerTop'], 690)

    def test_probe_reads_semantic_metadata_and_both_shape_name_formats(self):
        with tempfile.TemporaryDirectory() as tmp:
            prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
            fixtures = [
                ('ps:s01-chrome-title', 'Growth funds expansion', 'action-title', None),
                ('ps:s01-body-text', 'Evidence supports the decision', 'paragraph', None),
                ('ps:s01-grid-cell-text-0-0', 'North', 'table-cell-text', None),
                ('ps:s01-bars-value-label-revenue-2026', '42', 'data-label', None),
                ('ps:s01-chrome:source', 'Source: audited results', 'source-text', None),
                ('ps:opaque-node', 'Footnote', 'footnote-text', {'role': 'footnote-text', 'owner': 's01:chrome'}),
                ('ps:opaque-value', '85', 'table-cell-text', {'role': 'table-cell-text', 'owner': 's01:table'}),
            ]
            for i, (name, text, role, semantic) in enumerate(fixtures):
                shape = slide.shapes.add_textbox(Inches(1), Inches(.3 + i * .7), Inches(6), Inches(.5))
                shape.name = name; shape.text = text; shape.text_frame.paragraphs[0].font.size = Pt(12)
                if semantic:
                    shape._element.find('.//' + qn('p:cNvPr')).set('descr', json.dumps(semantic))
            file = Path(tmp) / 'legacy.pptx'; prs.save(file)
            scene = probe.probe(file)['slides'][0]
            self.assertEqual([(n['text'], n['role']) for n in scene['nodes']], [(text, role) for _, text, role, _ in fixtures])
            self.assertEqual({c['component'] for c in scene['componentInstances']}, {'slide-chrome', 'paragraph', 'table', 'chart.column'})

    def test_native_chart_axis_and_point_colors_survive_saved_pptx(self):
        with tempfile.TemporaryDirectory() as tmp:
            scene = run_node('''
import {compileDeck} from './skills/professional-slides/runtime/core.mjs';
import {REGISTRY} from './skills/professional-slides/runtime/registry.mjs';
const cases=[];
for (const type of ['column','bar']) for (const labels of [true,false]) for (const axis of [true,false])
 for (const extra of [{},{colorIndices:[2]},{highlights:[{category:'B',style:'bar'}]},{forecastFrom:'B'}])
  cases.push({component:'chart.'+type,props:{categories:['A','B'],series:[{name:'Revenue',values:[20,40]}],dataLabels:labels,showValueAxis:axis,...extra}});
const slides=cases.map((c,i)=>({id:'s'+i,composition:{nodeType:'component',id:'plot',...c,frame:{x:60,y:160,width:1000,height:460}}}));
console.log(JSON.stringify(compileDeck({id:'charts',slides},REGISTRY)));
''')
            file = Path(tmp) / 'charts.pptx'; Emitter(scene).run(file); prs = Presentation(file)
            for saved, source in zip(prs.slides, scene['slides']):
                chart = next(shape.chart for shape in saved.shapes if shape.has_chart)
                native = source['componentInstances'][0]['nativeChart']
                self.assertEqual(chart.value_axis.visible, native['showValueAxis'])
                self.assertEqual(chart.value_axis.visible, any(n['role'] == 'axis-label' for n in source['nodes']))
                marks = [n for n in source['nodes'] if n['role'] == 'chart-mark']
                series = chart.series[0]
                for point, mark in zip(series.points, marks):
                    color = point.format.fill.fore_color.rgb if point.format.fill.type else series.format.fill.fore_color.rgb
                    self.assertEqual(str(color), mark['style']['fill']['value'].lstrip('#').upper())


if __name__ == '__main__':
    unittest.main()
