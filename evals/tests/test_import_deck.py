"""An existing deck becomes an inventory and a starter pages file
(runtime/import-deck.py): every slide by stable id with its copy, tables,
charts, pictures and notes, and a starter the author revises page by page
that is never overwritten without --force."""
import hashlib
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from node_probe import HAS_PILLOW, HAS_PPTX, ROOT, RUNTIME, RUNTIME_PYTHON

PYTHON = RUNTIME_PYTHON or sys.executable
SCRIPT = RUNTIME / "import-deck.py"

# Five slides: a title slide, bullets at two levels with notes, a table, a
# clustered column chart, and a hidden slide whose picture and caption sit in
# a group drawn at twice its child frame.
BUILD = r'''
import sys
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn
from pptx.util import Emu

out = Path(sys.argv[1])
px = lambda n: Emu(n * 9525)
swatch = out.parent / "swatch.png"
Image.new("RGB", (40, 20), (14, 122, 94)).save(swatch)
prs = Presentation()
prs.slide_width, prs.slide_height = px(1280), px(720)
layouts = {layout.name: layout for layout in prs.slide_layouts}

cover = prs.slides.add_slide(layouts["Title Slide"])
cover.shapes.title.text = "Northvale growth plan"
cover.placeholders[1].text = "Board meeting\vMarch 2026"

bullets = prs.slides.add_slide(layouts["Title and Content"])
bullets.shapes.title.text = "Revenue grew nine percent while costs held flat"
body = bullets.placeholders[1].text_frame
body.text = "Growth came from the core"
for text, level in (("North grew 12%", 1), ("Costs held flat", 0)):
    paragraph = body.add_paragraph()
    paragraph.text, paragraph.level = text, level
bullets.notes_slide.notes_text_frame.text = "Say the margin number out loud"

table_slide = prs.slides.add_slide(layouts["Title Only"])
table_slide.shapes.title.text = "Three regions carry the growth"
table = table_slide.shapes.add_table(3, 2, px(60), px(200), px(600), px(200)).table
for r, row in enumerate([["Region", "Growth"], ["North", "12%"], ["South", "9%"]]):
    for c, text in enumerate(row):
        table.cell(r, c).text = text

chart_slide = prs.slides.add_slide(layouts["Title Only"])
chart_slide.shapes.title.text = "Revenue outgrew cost in each of three years"
data = CategoryChartData()
data.categories = ["2023", "2024", "2025"]
data.add_series("Revenue", (40, 46, 52))
data.add_series("Cost", (30, 31, 33))
chart = chart_slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, px(60), px(200), px(800), px(400), data).chart
chart.has_title = True
chart.chart_title.text_frame.text = "Revenue and cost"

photo = prs.slides.add_slide(layouts["Title Only"])
photo.shapes.title.text = "The depot today"
group = photo.shapes.add_group_shape()
picture = group.shapes.add_picture(str(swatch), px(100), px(300), px(160), px(80))
picture._element.nvPicPr.cNvPr.set("descr", "Green swatch")
group.shapes.add_textbox(px(100), px(390), px(160), px(30)).text_frame.text = "Caption in a group"
ext = group._element.find(qn("p:grpSpPr")).find(qn("a:xfrm")).find(qn("a:ext"))
ext.set("cx", str(int(ext.get("cx")) * 2))
ext.set("cy", str(int(ext.get("cy")) * 2))
photo._element.set("show", "0")

prs.save(str(out))
'''


def run(*args):
    return subprocess.run([PYTHON, str(SCRIPT), *map(str, args)], capture_output=True, text=True, cwd=ROOT, timeout=120)


@unittest.skipUnless(HAS_PPTX and HAS_PILLOW, "needs python-pptx and Pillow (python3 -m pip install -r requirements.txt)")
class ImportDeckTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp_dir = tempfile.TemporaryDirectory()
        cls.tmp = Path(cls.tmp_dir.name)
        cls.pptx = cls.tmp / "Board Pack (Q3) v2.pptx"
        subprocess.run([PYTHON, "-c", BUILD, str(cls.pptx)], check=True, capture_output=True, text=True, timeout=120)
        cls.out = cls.tmp / "out"
        cls.result = run(cls.pptx, cls.out)
        if cls.result.returncode != 0:
            raise AssertionError(cls.result.stderr)
        cls.summary = json.loads(cls.result.stdout)
        cls.inventory = json.loads((cls.out / "board-pack-q3-v2.inventory.json").read_text())
        cls.starter_text = (cls.out / "board-pack-q3-v2.pages.json").read_text()
        cls.starter = json.loads(cls.starter_text)
        cls.slides = {s["id"]: s for s in cls.inventory["slides"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp_dir.cleanup()

    def test_summary_counts_what_was_read(self):
        self.assertEqual(self.summary["inventory"], str(self.out / "board-pack-q3-v2.inventory.json"))
        self.assertEqual(self.summary["pages"], str(self.out / "board-pack-q3-v2.pages.json"))
        self.assertEqual({k: self.summary[k] for k in ("slides", "pictures", "charts", "tables")}, {"slides": 5, "pictures": 1, "charts": 1, "tables": 1})

    def test_inventory_names_the_source_by_basename_and_hash(self):
        self.assertEqual(self.inventory["schema"], "professional-slides.inventory/v1")
        self.assertEqual(self.inventory["id"], "board-pack-q3-v2")
        self.assertEqual(self.inventory["source"], {"file": "Board Pack (Q3) v2.pptx", "sha256": hashlib.sha256(self.pptx.read_bytes()).hexdigest()})
        self.assertEqual(self.inventory["slideSize"], {"width": 1280, "height": 720})
        self.assertEqual([s["id"] for s in self.inventory["slides"]], ["s01", "s02", "s03", "s04", "s05"])
        self.assertEqual([s["index"] for s in self.inventory["slides"]], [1, 2, 3, 4, 5])
        self.assertEqual([s["layout"] for s in self.inventory["slides"]], ["Title Slide", "Title and Content", "Title Only", "Title Only", "Title Only"])
        self.assertEqual([s["hidden"] for s in self.inventory["slides"]], [False, False, False, False, True])

    def test_titles_subtitles_paragraphs_and_notes(self):
        s01, s02 = self.slides["s01"], self.slides["s02"]
        self.assertEqual((s01["title"], s01["subtitle"], s01["paragraphs"]), ("Northvale growth plan", "Board meeting March 2026", []))
        self.assertEqual(s02["title"], "Revenue grew nine percent while costs held flat")
        self.assertIsNone(s02["subtitle"])
        self.assertEqual(s02["paragraphs"], [{"text": "Growth came from the core", "level": 0}, {"text": "North grew 12%", "level": 1}, {"text": "Costs held flat", "level": 0}])
        self.assertEqual(s02["notes"], "Say the margin number out loud")
        self.assertIsNone(s01["notes"])

    def test_table_cells_stay_out_of_the_paragraphs(self):
        s03 = self.slides["s03"]
        self.assertEqual(s03["tables"], [{"rows": 3, "cols": 2, "cells": [["Region", "Growth"], ["North", "12%"], ["South", "9%"]]}])
        self.assertEqual(s03["paragraphs"], [])

    def test_chart_categories_and_series_values(self):
        (chart,) = self.slides["s04"]["charts"]
        self.assertEqual(chart["type"], "COLUMN_CLUSTERED")
        self.assertEqual(chart["title"], "Revenue and cost")
        self.assertEqual(chart["categories"], ["2023", "2024", "2025"])
        self.assertEqual(chart["series"], [{"name": "Revenue", "values": [40, 46, 52]}, {"name": "Cost", "values": [30, 31, 33]}])
        self.assertNotIn("unreadable", chart)
        self.assertEqual(self.slides["s04"]["paragraphs"], [])

    def test_grouped_picture_is_extracted_at_its_displayed_size(self):
        s05 = self.slides["s05"]
        self.assertEqual(s05["pictures"], [{"file": "assets/board-pack-q3-v2/s05-1.png", "width": 320, "height": 160, "alt": "Green swatch"}])
        self.assertTrue((self.out / "assets" / "board-pack-q3-v2" / "s05-1.png").read_bytes().startswith(b"\x89PNG"))
        self.assertEqual(s05["paragraphs"], [{"text": "Caption in a group", "level": 0}])

    def test_starter_takes_the_cover_and_drafts_the_rest(self):
        deck = self.starter["deck"]
        self.assertEqual(deck, {"schema": "professional-slides.deck/v3", "id": "board-pack-q3-v2", "workflow": "existing_deck_revision",
                                "inventory": "board-pack-q3-v2.inventory.json", "request": "",
                                "cover": {"title": "Northvale growth plan", "subtitle": "Board meeting March 2026"}})
        pages = self.starter["pages"]
        self.assertEqual([(p["id"], p["sourceSlide"]) for p in pages], [("s02", 2), ("s03", 3), ("s04", 4), ("s05", 5)])
        self.assertEqual(pages[0]["title"], "Revenue grew nine percent while costs held flat")
        self.assertEqual(pages[0]["draft"], {"text": ["Growth came from the core", "  North grew 12%", "Costs held flat"], "notes": "Say the margin number out loud"})
        self.assertEqual(pages[1]["draft"], {"tables": [{"rows": 3, "cols": 2}]})
        self.assertEqual(pages[2]["draft"], {"charts": [{"type": "COLUMN_CLUSTERED", "categories": 3, "series": ["Revenue", "Cost"]}]})
        self.assertEqual(pages[3]["draft"], {"text": ["Caption in a group"], "pictures": ["assets/board-pack-q3-v2/s05-1.png"], "hidden": True})

    def test_a_starter_is_never_overwritten_without_force(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "out"
            self.assertEqual(run(self.pptx, out).returncode, 0)
            starter = out / "board-pack-q3-v2.pages.json"
            revised = json.loads(starter.read_text())
            revised["deck"]["request"] = "Tighten the board pack"
            starter.write_text(json.dumps(revised))
            again = run(self.pptx, out)
            self.assertEqual(again.returncode, 2)
            self.assertIn("--force", again.stderr)
            self.assertEqual(again.stdout, "")
            self.assertEqual(json.loads(starter.read_text())["deck"]["request"], "Tighten the board pack")
            forced = run(self.pptx, out, "--force")
            self.assertEqual(forced.returncode, 0, forced.stderr)
            self.assertEqual(starter.read_text(), self.starter_text)

    def test_a_file_that_is_not_a_deck_is_refused(self):
        with tempfile.TemporaryDirectory() as tmp:
            fake = Path(tmp) / "notes.pptx"
            fake.write_text("not a deck")
            result = run(fake, Path(tmp) / "out")
            self.assertEqual(result.returncode, 2)
            self.assertIn("not a PowerPoint deck", result.stderr)
            self.assertFalse((Path(tmp) / "out").exists())
            missing = run(Path(tmp) / "absent.pptx", Path(tmp) / "out")
            self.assertEqual(missing.returncode, 2)
            self.assertIn("No such deck", missing.stderr)

    def test_id_overrides_the_slug(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "out"
            result = run(self.pptx, out, "--id", "q3-pack")
            self.assertEqual(result.returncode, 0, result.stderr)
            inventory = json.loads((out / "q3-pack.inventory.json").read_text())
            starter = json.loads((out / "q3-pack.pages.json").read_text())
            self.assertEqual(inventory["id"], "q3-pack")
            self.assertEqual((starter["deck"]["id"], starter["deck"]["inventory"]), ("q3-pack", "q3-pack.inventory.json"))
            self.assertFalse((out / "board-pack-q3-v2.pages.json").exists())


if __name__ == "__main__":
    unittest.main()
