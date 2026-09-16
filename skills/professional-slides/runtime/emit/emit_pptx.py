#!/usr/bin/env python3
"""Emit an editable PowerPoint from a resolved scene.

Design goals (see deep-audit-and-revamp.md, items 1, 2, 3, 18):
  * The action title is a real title placeholder, so Outline view, Reset Slide and
    template swaps work.
  * Every text box wraps (`wrap="square"`) and carries `<a:normAutofit/>`; paragraphs
    are the author's paragraphs, not the layout engine's wrapped lines. The engine's
    measurement sizes the box; PowerPoint owns the final wrap.
  * Charts are native, workbook-backed chart objects when the component is one
    PowerPoint can express; assembled-shape diagrams are grouped so they move as one.
  * No vendor runtime: python-pptx only. Theme fonts and colours are patched into
    theme1.xml after save so the deck's palette is the PowerPoint theme.

Usage: emit_pptx.py scene.json out.pptx [--no-native-charts] [--no-groups]
"""
from __future__ import annotations

import argparse
import base64
import io
import json
import re
import sys
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_AXIS_CROSSES, XL_TICK_MARK
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt
from lxml import etree

PX = 914400 / 96  # EMU per CSS px at 96 px/in; slide is 1280x720 px = 13.333x7.5 in

# Roles a native chart replaces. Headings, units and rules stay as editable text.
CHART_PLOT_ROLES = {"chart-mark", "data-label", "category-label", "chart-axis", "axis-label", "chart-gridline",
                    "legend-swatch", "legend-label", "chart-line", "chart-point", "chart-segment", "chart-area",
                    "chart-wedge", "pie-label", "pie-leader", "chart-baseline", "chart-tick", "reference-line",
                    "reference-label", "chart-annotation", "annotation-leader", "chart-callout", "chart-highlight",
                    "value-label", "series-label", "end-label", "stack-label", "total-label", "axis-title",
                    "chart-hole", "chart-marker", "chart-connector", "chart-bar-axis", "chart-quadrant", "chart-quadrant-label"}
# Single-line labels are sized to their ink; they never wrap, so PowerPoint must not
# re-wrap them on a one-pixel advance difference. Prose keeps wrap="square".
LABEL_ROLES = {"legend-label", "data-label", "category-label", "axis-label", "value-label", "metric-value", "metric-label",
               "metric-delta", "page-number", "source-text", "chart-unit", "process-label", "tracker-label", "table-cell",
               "table-header", "pie-label", "reference-label", "end-label", "stack-label", "total-label", "scale-endpoint", "page-tag", "cover-date", "cover-logo", "table-status-label", "table-progress-label", "chart-badge"}
CHROME_COMPONENTS = {"slide-chrome", "page-template", "section", "paragraph", "section-heading",
                     "bullet-list", "insight", "evidence-note", "chart-title", "footnote", "cover"}
NATIVE = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "stacked-column": XL_CHART_TYPE.COLUMN_STACKED,
    "stacked-bar": XL_CHART_TYPE.BAR_STACKED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "donut": XL_CHART_TYPE.DOUGHNUT,
    "area": XL_CHART_TYPE.AREA,
    "scatter": XL_CHART_TYPE.XY_SCATTER,
    "range": XL_CHART_TYPE.BAR_STACKED,
}


def _luminance(hex_color):
    """Relative luminance (0..1) of a #RRGGBB colour, for label contrast."""
    try:
        h = str(hex_color).lstrip("#")
        r, g, b = (int(h[i:i + 2], 16) / 255.0 for i in (0, 2, 4))
    except Exception:
        return 1.0
    lin = lambda c: c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
    return 0.2126 * lin(r) + 0.7152 * lin(g) + 0.0722 * lin(b)


def emu(px: float) -> int:
    return int(round(px * PX))


def rgb(value: str | None) -> RGBColor | None:
    if not value:
        return None
    v = value.lstrip("#")
    if len(v) == 3:
        v = "".join(c * 2 for c in v)
    return RGBColor.from_string(v.upper())


def style_value(style: dict, key: str, default=None):
    v = style.get(key)
    if isinstance(v, dict) and "value" in v:
        return v["value"]
    return default if v is None else v


def token_colors(scene: dict) -> dict:
    tokens = scene.get("tokens") or (scene["slides"][0]["tokens"] if scene.get("slides") else {})
    return {k: v.get("value") for k, v in tokens.items() if v.get("kind") == "color"}


class Emitter:
    def __init__(self, scene: dict, *, native_charts=True, groups=True):
        self.scene = scene
        self.native_charts = native_charts
        self.groups = groups
        self.prs = Presentation()
        self.prs.slide_width = emu(1280)
        self.prs.slide_height = emu(720)
        self.title_layout = self.prs.slide_layouts[5]  # Title Only
        self.blank_layout = self.prs.slide_layouts[6]
        self.colors = token_colors(scene)
        self.series_colors = [self.colors.get(f"color.chartSeries{i}") for i in range(1, 7)]
        self.stats = {"slides": 0, "text": 0, "native_charts": 0, "grouped": 0, "shapes": 0, "images": 0}

    # ---------------------------------------------------------------- text
    @staticmethod
    def paragraphs_of(node: dict) -> list[str]:
        layout = (node.get("data") or {}).get("textLayout") or {}
        source = layout.get("source")
        if source is None:
            # Legacy scene without source: collapse the engine's wrapped lines. A blank
            # line marks a real paragraph break; other breaks are wraps.
            lines = layout.get("lines") or str(node.get("text", "")).split("\n")
            paras, cur = [], []
            for line in lines:
                if line.strip() == "":
                    paras.append(" ".join(cur)); cur = []
                else:
                    cur.append(line.strip())
            paras.append(" ".join(cur))
            return paras
        return str(source).split("\n")

    def fill_text_frame(self, tf, node: dict):
        style = node.get("style") or {}
        size = float(style_value(style, "fontSize", 14))
        family = style_value(style, "fontFamily", "Arial")
        color = rgb(style_value(style, "color", "#000000"))
        bold = bool(style.get("bold"))
        align = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(style.get("align"), PP_ALIGN.LEFT)
        layout0 = (node.get("data") or {}).get("textLayout") or {}
        single_label = node.get("role") in LABEL_ROLES and len(layout0.get("lines") or [1]) == 1
        tf.word_wrap = not single_label
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE  # <a:normAutofit/>
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "mid": MSO_ANCHOR.MIDDLE, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}.get(style.get("valign"), MSO_ANCHOR.TOP)
        for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, side, 0)
        layout = (node.get("data") or {}).get("textLayout") or {}
        runs = layout.get("sourceRuns") or node.get("runs")
        paragraphs = self.paragraphs_of(node)
        first = True
        if runs:
            # Runs vary only in bold; rebuild them per paragraph from the source text.
            text_all = "".join(r["text"] for r in runs)
            src_paras = text_all.split("\n")
            offset = 0
            for pi, para_text in enumerate(src_paras):
                p = tf.paragraphs[0] if first else tf.add_paragraph()
                first = False
                p.alignment = align
                start, end = offset, offset + len(para_text)
                pos = 0
                for r in runs:
                    r_start, r_end = pos, pos + len(r["text"])
                    pos = r_end
                    lo, hi = max(start, r_start), min(end, r_end)
                    if hi <= lo:
                        continue
                    run = p.add_run()
                    run.text = text_all[lo:hi]
                    accent = rgb(self.colors.get("color.accent") or self.colors.get("color.componentPrimary")) if r.get("accent") else None
                    self._font(run.font, family, size, bold or r.get("bold", False), accent or color)
                offset = end + 1
        else:
            for para_text in paragraphs:
                p = tf.paragraphs[0] if first else tf.add_paragraph()
                first = False
                p.alignment = align
                run = p.add_run()
                run.text = para_text
                self._font(run.font, family, size, bold, color)
        for p in tf.paragraphs:
            p.line_spacing = 1.0  # font-native line box; matches the engine's 1.12em for Arial

    @staticmethod
    def _font(font, family, size, bold, color):
        font.name = family
        font.size = Pt(size)
        font.bold = bool(bold)
        if color is not None:
            font.color.rgb = color

    def add_text(self, slide, node: dict, *, as_title=False):
        f = node["frame"]
        if as_title:
            shape = slide.shapes.title
            shape.left, shape.top, shape.width, shape.height = emu(f["x"]), emu(f["y"]), emu(f["width"]), emu(f["height"])
            tf = shape.text_frame
            # clear placeholder default paragraph
            for p in list(tf.paragraphs)[1:]:
                p._p.getparent().remove(p._p)
            tf.paragraphs[0].text = ""
        else:
            shape = slide.shapes.add_textbox(emu(f["x"]), emu(f["y"]), emu(f["width"]), emu(f["height"]))
            tf = shape.text_frame
        shape.name = f"ps:{node['id']}"
        self.fill_text_frame(tf, node)
        self.stats["text"] += 1
        return shape

    # --------------------------------------------------------------- shapes
    def add_rect(self, slide, node: dict):
        f = node["frame"]; style = node.get("style") or {}
        radius = float(style_value(style, "radius", 0) or 0)
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius > 0 else MSO_SHAPE.RECTANGLE
        shape = slide.shapes.add_shape(kind, emu(f["x"]), emu(f["y"]), emu(max(f["width"], 0.5)), emu(max(f["height"], 0.5)))
        shape.name = f"ps:{node['id']}"
        if radius > 0 and f["width"] and f["height"]:
            shape.adjustments[0] = min(0.5, radius / min(f["width"], f["height"]))
        fill = style_value(style, "fill", None)
        if fill and fill != "none":
            shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill)
        else:
            shape.fill.background()
        stroke = style_value(style, "stroke", None)
        if stroke and stroke != "none":
            shape.line.color.rgb = rgb(stroke)
            shape.line.width = Pt(float(style_value(style, "lineWidth", 1) or 1) * 0.75)
        else:
            shape.line.fill.background()
        self._strip_style(shape)
        self._strip_text(shape)
        self.stats["shapes"] += 1
        return shape

    def add_ellipse(self, slide, node: dict):
        f = node["frame"]; style = node.get("style") or {}
        shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, emu(f["x"]), emu(f["y"]), emu(f["width"]), emu(f["height"]))
        shape.name = f"ps:{node['id']}"
        fill = style_value(style, "fill", None)
        if fill and fill != "none":
            shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill)
        else:
            shape.fill.background()
        stroke = style_value(style, "stroke", None)
        if stroke and stroke != "none":
            shape.line.color.rgb = rgb(stroke); shape.line.width = Pt(float(style_value(style, "lineWidth", 1) or 1) * 0.75)
        else:
            shape.line.fill.background()
        self._strip_style(shape)
        self._strip_text(shape)
        self.stats["shapes"] += 1
        return shape

    # Preset geometries the scene uses, mapped onto PowerPoint's own autoshapes so
    # a chevron stays a chevron with its yellow adjustment handle.
    PRESETS = {
        "chevron": MSO_SHAPE.CHEVRON, "rightArrow": MSO_SHAPE.RIGHT_ARROW, "leftArrow": MSO_SHAPE.LEFT_ARROW,
        "notchedRightArrow": MSO_SHAPE.NOTCHED_RIGHT_ARROW, "snip1Rect": MSO_SHAPE.SNIP_1_RECTANGLE,
        "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE, "rtTriangle": MSO_SHAPE.RIGHT_TRIANGLE, "diamond": MSO_SHAPE.DIAMOND,
        "pentagon": MSO_SHAPE.REGULAR_PENTAGON, "hexagon": MSO_SHAPE.HEXAGON, "parallelogram": MSO_SHAPE.PARALLELOGRAM,
        "trapezoid": MSO_SHAPE.TRAPEZOID, "downArrow": MSO_SHAPE.DOWN_ARROW, "upArrow": MSO_SHAPE.UP_ARROW,
        "homePlate": MSO_SHAPE.PENTAGON, "roundRect": MSO_SHAPE.ROUNDED_RECTANGLE,
    }

    def _paint(self, shape, style: dict):
        fill = style_value(style, "fill", None)
        if fill and fill != "none":
            shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill)
            opacity = style.get("opacity")
            if opacity is not None and float(opacity) < 1:
                sf = shape.fill._xPr.find(qn("a:solidFill"))
                clr = sf[0] if sf is not None and len(sf) else None
                if clr is not None:
                    etree.SubElement(clr, qn("a:alpha")).set("val", str(int(float(opacity) * 100000)))
        else:
            shape.fill.background()
        stroke = style_value(style, "stroke", None)
        if stroke and stroke != "none":
            shape.line.color.rgb = rgb(stroke)
            shape.line.width = Pt(float(style_value(style, "lineWidth", 1) or 1) * 0.75)
            if style.get("lineCap") == "round" or style.get("iconStroke"):
                ln = shape.line._get_or_add_ln()
                ln.set("cap", "rnd")
                join = etree.SubElement(ln, qn("a:round"))
        else:
            shape.line.fill.background()
        rotate = float(style.get("rotate") or 0)
        if rotate:
            shape.rotation = rotate
        if style.get("flipH") or style.get("flipV"):
            xfrm = shape._element.spPr.find(qn("a:xfrm"))
            if xfrm is not None:
                if style.get("flipH"): xfrm.set("flipH", "1")
                if style.get("flipV"): xfrm.set("flipV", "1")

    def _freeform(self, slide, node: dict, paths, *, closed=True):
        """paths: list of point lists in scene px (absolute). One freeform, several subpaths."""
        f = node["frame"]
        first = paths[0][0]
        fb = slide.shapes.build_freeform(emu(first[0]), emu(first[1]), scale=1.0)
        for i, path in enumerate(paths):
            pts = [(emu(x), emu(y)) for x, y in path]
            if i:
                fb.move_to(*pts[0])
            fb.add_line_segments(pts[1:], close=closed)
        shape = fb.convert_to_shape(0, 0)
        shape.name = f"ps:{node['id']}"
        self._paint(shape, node.get("style") or {})
        self._strip_style(shape)
        self._strip_text(shape)
        self.stats["shapes"] += 1
        return shape

    def add_shape(self, slide, node: dict):
        f = node["frame"]; style = node.get("style") or {}; data = node.get("data") or {}
        geometry = data.get("geometry", "rect")
        if geometry == "customPolygon":
            paths = [[(f["x"] + float(x) * f["width"], f["y"] + float(y) * f["height"]) for x, y in path] for path in data.get("paths") or []]
            return self._freeform(slide, node, paths) if paths else self.add_rect(slide, node)
        if geometry == "iconPath":
            # Icon strokes: unit-square polylines, open unless the path says closed.
            paths = data.get("paths") or []
            pts = [[(f["x"] + float(x) * f["width"], f["y"] + float(y) * f["height"]) for x, y in path["points"]] for path in paths]
            closed = [bool(path.get("closed")) for path in paths]
            first = pts[0][0]
            fb = slide.shapes.build_freeform(emu(first[0]), emu(first[1]), scale=1.0)
            for i, path in enumerate(pts):
                ep = [(emu(x), emu(y)) for x, y in path]
                if i:
                    fb.move_to(*ep[0])
                fb.add_line_segments(ep[1:], close=closed[i])
            shape = fb.convert_to_shape(0, 0)
            shape.name = f"ps:{node['id']}"
            self._paint(shape, {**style, "iconStroke": True})
            self._strip_style(shape); self._strip_text(shape)
            self.stats["shapes"] += 1
            return shape
        if geometry == "quoteCallout":
            body = f["height"] * float(data.get("bodyRatio", 0.8))
            cc = f["x"] + f["width"] * float(data.get("caretCenterRatio", 0.2))
            ch = f["width"] * float(data.get("caretWidthRatio", 0.1)) / 2
            left, right, top, bottom, tip = f["x"], f["x"] + f["width"], f["y"], f["y"] + body, f["y"] + f["height"]
            return self._freeform(slide, node, [[(left, top), (right, top), (right, bottom), (cc + ch, bottom), (cc, tip), (cc - ch, bottom), (left, bottom)]])
        preset = self.PRESETS.get(geometry)
        if preset is None:
            return self.add_rect(slide, node)
        shape = slide.shapes.add_shape(preset, emu(f["x"]), emu(f["y"]), emu(max(f["width"], 0.5)), emu(max(f["height"], 0.5)))
        shape.name = f"ps:{node['id']}"
        try:
            if geometry == "chevron" and f["width"] and f["height"]:
                # PowerPoint's adj is a share of the shorter side; the scene draws the
                # point inset at half that side, which is the preset default (50000).
                shape.adjustments[0] = 0.5
            elif geometry == "rightArrow" and f["width"] and f["height"]:
                shape.adjustments[0] = 0.64
                shape.adjustments[1] = min(1.0, 0.28 * f["width"] / min(f["width"], f["height"]))
            elif geometry == "snip1Rect" and f["width"] and f["height"]:
                shape.adjustments[0] = min(0.5, 0.12 * f["width"] / min(f["width"], f["height"]))
        except (IndexError, ValueError):
            pass
        self._paint(shape, style)
        self._strip_style(shape)
        self._strip_text(shape)
        self.stats["shapes"] += 1
        return shape

    def add_wedge(self, slide, node: dict):
        import math
        f = node["frame"]; data = node.get("data") or {}
        cx, cy = f["x"] + f["width"] / 2, f["y"] + f["height"] / 2
        r = min(f["width"], f["height"]) / 2
        a0, a1 = float(data.get("startAngle", 0)), float(data.get("endAngle", 90))
        steps = max(2, int(abs(a1 - a0) / 4))
        pts = [(cx, cy)] + [(cx + r * math.cos(math.radians(a0 + (a1 - a0) * i / steps)), cy + r * math.sin(math.radians(a0 + (a1 - a0) * i / steps))) for i in range(steps + 1)]
        return self._freeform(slide, node, [pts])

    def add_line(self, slide, node: dict):
        f = node["frame"]; style = node.get("style") or {}
        d = node.get("data") or {}
        # Line endpoints live on the node (or its data); the frame is only the
        # bounding box, so a rising line must not be read as top-left → bottom-right.
        pick = lambda key, fallback: node.get(key, d.get(key, f.get(key, fallback)))
        x1 = pick("x1", f["x"]); y1 = pick("y1", f["y"])
        x2 = pick("x2", f["x"] + f["width"]); y2 = pick("y2", f["y"] + f["height"])
        conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, emu(x1), emu(y1), emu(x2), emu(y2))
        conn.name = f"ps:{node['id']}"
        self._strip_style(conn)
        stroke = style_value(style, "stroke", None) or style_value(style, "color", "#000000")
        conn.line.color.rgb = rgb(stroke)
        conn.line.width = Pt(float(style_value(style, "lineWidth", 1) or 1) * 0.75)
        dash = style.get("dash") or style.get("lineDash")
        if dash and str(dash).lower() not in ("solid", "none", "false"):
            from pptx.enum.dml import MSO_LINE_DASH_STYLE
            conn.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        data = node.get("data") or {}
        if data.get("endArrow") or data.get("startArrow"):
            ln = conn.line._get_or_add_ln()
            if data.get("startArrow"):
                etree.SubElement(ln, qn("a:headEnd")).set("type", "triangle")
            if data.get("endArrow"):
                etree.SubElement(ln, qn("a:tailEnd")).set("type", "triangle")
        self.stats["shapes"] += 1
        return conn

    def add_image(self, slide, node: dict):
        f = node["frame"]
        data = (node.get("data") or {})
        uri = data.get("dataUri") or node.get("dataUri") or data.get("src")
        if not uri:
            return None
        if str(uri).startswith("data:"):
            blob = base64.b64decode(uri.split(",", 1)[1])
            stream = io.BytesIO(blob)
        else:
            stream = str(uri)
        pic = slide.shapes.add_picture(stream, emu(f["x"]), emu(f["y"]), emu(f["width"]), emu(f["height"]))
        pic.name = f"ps:{node['id']}"
        self.stats["images"] += 1
        return pic

    @staticmethod
    def _strip_style(shape):
        """Remove the theme style reference so no shadow/effect is inherited."""
        el = shape._element
        for style in el.findall(qn("p:style")):
            el.remove(style)

    @staticmethod
    def _strip_text(shape):
        # autoshapes get an empty <p:txBody>; leave it but ensure no phantom text
        if shape.has_text_frame:
            shape.text_frame.text = ""

    # --------------------------------------------------------------- charts
    def add_native_chart(self, slide, instance: dict):
        spec = instance["nativeChart"]
        f = dict(spec.get("frame") or instance["frame"])
        band = instance.get("_headingBottom")
        if band is not None and band > f["y"]:
            gap = 8
            f["height"] -= (band + gap - f["y"]); f["y"] = band + gap
        ctype = NATIVE.get(spec["type"])
        if ctype is None:
            return None
        if spec["type"] in ("pie", "donut"):
            # A square frame keeps the hole at the frame centre in every renderer,
            # so the centre KPI box lands in the hole.
            side = min(f["width"], f["height"])
            f = {"x": f["x"] + (f["width"] - side) / 2, "y": f["y"] + (f["height"] - side) / 2, "width": side, "height": side}
        if spec["type"] == "scatter":
            cd = XyChartData()
            for s in spec.get("series") or [{"name": "", "values": []}]:
                ser = cd.add_series(s.get("name") or "Series")
                for p in spec.get("points") or []:
                    ser.add_data_point(p.get("x"), p.get("y"))
        else:
            cd = CategoryChartData()
            cd.categories = spec["categories"]
            for s in spec["series"]:
                cd.add_series(s.get("name") or "Series", s["values"])
        gf = slide.shapes.add_chart(ctype, emu(f["x"]), emu(f["y"]), emu(f["width"]), emu(f["height"]), cd)
        gf.name = f"ps:{instance['instanceId']}:chart"
        chart = gf.chart
        chart.font.name = "Arial"
        chart.font.size = Pt(10)
        chart.has_title = False
        chart.has_legend = bool(spec.get("legend"))
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.TOP
            chart.legend.include_in_layout = False
            chart.legend.font.size = Pt(10)
        plot = chart.plots[0]
        kind = spec["type"]
        is_range = kind == "range"
        fmt = spec.get("valueFormat") or {}
        decimals = int(fmt.get("decimals", 0)) if isinstance(fmt, dict) else 0
        number_format = "0" if decimals == 0 else "0." + "0" * decimals
        if kind not in ("pie", "donut", "scatter"):
            plot.gap_width = 60 if not is_range else 80
            if kind in ("column", "bar"):
                plot.overlap = 0
            if is_range:
                plot.overlap = 100
            va = chart.value_axis
            va.has_major_gridlines = bool(spec.get("gridlines"))
            va.visible = False if spec.get("dataLabels", True) else True
            va.tick_labels.font.size = Pt(10)  # chart furniture floor is 10 pt
            if spec.get("yMin") is not None:
                va.minimum_scale = spec["yMin"]
            if spec.get("yMax") is not None:
                va.maximum_scale = spec["yMax"]
            ca = chart.category_axis
            ca.tick_labels.font.size = Pt(10)
            if kind in ("bar", "stacked-bar", "range"):
                ca.reverse_order = True  # first category at the top, as authored
                va.crosses = XL_AXIS_CROSSES.MAXIMUM if False else va.crosses
            ca.has_major_gridlines = False
            ca.major_tick_mark = XL_TICK_MARK.NONE
            va.major_tick_mark = XL_TICK_MARK.NONE
            ca.format.line.color.rgb = rgb(self.colors.get("color.rule", "#929BA3"))
            va.format.line.fill.background()
        if spec.get("dataLabels", True) and not is_range:
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.font.size = Pt(11)   # data labels are the chart's loudest number
            dl.font.bold = True
            dl.number_format = number_format
            dl.number_format_is_linked = False
            if kind in ("column", "bar"):
                dl.position = XL_LABEL_POSITION.OUTSIDE_END
            elif kind in ("pie", "donut"):
                dl.position = XL_LABEL_POSITION.OUTSIDE_END if kind == "pie" else XL_LABEL_POSITION.CENTER
        # series colours from the palette (comparator series grey when the runtime would)
        idx = spec.get("colorIndices")
        accent = self.colors.get("color.accent") or self.colors.get("color.componentPrimary")
        forecast = self.colors.get("color.chartSeries6")
        highlight_indices = set(spec.get("highlightIndices") or [])
        forecast_index = spec.get("forecastIndex", -1)
        single = len(spec["series"]) == 1
        max_value = max((abs(v) for s in spec["series"] for v in s["values"] if isinstance(v, (int, float))), default=0)
        for i, ser in enumerate(plot.series):
            ci = idx[i] if idx and i < len(idx) else i
            color = self.series_colors[ci % len(self.series_colors)] if self.series_colors else None
            if len(spec["series"]) == 2 and i == 1 and not idx and not is_range and kind not in ("line", "area"):
                color = self.colors.get("color.chartComparator", color)
            if is_range and i == 0:
                # invisible base: the bar floats from low to high
                ser.format.fill.background(); ser.format.line.fill.background()
                continue
            if color and kind not in ("pie", "donut"):
                fill = ser.format.fill
                fill.solid(); fill.fore_color.rgb = rgb(color)
                if kind in ("stacked-column", "stacked-bar") and spec.get("dataLabels", True):
                    # Segment labels sit inside their segment, white on a dark
                    # fill and ink on a light one, so every segment reads.
                    sdl = ser.data_labels
                    sdl.show_value = True
                    sdl.number_format = number_format; sdl.number_format_is_linked = False
                    sdl.font.size = Pt(11); sdl.font.bold = True
                    sdl.position = XL_LABEL_POSITION.CENTER
                    sdl.font.color.rgb = rgb(self.colors.get("color.onPrimary", "#FFFFFF") if _luminance(color) < 0.45 else self.colors.get("color.ink", "#000000"))
                if kind == "line":
                    ser.format.line.color.rgb = rgb(color)
                    ser.format.line.width = Pt(2.25)
                    ser.smooth = False
                if kind in ("column", "bar", "stacked-column", "stacked-bar") and single and spec.get("dataLabels", True):
                    # Series-level labels first (a per-point override otherwise
                    # creates a series block that hides the other labels), then
                    # highlight the answer, lighten the forecast, and put the label
                    # inside a bar that is wide enough to carry it in white.
                    sdl = ser.data_labels
                    sdl.show_value = True
                    sdl.number_format = number_format; sdl.number_format_is_linked = False
                    sdl.font.size = Pt(11); sdl.font.bold = True
                    sdl.font.color.rgb = rgb(self.colors.get("color.ink", "#000000"))
                    if kind in ("column", "bar"):
                        sdl.position = XL_LABEL_POSITION.OUTSIDE_END
                    for j, pt in enumerate(ser.points):
                        value = spec["series"][i]["values"][j]
                        if j in highlight_indices and accent:
                            pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(accent)
                        elif forecast_index is not None and forecast_index >= 0 and j >= forecast_index and forecast:
                            pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(forecast)
                        if kind in ("column", "bar") and max_value and isinstance(value, (int, float)) and abs(value) >= 0.4 * max_value:
                            lab = pt.data_label
                            lab.position = XL_LABEL_POSITION.INSIDE_END
                            lab.font.size = Pt(11); lab.font.bold = True
                            lab.font.color.rgb = rgb(self.colors.get("color.onPrimary", "#FFFFFF"))
                            lab.number_format = number_format; lab.number_format_is_linked = False
                if kind == "line" and spec.get("endLabels"):
                    # Series name at the last point instead of a legend.
                    last = len(spec["series"][i]["values"]) - 1
                    if last >= 0:
                        lab = ser.points[last].data_label
                        lab.position = XL_LABEL_POSITION.RIGHT
                        tf = lab.text_frame
                        tf.text = str(spec["series"][i].get("name") or "")
                        for p in tf.paragraphs:
                            for r in p.runs:
                                r.font.size = Pt(10); r.font.bold = True; r.font.color.rgb = rgb(color)
                if is_range:
                    lows, highs = spec.get("low") or [], spec.get("high") or []
                    for j, pt in enumerate(ser.points):
                        if j in highlight_indices and accent:
                            pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(accent)
                        lab = pt.data_label
                        lab.position = XL_LABEL_POSITION.INSIDE_END
                        tf = lab.text_frame
                        lo = lows[j] if j < len(lows) else ""; hi = highs[j] if j < len(highs) else ""
                        tf.text = f"{lo:g}–{hi:g}" if isinstance(lo, (int, float)) and isinstance(hi, (int, float)) else f"{lo}–{hi}"
                        for p in tf.paragraphs:
                            for r in p.runs:
                                r.font.size = Pt(10); r.font.bold = True; r.font.color.rgb = rgb(self.colors.get("color.onPrimary", "#FFFFFF"))
            elif kind in ("pie", "donut"):
                for j, pt in enumerate(ser.points):
                    c = self.series_colors[j % len(self.series_colors)]
                    if c:
                        pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(c)
        if kind == "donut" and spec.get("center"):
            # The centre KPI: a text box over the hole (value bold, label under).
            center = spec["center"] if isinstance(spec["center"], dict) else {"value": str(spec["center"])}
            w, h = f["width"] * 0.34, f["height"] * 0.3
            box = slide.shapes.add_textbox(emu(f["x"] + (f["width"] - w) / 2), emu(f["y"] + (f["height"] - h) / 2), emu(w), emu(h))
            box.name = f"ps:{instance['instanceId']}:center"
            tf = box.text_frame; tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
            r = p.add_run(); r.text = str(center.get("value", ""))
            self._font(r.font, "Arial", 20, True, rgb(self.colors.get("color.ink", "#000000")))
            if center.get("label"):
                p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
                r2 = p2.add_run(); r2.text = str(center["label"])
                self._font(r2.font, "Arial", 10, False, rgb(self.colors.get("color.textSecondary", "#404040")))
        self.stats["native_charts"] += 1
        return gf

    # ---------------------------------------------------------------- slide
    def emit_slide(self, sl: dict):
        nodes = sl["nodes"]
        instances = {ci["instanceId"]: ci for ci in sl.get("componentInstances", [])}
        title_node = next((n for n in nodes if n["type"] == "text" and n.get("role") in ("action-title", "cover-title", "deck-title", "section-title")), None)
        slide = self.prs.slides.add_slide(self.title_layout if title_node else self.blank_layout)
        skip_instances = set()
        if self.native_charts:
            for iid, ci in instances.items():
                if ci.get("nativeChart"):
                    heads = [n for n in nodes if (n.get("data") or {}).get("componentInstance") == iid
                             and n.get("role") in ("section-heading", "section-heading-rule", "chart-unit", "chart-heading", "chart-title")]
                    ci["_headingBottom"] = max((n["frame"]["y"] + n["frame"]["height"] for n in heads), default=None)
                    self.add_native_chart(slide, ci)
                    skip_instances.add(iid)
        # group membership: diagram-like components with several primitives
        member_shapes: dict[str, list] = {}
        title_shape = None
        for node in nodes:
            inst = (node.get("data") or {}).get("componentInstance")
            if inst in skip_instances and node.get("role") in CHART_PLOT_ROLES:
                continue
            if node is title_node:
                shape = title_shape = self.add_text(slide, node, as_title=True)
            elif node["type"] == "text":
                shape = self.add_text(slide, node)
            elif node["type"] == "rect":
                shape = self.add_rect(slide, node)
            elif node["type"] == "ellipse":
                shape = self.add_ellipse(slide, node)
            elif node["type"] == "line":
                shape = self.add_line(slide, node)
            elif node["type"] == "image":
                shape = self.add_image(slide, node)
            elif node["type"] == "shape":
                shape = self.add_shape(slide, node)
            elif node["type"] == "wedge":
                shape = self.add_wedge(slide, node)
            else:
                shape = self.add_rect(slide, node)
            if shape is not None and inst:
                member_shapes.setdefault(inst, []).append(shape)
        if title_shape is not None:
            # The layout creates the placeholder first; surfaces drawn after it
            # (a dark cover) would hide it. Put it on top of the z-order.
            tree = title_shape._element.getparent()
            tree.remove(title_shape._element); tree.append(title_shape._element)
        if self.groups:
            for iid, shapes in member_shapes.items():
                comp = instances.get(iid, {}).get("component", "")
                if len(shapes) >= 3 and comp not in CHROME_COMPONENTS and not comp.startswith("chart-"):
                    self._group(slide, shapes, f"ps:{iid}:group")
        if sl.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(sl["notes"])
        self.stats["slides"] += 1

    def _group(self, slide, shapes, name):
        grp = slide.shapes.add_group_shape(shapes)
        grp.name = name
        self.stats["grouped"] += 1

    # ---------------------------------------------------------------- theme
    def patch_theme(self, path: Path):
        """Write the scene palette into theme1.xml so the deck's theme *is* the palette."""
        slots = {}
        tokens = self.scene["slides"][0]["tokens"] if self.scene.get("slides") else {}
        for tid, t in tokens.items():
            if t.get("kind") == "color" and t.get("themeSlot"):
                slots[t["themeSlot"]] = t["value"].lstrip("#").upper()
        fonts = {t.get("value") for t in tokens.values() if t.get("kind") == "fontFamily"}
        family = "Arial" if "Arial" in fonts or not fonts else sorted(fonts)[0]
        buf = io.BytesIO()
        with zipfile.ZipFile(path) as zin:
            with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
                for item in zin.infolist():
                    data = zin.read(item.filename)
                    if item.filename == "ppt/theme/theme1.xml":
                        root = etree.fromstring(data)
                        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
                        for slot, hexv in slots.items():
                            el = root.find(f".//a:clrScheme/a:{slot}", ns)
                            if el is not None:
                                for child in list(el):
                                    el.remove(child)
                                etree.SubElement(el, qn("a:srgbClr")).set("val", hexv)
                        for tag in ("majorFont", "minorFont"):
                            latin = root.find(f".//a:fontScheme/a:{tag}/a:latin", ns)
                            if latin is not None:
                                latin.set("typeface", family)
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    zout.writestr(item, data)
        path.write_bytes(buf.getvalue())

    def run(self, out: Path):
        for sl in self.scene["slides"]:
            self.emit_slide(sl)
        self.prs.save(str(out))
        self.patch_theme(out)
        return self.stats


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("scene"); ap.add_argument("out")
    ap.add_argument("--no-native-charts", action="store_true")
    ap.add_argument("--no-groups", action="store_true")
    a = ap.parse_args(argv)
    scene = json.loads(Path(a.scene).read_text())
    stats = Emitter(scene, native_charts=not a.no_native_charts, groups=not a.no_groups).run(Path(a.out))
    print(json.dumps({"out": a.out, **stats}))


if __name__ == "__main__":
    main()
