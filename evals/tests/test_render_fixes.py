"""Five finish defects an Emirates deck shipped with.

Callout leaders landed on a column's right edge, a pie's native labels
wrapped "75.8" one character per line, grey forecast columns carried no key,
a reference label with callouts in the plot threw instead of moving outside,
and table logos were held to one body line so square marks became specks.
Each test renders the component the way the deck did and reads the geometry.
"""
import tempfile
import unittest
from pathlib import Path

from node_probe import run_node, requires_python_package

needs_pptx = requires_python_package('pptx')
try:
    from test_pr_review_export import Emitter
    from pptx import Presentation
except ImportError:  # pragma: no cover - exercised only without python-pptx
    Emitter = Presentation = None


QATAR = """{heading:'Qatar airline passengers',unit:'million',categories:['FY24','FY25','FY26','FY27','FY28','FY29'],
 series:[{name:'Passengers',values:[40,43.1,41.8,45.3,49.1,53.2]}],forecastFrom:'FY27'}"""


class CalloutLeaderTests(unittest.TestCase):
    def test_leader_lands_on_the_column_top_centre_and_the_bar_end_centre(self):
        result = run_node("""
import {REGISTRY} from './skills/professional-slides/runtime/registry.mjs';
const frame={x:72,y:180,width:1136,height:460};
const render=(type,props)=>REGISTRY.get(type).render({id:'c',frame,props}).nodes;
const column=render('chart.column',{heading:'Mix',unit:'%',categories:['Europe','E Asia','Americas','Africa'],series:[{name:'Share',values:[31,28,16,8]}],annotations:[{category:'E Asia',text:'Europe and East Asia together: 59%'}]});
const bar=render('chart.bar',{heading:'Mix',unit:'%',categories:['Europe','E Asia','Americas'],series:[{name:'Share',values:[31,28,16]}],annotations:[{category:'Europe',text:'The largest region'}]});
const pick=(nodes,category)=>({mark:nodes.find(n=>n.role==='chart-mark'&&n.data.category===category).frame,
  label:nodes.find(n=>n.role==='data-label'&&n.data.category===category).frame,
  leader:nodes.find(n=>n.role==='annotation-leader').data});
console.log(JSON.stringify({column:pick(column,'E Asia'),bar:pick(bar,'Europe')}));
""")
        column, bar = result['column'], result['bar']
        # Top-centre of the column, stopping above its value label rather than
        # striking through the number.
        self.assertAlmostEqual(column['leader']['x2'], column['mark']['x'] + column['mark']['width'] / 2, places=3)
        self.assertLess(column['leader']['y2'], column['label']['y'] + column['label']['height'] / 2)
        self.assertGreater(column['leader']['y2'], column['label']['y'] - 16)
        # End-centre of a horizontal bar.
        self.assertAlmostEqual(bar['leader']['x2'], bar['mark']['x'] + bar['mark']['width'], places=3)
        self.assertAlmostEqual(bar['leader']['y2'], bar['mark']['y'] + bar['mark']['height'] / 2, places=3)


class PieLabelTests(unittest.TestCase):
    def test_labels_take_their_measured_width_and_a_thin_slice_goes_outside(self):
        result = run_node("""
import {REGISTRY} from './skills/professional-slides/runtime/registry.mjs';
import {nativeChartSpec} from './skills/professional-slides/runtime/core.mjs';
const frame={x:72,y:200,width:673,height:370};
const pie=REGISTRY.get('chart.pie').render({id:'p',frame,props:{labels:['International','Domestic'],values:[75.8,65.1]}}).nodes;
const donutProps={labels:['Core','Growth','New'],values:[90,6,4]};
const donut=REGISTRY.get('chart.donut').render({id:'d',frame,props:donutProps}).nodes;
const labels=nodes=>nodes.filter(n=>n.role==='data-label').map(n=>({text:n.text,placement:n.data.placement,frame:n.frame}));
const circle=donut.find(n=>n.role==='chart-segment').frame;
console.log(JSON.stringify({pie:labels(pie),donut:labels(donut),circle,
  native:nativeChartSpec('chart.donut',donutProps,frame,donut)}));
""")
        self.assertEqual([l['text'] for l in result['pie']], ['54%', '46%'])
        self.assertTrue(all(l['placement'] == 'inside' for l in result['pie']))
        # The frame is the text's width plus a margin, not a fixed 64px box.
        for label in result['pie']:
            self.assertLess(label['frame']['width'], 64)
            self.assertGreater(label['frame']['width'], label['frame']['height'])
        outside = [l for l in result['donut'] if l['placement'] == 'outside']
        self.assertEqual([l['text'] for l in outside], ['6%', '4%'])
        c = result['circle']
        for label in outside:
            f = label['frame']
            cx, cy = f['x'] + f['width'] / 2, f['y'] + f['height'] / 2
            self.assertGreater(((cx - c['x'] - c['width'] / 2) ** 2 + (cy - c['y'] - c['height'] / 2) ** 2) ** 0.5, c['width'] / 2)
        # An outside share is placed against the drawn circle, so it stays drawn.
        self.assertIsNone(result['native'])

    @needs_pptx
    def test_native_pie_prints_the_scenes_shares_inside(self):
        scene = run_node("""
import {compileDeck} from './skills/professional-slides/runtime/core.mjs';
import {REGISTRY} from './skills/professional-slides/runtime/registry.mjs';
const slides=[{id:'p',composition:{nodeType:'component',id:'pie',component:'chart.pie',props:{labels:['International','Domestic'],values:[75.8,65.1]},frame:{x:72,y:200,width:673,height:370}}}];
console.log(JSON.stringify(compileDeck({id:'pie',slides},REGISTRY)));
""")
        self.assertIsNotNone(scene['slides'][0]['componentInstances'][0].get('nativeChart'))
        with tempfile.TemporaryDirectory() as tmp:
            file = Path(tmp) / 'pie.pptx'
            Emitter(scene).run(file)
            chart = next(s.chart for s in Presentation(file).slides[0].shapes if s.has_chart)
            xml = chart._chartSpace.xml
        # Percentages, formatted as the scene prints them, centred on the slice
        # - never the raw value outside the rim.
        self.assertIn('<c:showPercent val="1"/>', xml)
        self.assertNotIn('<c:showVal val="1"/>', xml)
        self.assertIn('formatCode="0%"', xml)
        self.assertNotIn('<c:dLblPos val="outEnd"/>', xml)


class ForecastKeyTests(unittest.TestCase):
    def test_forecast_columns_are_keyed_and_divided(self):
        result = run_node(f"""
import {{REGISTRY}} from './skills/professional-slides/runtime/registry.mjs';
const frame={{x:72,y:162,width:1136,height:431}};
const render=extra=>REGISTRY.get('chart.column').render({{id:'c',frame,props:{{...{QATAR},legend:false,...extra}}}}).nodes;
const summary=nodes=>({{legend:nodes.filter(n=>n.role==='legend-label').map(n=>n.text),
  swatches:nodes.filter(n=>n.role==='legend-swatch').map(n=>n.style.fill.tokenId),
  divider:nodes.find(n=>n.role==='chart-forecast-divider')?.data,
  marks:nodes.filter(n=>n.role==='chart-mark').map(n=>n.frame)}});
console.log(JSON.stringify({{plain:summary(render({{}})),named:summary(render({{forecastLabel:'Required path'}})),off:summary(render({{forecastKey:false}}))}}));
""")
        plain = result['plain']
        # The composer's `legend: false` (no series list) does not remove the key.
        self.assertEqual(plain['legend'], ['Actual', 'Forecast'])
        self.assertEqual(plain['swatches'][1], 'color.chartSeries6')
        self.assertEqual(result['named']['legend'], ['Actual', 'Required path'])
        self.assertEqual(result['off']['legend'], [])
        # The dashed boundary sits between FY26 and FY27, clear of both columns.
        divider, marks = plain['divider'], plain['marks']
        self.assertEqual(divider['x1'], divider['x2'])
        self.assertGreater(divider['x1'], marks[2]['x'] + marks[2]['width'])
        self.assertLess(divider['x1'], marks[3]['x'])
        # Dropping the key keeps the non-colour cue.
        self.assertIsNotNone(result['off']['divider'])


class ReferenceLabelFallbackTests(unittest.TestCase):
    def test_a_crowded_reference_label_moves_outside_instead_of_throwing(self):
        result = run_node(f"""
import {{REGISTRY}} from './skills/professional-slides/runtime/registry.mjs';
const frame={{x:72,y:162,width:1136,height:431}};
const props={{...{QATAR},referenceLines:[{{value:53.2,label:'Emirates FY26: 53.2m'}}],
  annotations:[{{category:'FY26',text:'The latest year fell about 3%, well below the required path'}},{{category:'FY28',text:'Illustrative path: about 8.4% a year'}}]}};
const nodes=REGISTRY.get('chart.column').render({{id:'c',frame,props}}).nodes;
const label=nodes.find(n=>n.role==='chart-reference-label').frame;
const right=Math.max(...nodes.filter(n=>n.role==='chart-mark').map(n=>n.frame.x+n.frame.width));
const line=nodes.find(n=>n.role==='chart-reference-line').data;
let explicit=null;
try {{ REGISTRY.get('chart.column').render({{id:'c',frame,props:{{...props,referenceLines:[{{...props.referenceLines[0],placement:'inside'}}]}}}}); }} catch (error) {{ explicit=error.message; }}
console.log(JSON.stringify({{label,right,lineEnd:line.x2,explicit}}));
""")
        # Outside the plot, in the right gutter, not over the marks.
        self.assertGreater(result['label']['x'], result['lineEnd'])
        self.assertGreater(result['label']['x'], result['right'])
        # An author who asked for inside keeps the error, now naming the way out.
        self.assertIn('No collision-free reference-line label position', result['explicit'])
        self.assertIn('outside-end', result['explicit'])


class TableLogoTests(unittest.TestCase):
    def test_logos_share_an_area_and_fractional_widths_are_shares(self):
        result = run_node("""
import {renderTable} from './skills/professional-slides/runtime/tables.mjs';
import {TABLE_VARIANTS} from './skills/professional-slides/runtime/table-fixtures.mjs';
const base=TABLE_VARIANTS['category-logo-comparison'].props;
const frame={x:60,y:60,width:1160,height:580};
// The composer weights text columns by measured width (84-444); the author's
// logo column says 0.3 - a share of the table, not a weight of 0.3 among them.
const props={...base,columns:base.columns.map((c,i)=>({...c,width:i===1?0.3:[120,0,300][i]}))};
const nodes=renderTable({id:'t',frame,props}).nodes;
const logos=nodes.filter(n=>n.role==='table-logo').map(n=>({frame:n.frame,cellHeight:n.data.cellHeight,aspect:n.data.width/n.data.height}));
const cells=nodes.filter(n=>n.data?.column===1&&n.role==='table-rule').map(n=>n.frame.width);
let narrow=null;
try { renderTable({id:'n',frame,props:{...base,columns:base.columns.map((c,i)=>i===1?{...c,label:'ID'}:c),columnWidths:[.4,.05,.55]}}); } catch (error) { narrow=error.message; }
console.log(JSON.stringify({logos,cells,narrow}));
""")
        logos = result['logos']
        self.assertEqual(len(logos), 2)
        areas = [l['frame']['width'] * l['frame']['height'] for l in logos]
        # Two lines of cell, not one: the tallest mark stands at least 30px.
        self.assertGreaterEqual(min(l['cellHeight'] for l in logos), 40)
        self.assertGreaterEqual(max(l['frame']['height'] for l in logos), 30)
        self.assertLess(abs(areas[0] - areas[1]) / max(areas), 0.1)
        for logo in logos:
            self.assertAlmostEqual(logo['frame']['width'] / logo['frame']['height'], logo['aspect'], places=3)
        # 0.3 of a 1160px table, not a one-pixel column.
        self.assertTrue(result['cells'])
        self.assertAlmostEqual(result['cells'][0], 0.3 * 1160, delta=16)
        self.assertRegex(result['narrow'], r'Logo column is too narrow')


if __name__ == '__main__':
    unittest.main()
