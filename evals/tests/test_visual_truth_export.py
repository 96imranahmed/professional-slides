"""Read the saved editable file at the semantic boundaries seen in real decks."""
import sys
import tempfile
import unittest
from pathlib import Path

from node_probe import ROOT, run_node, requires_python_package

# python-pptx is optional at test time (see requirements.txt); the emitter and
# the readback both import it, so the class skips when it is absent rather than
# failing collection.
sys.path.insert(0, str(ROOT / 'skills/professional-slides/runtime/emit'))
try:
    from pptx import Presentation
    from emit_pptx import Emitter
    from readback_pptx import readback
except ImportError:  # pragma: no cover - exercised only without python-pptx
    Presentation = Emitter = readback = None


@requires_python_package('pptx')
class VisualTruthExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.scene = run_node(r"""
import {toDeckPlan} from './skills/professional-slides/runtime/compose.mjs';
import {planDeck} from './skills/professional-slides/runtime/planner.mjs';
const slides=[
 {id:'signed',title:'A loss and a gain must retain their direction and common scale',arrange:'row',exhibits:[[8,-12],[39,8]].map((values,i)=>({type:'chart.bar',heading:i?'Alternative B':'Alternative A',unit:'£k',categories:['First','Second'],series:[{name:'Residual',values}],dataTable:false}))},
 {id:'dates',title:'A 20-day interval occupies twice the width of a 10-day interval',exhibit:{type:'timeline',variant:'dated-lanes',axis:{kind:'date',precision:'day',scale:'elapsed-days',domain:['2026-01-01','2026-02-01']},lanes:[{id:'board',label:'Board'}],events:[{id:'a',laneId:'board',date:'2026-01-01',label:'First'},{id:'b',laneId:'board',date:'2026-01-11',label:'Second'},{id:'c',laneId:'board',date:'2026-01-31',label:'Third'}]}},
 {id:'nested',title:'Capacity and event classes retain their distinct visual roles',arrange:'row',exhibits:[
  {type:'chart.bar',heading:'Available capacity',unit:'seats',categories:['A','B'],series:[{name:'Seats',values:[24,28]}],referenceLines:[{value:32,label:'Room limit'}]},
  {type:'table',columns:[{label:'Class',type:'category'},'Meaning'],rows:[['Assignment','Given access'],['Outcome','Completed service']]}]},
 {id:'native-axis',title:'Stored supply declines across equal daily intervals',exhibit:{type:'chart.line',heading:'Remaining storage',unit:'ML',categories:['Day 0','Day 1','Day 2'],series:[{name:'Storage',values:[42,36,30]}],showValueAxis:true,dataLabels:false,endLabels:false,yMin:0,yMax:50}},
 {id:'numeric-line',title:'The selected intermediate observation retains its true elapsed position',exhibit:{type:'chart.line',heading:'Observed output',unit:'units',xAxis:{unit:'days',min:0,max:3,ticks:[{value:0,label:'Day 0'},{value:1,label:'Day 1'},{value:3,label:'Day 3'}]},gapPolicy:'connect-observations',series:[{name:'Output',points:[{key:'a',x:0,y:2},{key:'b',x:1,y:4,label:true},{key:'c',x:3,y:6}]}],dataLabels:false,yMin:0,yMax:10}},
 ...[false,true].map(longChart=>({id:longChart?'chart-wrap':'table-wrap',title:'Peer heading rules align while every label remains visible',layout:'two-up',exhibits:[
  {type:'chart.column',heading:longChart?'Output across the full observation period for the two defined operating alternatives':'Output',unit:'units',categories:['A','B'],series:[{name:'Output',values:[2,4]}]},
  {type:'table',columns:longChart?['Choice','Cash','Condition']:['First\nchoice','Initial cash\nand water','Timing condition'],rows:[['Full','$31m','Starts now'],['Wait','$0m','Starts later']]}
 ]}))
];
console.log(JSON.stringify(planDeck(toDeckPlan({schema:'professional-slides.deck/v3',id:'truth',slides})).deck));
""")
        cls.tmp = tempfile.TemporaryDirectory()
        cls.path = Path(cls.tmp.name) / 'truth.pptx'
        Emitter(cls.scene).run(cls.path)
        cls.saved = Presentation(cls.path)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def shape(self, page, node):
        def walk(shapes):
            for shape in shapes:
                yield shape
                if hasattr(shape, 'shapes'):
                    yield from walk(shape.shapes)
        return next(s for s in walk(self.saved.slides[page].shapes) if s.name == 'ps:' + node['id'])

    def test_signed_marks_and_equal_values_keep_their_physical_meaning(self):
        nodes = self.scene['slides'][0]['nodes']
        marks = [n for n in nodes if n['role'] == 'chart-mark']
        left = [n for n in marks if n['data']['componentInstance'].endswith('exhibit-0')]
        right = [n for n in marks if n['data']['componentInstance'].endswith('exhibit-1')]
        positive, negative = [self.shape(0, n) for n in left]
        self.assertLess(negative.left, positive.left)
        self.assertAlmostEqual(negative.left + negative.width, positive.left, delta=2)
        self.assertAlmostEqual(positive.width, self.shape(0, right[1]).width, delta=2)

    def test_unequal_date_intervals_survive_export(self):
        nodes = self.scene['slides'][1]['nodes']
        shapes = [self.shape(1, next(n for n in nodes if n['id'].endswith('dates-exhibit:' + key))) for key in ['a', 'b', 'c']]
        a, b, c = [s.left for s in shapes]
        self.assertAlmostEqual((c-b)/(b-a), 2, places=4)

    def test_reference_and_category_surfaces_remain_visible_in_saved_file(self):
        nodes = self.scene['slides'][2]['nodes']
        reference = next(n for n in nodes if n['role'] == 'chart-reference-label')
        self.assertEqual(self.shape(2, reference).text, 'Room limit')
        for node in nodes:
            if node['role'] == 'table-cell' and node.get('data', {}).get('cellType') == 'category':
                self.assertIsNotNone(self.shape(2, node).fill.fore_color.rgb)
        category_text = [n for n in nodes if n['type'] == 'text' and n.get('data', {}).get('cellType') == 'category']
        self.assertEqual(len(category_text), 2)
        for node in category_text:
            run = self.shape(2, node).text_frame.paragraphs[0].runs[0]
            self.assertEqual(str(run.font.color.rgb), 'FFFFFF')

    def test_native_ticks_keep_the_composed_domain_and_step(self):
        slide = self.scene['slides'][3]
        instance = next(c for c in slide['componentInstances'] if c.get('nativeChart'))
        saved_chart = next(s.chart for s in self.saved.slides[3].shapes if getattr(s, 'has_chart', False))
        self.assertEqual(saved_chart.value_axis.minimum_scale, 0)
        self.assertEqual(saved_chart.value_axis.maximum_scale, 50)
        self.assertEqual(saved_chart.value_axis.major_unit, 10)
        self.assertEqual(instance['nativeChart']['yMajorUnit'], 10)
        altered = Presentation(self.path)
        chart = next(s.chart for s in altered.slides[3].shapes if getattr(s, 'has_chart', False))
        chart.value_axis.major_unit = 5
        bad = Path(self.tmp.name) / 'bad-axis.pptx'
        altered.save(bad)
        self.assertTrue(any(f['code'] == 'NATIVE_AXIS_DRIFT' for f in readback(self.scene, bad)['findings']))

    def test_numeric_x_line_automatically_preserves_spacing_and_selected_label(self):
        slide = self.scene['slides'][4]
        self.assertFalse(any(c.get('nativeChart') for c in slide['componentInstances']))
        self.assertFalse(any(getattr(s, 'has_chart', False) for s in self.saved.slides[4].shapes))
        markers = sorted((n for n in slide['nodes'] if n['role'] == 'chart-marker'), key=lambda n:n['data']['xValue'])
        a, b, c = [self.shape(4, n).left for n in markers]
        self.assertAlmostEqual((c-b)/(b-a), 2, places=4)
        labels = [n for n in slide['nodes'] if n['role'] == 'data-label']
        self.assertEqual([n['data']['pointKey'] for n in labels], ['b'])
        self.assertEqual(self.shape(4, labels[0]).text, '4')

    def test_mixed_peer_heading_rules_align_in_saved_file_for_either_longer_header(self):
        for page in [5, 6]:
            nodes = self.scene['slides'][page]['nodes']
            chart_rule = next(n for n in nodes if n['role'] == 'section-heading-rule')
            table_rule = next(n for n in nodes if ':header-rule:' in n['id'])
            rule_y = self.shape(page, chart_rule).top
            self.assertEqual(rule_y, self.shape(page, table_rule).top)
            headings = [n for n in nodes if n['role'] in ['section-heading', 'chart-unit', 'table-header-text']]
            for node in headings:
                shape = self.shape(page, node)
                self.assertLessEqual(shape.top + shape.height, rule_y)
            body = [n for n in nodes if n['role'] == 'table-cell-text']
            self.assertTrue(body)
            self.assertGreater(min(self.shape(page, n).top for n in body), rule_y)
