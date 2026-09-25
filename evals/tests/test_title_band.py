"""The title band: tracker, title, optional standfirst, the rule that closes
it, and the body below.

Three in four well-made analytical pages close the title band with a rule or
a band and a fifth carry a standfirst under the title; the generated decks had
neither, and nine two-line titles in ten ended on a one-to-three-word line.
These pin the rule the design system draws, the standfirst's place and count,
and the balanced break the render and the PPTX share.
"""

from __future__ import annotations

import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from node_probe import run_node, requires_python_package

ROOT = Path(__file__).resolve().parents[2]
EMIT = ROOT / "skills" / "professional-slides" / "runtime" / "emit"
GATES = ROOT / "skills" / "professional-slides" / "runtime" / "gates"
sys.path.insert(0, str(GATES))
import page_gates  # noqa: E402

LONG = "Every segment recovered, and only new-age tech is still below its 2021 level across the region today"

CHROME = r"""
import assert from 'node:assert/strict';
import {compileDeck, component, absolute} from './skills/professional-slides/runtime/core.mjs';
import {createRegistry} from './skills/professional-slides/runtime/registry.mjs';
import {applyDesign} from './skills/professional-slides/runtime/design-systems.mjs';
const REGISTRY=createRegistry(), frame={x:0,y:0,width:1280,height:720};
const page=(deck,props)=>{
  const spec=deck.design?applyDesign({id:'t',...deck}):deck;
  const out=compileDeck({palette:spec.palette,slides:[{id:'s1',chrome:{source:'Source: fixture',...props},composition:absolute({id:'empty',children:[]})}]},REGISTRY).slides[0];
  const role=(r)=>out.nodes.find(n=>n.role===r);
  return {out,role,title:role('action-title'),rule:role('title-rule'),sub:role('action-subtitle')};
};
"""


class TitleRuleTests(unittest.TestCase):
    def test_each_design_system_draws_its_own_rule(self):
        result = run_node(CHROME + r"""
const consulting=page({design:'consulting'},{title:'Short title'}).rule;
const editorial=page({design:'editorial'},{title:'Short title'}).rule;
const evergreen=page({palette:'evergreen'},{title:'Short title'}).rule;
const journal=page({design:'journal'},{title:'Short title'});
const keynote=page({design:'keynote'},{title:'Short title'});
console.log(JSON.stringify({
  consulting:[consulting.data.x1,consulting.data.x2,consulting.style.lineWidth.value,consulting.style.stroke.tokenId],
  editorial:[editorial.data.x1,editorial.data.x2],
  evergreen:[evergreen.data.x2-evergreen.data.x1,evergreen.style.lineWidth.value,evergreen.style.stroke.tokenId],
  journal:!!journal.rule, keynote:!!keynote.rule, tab:!!journal.role('title-tab'), block:!!keynote.role('title-band')}));
""")
        self.assertEqual(result["consulting"], [60, 1220, 1, "color.rule"])   # margin to margin, a grey hairline
        self.assertEqual(result["editorial"], [0, 1280])                      # edge to edge
        self.assertEqual(result["evergreen"], [64, 4, "color.accent"])        # a short accent bar
        # A tab or a block already closes the band: no second rule under it.
        self.assertFalse(result["journal"])
        self.assertFalse(result["keynote"])
        self.assertTrue(result["tab"] and result["block"])

    def test_the_title_sits_on_a_fixed_rule_and_the_body_starts_below_it(self):
        result = run_node(CHROME + r"""
const deck={design:'consulting'};
const one=page(deck,{title:'Short title'}), two=page(deck,{title:'""" + LONG + r"""'});
const tracked=page(deck,{title:'Short title',tracker:{trackerId:'map',items:['A','B','C'].map(id=>({id,label:`Section ${id}`})),selectedId:'B',construction:'compact-label',mode:'light'}});
const bottom=(p)=>p.title.frame.y+p.title.data.textLayout.height;
console.log(JSON.stringify({
  rules:[one.rule.frame.y,two.rule.frame.y], gaps:[one.rule.frame.y-bottom(one),two.rule.frame.y-bottom(two)],
  body:[one.out.contentFrame.y,two.out.contentFrame.y], lines:[one.title.data.textLayout.lines.length,two.title.data.textLayout.lines.length],
  tracked:{tracker:Math.max(...tracked.out.nodes.filter(n=>/^tracker/.test(n.role)&&n.frame).map(n=>n.frame.y+n.frame.height)),title:tracked.title.frame.y,rule:tracked.rule.frame.y,body:tracked.out.contentFrame.y}}));
""")
        self.assertEqual(result["lines"], [1, 2])
        # One rule height and one body top whether the title takes one line or two,
        # and the title set down on the rule rather than the rule left under air.
        self.assertEqual(result["rules"][0], result["rules"][1])
        self.assertEqual(result["gaps"], [12, 12])
        self.assertEqual(result["body"], [140, 140])
        self.assertEqual(result["body"][0] - result["rules"][0], 12)
        tracked = result["tracked"]
        self.assertLess(tracked["tracker"], tracked["title"])
        self.assertLess(tracked["rule"], tracked["body"])

    def test_the_rule_is_set_by_tokens_a_palette_may_override(self):
        result = run_node(CHROME + r"""
const custom=page({palette:{base:'midnight',colors:{'style.titleRuleLength':'full','style.titleRuleColor':'ink','line.titleRule':2,'layout.titleRuleGap':16}}},{title:'Short title'});
const errors=[];
for (const colors of [{'line.titleRule':'thick'},{'layout.titleRuleGap':-4},{'spacing.x':3}]) { try { page({palette:{base:'midnight',colors}},{title:'T'}); } catch (e) { errors.push(e.message); } }
for (const colors of [{'style.titleRuleLength':'half'},{'style.titleRuleColor':'green'}]) { try { page({palette:{base:'midnight',colors}},{title:'T'}); } catch (e) { errors.push(e.message); } }
const off=page({palette:{base:'midnight',colors:{'style.titleRule':'none'}}},{title:'Short title'});
console.log(JSON.stringify({x:[custom.rule.data.x1,custom.rule.data.x2],width:custom.rule.style.lineWidth.value,stroke:custom.rule.style.stroke.tokenId,
  gap:custom.out.contentFrame.y-custom.rule.frame.y,errors:errors.length,off:!!off.rule,offTitle:off.title.frame.y}));
""")
        self.assertEqual(result["x"], [0, 1280])
        self.assertEqual(result["width"], 2)
        self.assertEqual(result["stroke"], "color.ink")
        self.assertEqual(result["gap"], 16)
        self.assertEqual(result["errors"], 5)
        # With the rule off the page keeps its open band and its title at the top.
        self.assertFalse(result["off"])
        self.assertEqual(result["offTitle"], 44)

    def test_the_rule_is_not_counted_as_body_ink(self):
        slide = {"nodes": [{"type": "line", "role": "title-rule", "frame": {"x": 60, "y": 128, "width": 1160, "height": 0},
                            "style": {"lineWidth": {"value": 1}}, "data": {"x1": 60, "x2": 1220, "y1": 128, "y2": 128}}]}
        rows = [0] * 720
        rows[127] = rows[128] = 1160
        rows[300] = 500
        cleaned = page_gates.without_title_rule(rows, slide)
        self.assertEqual(cleaned[127] + cleaned[128], 0)
        self.assertEqual(cleaned[300], 500)
        self.assertEqual(rows[128], 1160)  # the void gates still read the rule


class StandfirstTests(unittest.TestCase):
    def test_the_standfirst_sits_between_title_and_rule_and_pushes_only_a_tall_band(self):
        result = run_node(CHROME + r"""
const deck={design:'consulting'}, sub='Announced deal value by segment, India, $B';
const one=page(deck,{title:'Short title',subtitle:sub}), two=page(deck,{title:'""" + LONG + r"""',subtitle:sub});
const check=(p)=>({under:p.sub.frame.y>=p.title.frame.y+p.title.data.textLayout.height, above:p.sub.frame.y+p.sub.frame.height<p.rule.frame.y,
  smaller:p.sub.style.fontSize.value<p.title.style.fontSize.value, secondary:p.sub.style.color.tokenId, body:p.out.contentFrame.y, rule:p.rule.frame.y});
console.log(JSON.stringify({one:check(one),two:check(two)}));
""")
        for key in ("one", "two"):
            self.assertTrue(result[key]["under"] and result[key]["above"] and result[key]["smaller"], key)
            self.assertEqual(result[key]["secondary"], "color.textSecondary")
        # A one-line title and its standfirst fit the band; two lines and a
        # standfirst push the rule and the body down together.
        self.assertEqual(result["one"]["body"], 140)
        self.assertGreater(result["two"]["body"], 140)
        self.assertEqual(result["two"]["body"] - result["two"]["rule"], 12)

    def test_a_page_type_takes_a_one_line_standfirst_that_adds_to_its_title(self):
        result = run_node(r"""
import {compilePage, pageSchema, describeTypes, SUBTITLE_WORDS} from './skills/professional-slides/runtime/page-types.mjs';
import {planRole} from './skills/professional-slides/runtime/derive-content.mjs';
import fs from 'node:fs';
const base=JSON.parse(fs.readFileSync('./skills/professional-slides/examples/page-types.pages.json','utf8')).pages.find(p=>p.id==='p03');
const error=(page)=>{ try { compilePage(page); return null; } catch (e) { return e.message; } };
const ok=compilePage({...base,subtitle:'Passenger journeys a year, all lines, FY19 to FY26'});
const schema=pageSchema().properties.pages.items.oneOf.filter(t=>t.properties.type);
console.log(JSON.stringify({kept:ok.subtitle, limit:SUBTITLE_WORDS,
  long:error({...base,subtitle:'Passenger journeys a year on every one of the five lines in the network, counted at the gates, FY19 to FY26'}),
  restates:error({...base,subtitle:base.title.split(' ').slice(0,8).join(' ')}),
  empty:error({...base,subtitle:' '}),
  statement:error({id:'s',type:'statement',form:'statement',commentary:'none',takeaway:false,why:'The claim stands alone as the message page',text:'The peak is full',title:'The peak is full',subtitle:'All lines'}),
  schema:schema.every(t=>t.properties.subtitle && !t.required.includes('subtitle')),
  described:/`subtitle`/.test(describeTypes()),
  roles:[planRole('action-subtitle'),planRole('takeaway-standfirst'),planRole('action-title')]}));
""")
        self.assertEqual(result["kept"], "Passenger journeys a year, all lines, FY19 to FY26")
        self.assertEqual(result["limit"], 16)
        self.assertIn("words", result["long"])
        self.assertIn("repeats the title", result["restates"])
        self.assertIn("one line", result["empty"])
        self.assertIn("no title band", result["statement"])
        self.assertTrue(result["schema"])
        self.assertTrue(result["described"])
        # The standfirst counts with the title, not toward the body floor; the
        # journal's takeaway set as a standfirst is still the page's takeaway.
        self.assertEqual(result["roles"], ["title", "body", "title"])


class BalancedTitleTests(unittest.TestCase):
    TITLES = [
        LONG,
        "Emirates carried more passengers than Qatar Airways and Etihad combined in the year to March",
        "Premium yields held at 2019 levels while economy yields fell from 66% to 54%",
        "The Dubai hub connects more city pairs than any rival, and the gap is widening on long-haul banks",
    ]

    def test_a_wrapped_title_is_balanced_and_never_ends_on_a_short_line(self):
        titles = json.dumps(self.TITLES)
        result = run_node(CHROME + r"""
import {measureText, balancedWrap} from './skills/professional-slides/runtime/text-layout.mjs';
const rows=[];
for (const text of """ + titles + r""") {
  const p=page({design:'consulting'},{title:text}), layout=p.title.data.textLayout;
  const greedy=measureText(text,1160,{fontFamily:'Georgia',fontSize:24,bold:true,wrapWidthRatio:0.98});
  rows.push({lines:layout.lines,greedy:greedy.lines.length,width:p.title.frame.width,balanced:!!p.title.data.balanced});
}
const kept=page({design:'consulting'},{title:'A title the author broke\nhere on purpose'}).title;
const wrap=balancedWrap((w)=>measureText('one two three four five six seven eight nine ten',w,{fontSize:12}),200);
const single=balancedWrap((w)=>measureText('short',w,{fontSize:12}),200);
console.log(JSON.stringify({rows,kept:[kept.data.textLayout.lines,!!kept.data.balanced],wrap:[wrap.width<200,wrap.layout.lines.length,measureText('one two three four five six seven eight nine ten',200,{fontSize:12}).lines.length],single:single.width}));
""")
        for row in result["rows"]:
            self.assertEqual(len(row["lines"]), row["greedy"], row)          # the same number of lines
            self.assertGreaterEqual(len(row["lines"][-1].split()), 4, row)  # no stranded last line
            self.assertTrue(row["balanced"], row)
            self.assertLess(row["width"], 1160, row)                         # the box takes the balanced width
        self.assertEqual(result["kept"], [["A title the author broke", "here on purpose"], False])
        self.assertEqual(result["wrap"][1], result["wrap"][2])
        self.assertTrue(result["wrap"][0])
        self.assertEqual(result["single"], 200)

    @requires_python_package("pptx")
    def test_the_pptx_breaks_a_balanced_title_where_the_render_does(self):
        scene = run_node(CHROME + r"""
const p=page({design:'consulting'},{title:'""" + LONG + r"""'});
const deck=compileDeck({palette:applyDesign({id:'t',design:'consulting'}).palette,slides:[{id:'s1',frame,composition:component({id:'chrome',component:'slide-chrome',frame,props:{title:'""" + LONG + r"""',source:'Source: fixture'}})}]},REGISTRY);
console.log(JSON.stringify(deck));
""")
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as tmp:
            scene_path, pptx_path = Path(tmp) / "scene.json", Path(tmp) / "deck.pptx"
            scene_path.write_text(json.dumps(scene))
            emit = subprocess.run([sys.executable, str(EMIT / "emit_pptx.py"), str(scene_path), str(pptx_path)], capture_output=True, text=True, cwd=str(ROOT))
            self.assertEqual(emit.returncode, 0, emit.stderr)
            read = subprocess.run([sys.executable, str(EMIT / "readback_pptx.py"), str(scene_path), str(pptx_path)], capture_output=True, text=True, cwd=str(ROOT))
            self.assertEqual(read.returncode, 0, read.stdout + read.stderr)
            self.assertTrue(json.loads(read.stdout)["accepted"], read.stdout)
            title = Presentation(str(pptx_path)).slides[0].shapes.title
            paragraphs = title.text_frame.paragraphs
            lines = next(n for n in scene["slides"][0]["nodes"] if n.get("role") == "action-title")["data"]["textLayout"]["lines"]
            # One paragraph for Outline view, broken with a soft return exactly
            # where the render broke it.
            self.assertEqual(len(paragraphs), 1)
            self.assertEqual(paragraphs[0].text.split("\x0b"), lines)


@requires_python_package("numpy")
class CensusRuleTests(unittest.TestCase):
    """The census finds the rule under a title with a tracker over it."""

    @classmethod
    def setUpClass(cls):
        spec = importlib.util.spec_from_file_location("reference_census", ROOT / "evals" / "scripts" / "reference_census.py")
        cls.census = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(cls.census)

    def page(self, blocks, rule_y):
        import numpy as np
        image = np.full((675, 1200), 255, dtype=np.uint8)
        for top, bottom, x1, x2 in blocks:
            image[top:bottom, x1:x2:3] = 20  # text: dense but broken rows
        image[rule_y, 56:1144] = 120
        return image

    def test_a_tracker_line_over_a_title_set_on_its_rule(self):
        tracker, title = (33, 39, 56, 300), (93, 122, 56, 900)
        self.assertTrue(self.census.title_rule(self.page([tracker, title], 131)))
        # An exhibit heading between the title and the rule is still not a title rule.
        heading = (150, 160, 56, 400)
        self.assertFalse(self.census.title_rule(self.page([tracker, title, heading], 172)))
        # Nor is a rule standing well clear of the title.
        self.assertFalse(self.census.title_rule(self.page([tracker, title], 190)))


if __name__ == "__main__":
    unittest.main()
