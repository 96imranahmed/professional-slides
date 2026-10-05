#!/usr/bin/env python3
"""Turn an existing deck into an inventory and a starter pages file.

    python3 runtime/import-deck.py deck.pptx out-dir/ [--id <id>] [--carry] [--force]
    python3 runtime/import-deck.py deck.pptx --inventory-only <file.json>

The existing-deck workflow (`workflow: "existing_deck_revision"`) starts here.
The .pptx is read, never written: its bytes are hashed and parsed in memory.
Four things land in out-dir:

    <id>.inventory.json  every slide by stable id (s01, s02, ...): layout,
                         hidden flag, title, subtitle, body paragraphs with
                         their indent levels in reading order (where the slide
                         shows them, through turned and mirrored groups), table
                         cells, chart series and values, pictures with the size
                         and turn they show at, and speaker notes; a paragraph
                         that is page furniture carries its `role` - a line
                         repeated on most slides ("furniture": a footer, a
                         tracker), the slide's own number ("page-number"), a
                         source or a note line ("source", "note") - and
                         `unheld` counts what the slide draws that the
                         inventory does not hold (shapes drawn without text, a
                         diagram, an embedded object, a film)
    <id>.pages.json      the starter: a deck/v3 head that names the inventory
                         and the rules version the revision is made under, and
                         a page per slide carrying its old copy as `draft`
                         (its furniture left out) and, when the slide was
                         hidden, `hidden: true`
    <id>.source.pptx     the deck itself, byte for byte: what a carried slide
                         is copied from when the revision is built
    assets/<id>/         the pictures, as s01-1.png, s01-2.jpg, ... (one folder
                         per imported deck, so two decks never share a name)

Two starters, for the two things a revision can be:

    --carry   a point change. Every slide is a page marked `carry: true`: the
              build copies it from the source deck unchanged, and nothing of
              it is recomposed, judged or reviewed. The author changes only
              the pages the request names - a text edit on a carried slide
              (`replace`), or a page given a `type` in place of `carry` and
              `draft`, which the runtime then composes - and adds or deletes
              pages. The starter compiles as written.
    (plain)   a rebuild. The deck's footer when one line repeats on its
              slides, a cover when slide 1 is a shown title slide, and every
              other slide a page to map: the author gives each, by its stable
              id, a page type (`type`, `form`, `commentary`, `title`,
              `exhibit`, ...) before author-deck.mjs compiles it, or marks it
              `carry: true` to keep the slide as it is. This starter does not
              compile as written, by design.

Either way the inventory keeps the full data (table cells, chart values) to
copy into exhibits, a mapped page keeps its `hidden` (or sets it false to show
the slide again), and the import says what the author has to decide before
anything is compiled: which slides draw what the inventory does not hold, and
whether a composed page can be set into this deck at all (its slide size).

A starter holds the author's revisions, so it is never overwritten: when
<id>.pages.json exists the import refuses (exit 2) unless --force is given.

`--inventory-only` reads a deck as the inventory reads it and writes that one
file - no starter, no pictures, no copy.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import re
import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.enum.shapes import PP_PLACEHOLDER
    from pptx.shapes.group import GroupShape
    from pptx.shapes.picture import Picture
except ImportError:  # pragma: no cover
    sys.exit("python-pptx is required: pip install python-pptx")

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
}
EMU_PER_PX = 9525  # 914400 EMU an inch at 96 px an inch
INVENTORY_SCHEMA = "professional-slides.inventory/v1"
DECK_SCHEMA = "professional-slides.deck/v3"
TITLES = {"TITLE", "CENTER_TITLE", "VERTICAL_TITLE"}
CHROME = {"DATE", "FOOTER", "SLIDE_NUMBER", "HEADER"}  # page furniture, not copy
EXTENSIONS = {
    "image/png": "png", "image/jpeg": "jpg", "image/jpg": "jpg", "image/gif": "gif",
    "image/bmp": "bmp", "image/tiff": "tiff", "image/svg+xml": "svg",
    "image/x-emf": "emf", "image/x-wmf": "wmf", "image/webp": "webp",
}
# A placement carries a shape's own coordinates onto the slide, as the affine
# (a, b, c, d, e, f): slide x = a * x + c * y + e, slide y = b * x + d * y + f.
IDENTITY = (1.0, 0.0, 0.0, 1.0, 0.0, 0.0)


class Refusal(Exception):
    """An input the import will not act on; main() reports it and exits 2."""


def slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-") or "deck"


def deck_id(value: str) -> str:
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]*", value):
        raise argparse.ArgumentTypeError(f"'{value}' is not a usable id: letters, digits, '.', '_' and '-' only")
    return value


def px(emu: float) -> int:
    return round(emu / EMU_PER_PX)


def clean(text: str | None) -> str:
    """Text as the author will revise it: soft line breaks read as spaces."""
    return (text or "").replace("\x0b", " ").strip()


def joined(text_frame) -> str | None:
    """A title or subtitle frame as one line: its paragraphs joined by spaces."""
    parts = [clean(p.text) for p in text_frame.paragraphs]
    return " ".join(p for p in parts if p) or None


def reason(error: Exception) -> str:
    return f"{type(error).__name__}: {error}" if str(error) else type(error).__name__


def placeholder_kind(shape) -> str | None:
    try:
        return shape.placeholder_format.type.name if shape.is_placeholder else None
    except Exception:
        return None


def then(outer: tuple, inner: tuple) -> tuple:
    """The placement that applies `inner`, then `outer`."""
    a1, b1, c1, d1, e1, f1 = outer
    a2, b2, c2, d2, e2, f2 = inner
    return (a1 * a2 + c1 * b2, b1 * a2 + d1 * b2, a1 * c2 + c1 * d2, b1 * c2 + d1 * d2,
            a1 * e2 + c1 * f2 + e1, b1 * e2 + d1 * f2 + f1)


def turned(xfrm, left: float, top: float, width: float, height: float) -> tuple:
    """A frame's own flips and rotation, both about its centre: flipH and flipV
    mirror it, then `rot` (60,000ths of a degree, clockwise) turns it."""
    if xfrm is None:
        return IDENTITY
    theta = math.radians(int(xfrm.get("rot", 0)) / 60000)
    fx = -1.0 if xfrm.get("flipH") in ("1", "true") else 1.0
    fy = -1.0 if xfrm.get("flipV") in ("1", "true") else 1.0
    a, b, c, d = math.cos(theta) * fx, math.sin(theta) * fx, -math.sin(theta) * fy, math.cos(theta) * fy
    cx, cy = left + width / 2, top + height / 2
    return (a, b, c, d, cx - a * cx - c * cy, cy - b * cx - d * cy)


def child_placement(placement: tuple, group) -> tuple:
    """The placement of a group's members: the group's child frame (chOff,
    chExt) stretched onto its own frame (off, ext), then the group's flips and
    rotation about that frame's centre, then wherever the group itself sits."""
    xfrm = group._element.find("p:grpSpPr/a:xfrm", NS)
    if xfrm is None:
        return placement
    def pair(tag, first, second):
        el = xfrm.find(f"a:{tag}", NS)
        return None if el is None else (int(el.get(first, 0)), int(el.get(second, 0)))
    off, ext, ch_off, ch_ext = pair("off", "x", "y"), pair("ext", "cx", "cy"), pair("chOff", "x", "y"), pair("chExt", "cx", "cy")
    if None in (off, ext, ch_off, ch_ext):
        return placement
    sx = ext[0] / ch_ext[0] if ch_ext[0] else 1.0
    sy = ext[1] / ch_ext[1] if ch_ext[1] else 1.0
    stretch = (sx, 0.0, 0.0, sy, off[0] - ch_off[0] * sx, off[1] - ch_off[1] * sy)
    return then(placement, then(turned(xfrm, off[0], off[1], ext[0], ext[1]), stretch))


def displayed(shape, placement: tuple) -> tuple:
    """Where the slide shows a leaf shape: (frame, rotation, mirrored). The
    frame (left, top, width, height, slide EMU) is the axis-aligned box its
    turned corners cover; the rotation is how far its upright edge is turned
    clockwise, in degrees, net of every group's mirrors; mirrored is whether an
    odd number of flips shows it reversed, read as a left-right flip."""
    left, top, width, height = shape.left or 0, shape.top or 0, shape.width or 0, shape.height or 0
    element = shape._element
    own = element.find("p:spPr/a:xfrm", NS)
    m = then(placement, turned(own if own is not None else element.find("p:xfrm", NS), left, top, width, height))
    corners = [(m[0] * x + m[2] * y + m[4], m[1] * x + m[3] * y + m[5]) for x in (left, left + width) for y in (top, top + height)]
    xs, ys = [x for x, _ in corners], [y for _, y in corners]
    # A left-right flip leaves the upright edge where it was, so its direction
    # is the turn whether or not the shape is mirrored.
    rotation = round(math.degrees(math.atan2(-m[2], m[3])), 2) % 360
    return (min(xs), min(ys), max(xs) - min(xs), max(ys) - min(ys)), int(rotation) if rotation.is_integer() else rotation, m[0] * m[3] - m[1] * m[2] < 0


def walk(shapes, placement=IDENTITY):
    """Every leaf shape on the slide, group members included, with where the
    slide shows it (displayed)."""
    for shape in shapes:
        if isinstance(shape, GroupShape):
            yield from walk(shape.shapes, child_placement(placement, shape))
            continue
        yield (shape, *displayed(shape, placement))


def reading_order(slide) -> list:
    """The slide's leaf shapes top to bottom, then left to right, by where the
    slide shows them."""
    return sorted(walk(slide.shapes), key=lambda item: (item[1][1], item[1][0]))


def read_table(shape) -> dict:
    table = shape.table
    cells = [[clean(cell.text) for cell in row.cells] for row in table.rows]
    return {"rows": len(table.rows), "cols": len(table.columns), "cells": cells}


def number(value):
    if value is None:
        return None
    try:
        value = float(value)
    except (TypeError, ValueError):
        return None
    return value if math.isfinite(value) else None


def plot_tag(chart) -> str | None:
    """The first plot's XML tag (barChart, radarChart, ...) for a chart type
    python-pptx does not name."""
    area = chart._chartSpace.find("c:chart/c:plotArea", NS)
    for child in area if area is not None else []:
        tag = child.tag.split("}")[-1]
        if tag.endswith("Chart"):
            return tag
    return None


def category_labels(plot) -> list:
    """The leaf category labels by index; a category left empty in the
    workbook has no point and reads as None."""
    categories = plot.categories
    labels = [None] * len(categories)
    for position, category in enumerate(categories):
        pt = getattr(category, "_pt", None)
        index = int(pt.get("idx")) if pt is not None and pt.get("idx") is not None else position
        if 0 <= index < len(labels):
            labels[index] = str(category)
    return labels


def read_chart(shape) -> dict:
    """What can be read of a chart. A part that cannot be read is left empty
    and the failure is recorded under `unreadable`."""
    record = {"type": None, "title": None, "categories": [], "series": []}
    problems = []
    try:
        chart = shape.chart
    except Exception as error:
        return {**record, "unreadable": f"chart: {reason(error)}"}
    try:
        record["type"] = chart.chart_type.name
    except Exception:
        try:
            record["type"] = plot_tag(chart)
        except Exception as error:
            problems.append(f"type: {reason(error)}")
    try:
        if chart.has_title and chart.chart_title.has_text_frame:
            record["title"] = joined(chart.chart_title.text_frame)
    except Exception as error:
        problems.append(f"title: {reason(error)}")
    try:
        if len(chart.plots):
            record["categories"] = category_labels(chart.plots[0])
    except Exception as error:
        problems.append(f"categories: {reason(error)}")
    try:
        count = len(chart.series)
    except Exception as error:
        problems.append(f"series: {reason(error)}")
        count = 0
    for i in range(count):
        entry = {"name": None, "values": []}
        try:
            series = chart.series[i]
            entry["name"] = clean(series.name) or None
            entry["values"] = [number(v) for v in series.values]
        except Exception as error:
            problems.append(f"series {i + 1}: {reason(error)}")
        record["series"].append(entry)
    if problems:
        record["unreadable"] = "; ".join(problems)
    return record


def read_picture(shape, frame, rotation, mirrored: bool, name: str, assets: Path | None) -> dict:
    """Write the picture to assets/ and record how the slide shows it: the box
    it covers, its turn and, when it shows reversed, that it is mirrored. The
    file is the image as stored, unturned."""
    alt = None
    try:
        props = shape._element.find(".//p:cNvPr", NS)
        alt = clean(props.get("descr")) or clean(props.get("name")) or None
    except Exception:
        pass
    record = {"file": None, "width": px(frame[2]), "height": px(frame[3]), "rotation": rotation, **({"mirrored": True} if mirrored else {}), "alt": alt}
    if assets is None:  # an inventory alone: the picture is recorded, not written
        return record
    try:
        image = shape.image
        ext = EXTENSIONS.get(image.content_type, image.ext)
        path = assets / f"{name}.{ext}"
        assets.mkdir(parents=True, exist_ok=True)
        path.write_bytes(image.blob)
        record["file"] = path.relative_to(assets.parent.parent).as_posix()
    except Exception as error:
        record["unreadable"] = reason(error)
    return record


def read_notes(slide) -> str | None:
    try:
        if not slide.has_notes_slide:
            return None
        frame = slide.notes_slide.notes_text_frame
        return None if frame is None else clean(frame.text) or None
    except Exception:
        return None


def read_slide(slide, index: int, sid: str, assets: Path | None) -> dict:
    record = {
        "index": index, "id": sid, "layout": None, "hidden": slide._element.get("show") in ("0", "false"),
        "title": None, "subtitle": None, "paragraphs": [], "tables": [], "charts": [], "pictures": [], "notes": read_notes(slide),
    }
    try:
        record["layout"] = slide.slide_layout.name
    except Exception:
        pass
    unheld: dict = {}
    for shape, frame, rotation, mirrored in reading_order(slide):
        if isinstance(shape, Picture):
            record["pictures"].append(read_picture(shape, frame, rotation, mirrored, f"{sid}-{len(record['pictures']) + 1}", assets))
            if shape._element.find(".//a:videoFile", NS) is not None or shape._element.find(".//a:audioFile", NS) is not None:
                unheld["media"] = unheld.get("media", 0) + 1
        elif getattr(shape, "has_chart", False):
            record["charts"].append(read_chart(shape))
        elif getattr(shape, "has_table", False):
            record["tables"].append(read_table(shape))
        elif getattr(shape, "has_text_frame", False):
            kind = placeholder_kind(shape)
            if kind in CHROME:
                continue
            if kind in TITLES or kind == "SUBTITLE":
                key = "title" if kind in TITLES else "subtitle"
                text = joined(shape.text_frame)
                if record[key] is None:
                    record[key] = text
                    continue
            said = False
            for paragraph in shape.text_frame.paragraphs:
                text = clean(paragraph.text)
                if text:
                    said = True
                    record["paragraphs"].append({"text": text, "level": paragraph.level})
            # A shape with no words and no placeholder is part of a drawing: a box, a bar, an arrow.
            if not said and kind is None:
                unheld["drawn"] = unheld.get("drawn", 0) + 1
        else:
            what = unheld_kind(shape)
            unheld[what] = unheld.get(what, 0) + 1
    if unheld:
        record["unheld"] = unheld
    return record


def unheld_kind(shape) -> str:
    """What a shape the inventory holds nothing of is: a diagram (SmartArt),
    an embedded object, or a line or connector of a drawing."""
    uri = shape._element.find(".//a:graphicData", NS)
    kind = (uri.get("uri") or "") if uri is not None else ""
    return "diagrams" if "diagram" in kind else "objects" if kind else "drawn"


PAGE_NUMBER = re.compile(r"^(?:(?:page|slide|p\.?)\s*)?0*(\d{1,3})(?:\s*(?:/|of)\s*\d{1,3})?$", re.I)
SOURCE_LINE = re.compile(r"^(sources?|notes?)\s*:", re.I)
# A line is furniture when it stands on this share of the slides, and on three at least.
FURNITURE_SHARE = 0.5
FURNITURE_MIN = 3


def mark_furniture(slides: list) -> list:
    """Give each paragraph that is page furniture its `role`, and return the
    lines that repeat: a footer, a confidentiality line, a tracker's labels.

    Furniture is what a rebuild writes again from the deck's own settings, so
    it is neither a slide's copy to map nor a difference when the rebuilt deck
    is read against the imported one. A line repeated word for word on half
    the slides (three at least) is furniture; a bare number is a page number
    when that many slides each carry their own number at one offset from
    their place; a line that opens "Source:" or "Note:" is the slide's own,
    and is marked so the author maps it to the page's `source` or `note`."""
    shown = [s for s in slides if s["paragraphs"]]
    floor = max(FURNITURE_MIN, math.ceil(len(slides) * FURNITURE_SHARE))
    counts: dict = {}
    for slide in shown:
        for text in {p["text"] for p in slide["paragraphs"]}:
            counts[text] = counts.get(text, 0) + 1
    repeated = {text for text, n in counts.items() if n >= floor and re.search(r"[^\W\d_]", text)}
    # The offset between a slide's place and the number it prints: decks number from the cover, or from the page after it.
    offsets: dict = {}
    for slide in shown:
        for found in {int(m.group(1)) - slide["index"] for m in (PAGE_NUMBER.match(p["text"]) for p in slide["paragraphs"]) if m}:
            offsets[found] = offsets.get(found, 0) + 1
    offset = max(offsets, key=offsets.get) if offsets and max(offsets.values()) >= floor else None
    for slide in slides:
        numbered = False
        for paragraph in slide["paragraphs"]:
            match = PAGE_NUMBER.match(paragraph["text"])
            if paragraph["text"] in repeated:
                paragraph["role"] = "furniture"
            elif offset is not None and match and int(match.group(1)) - slide["index"] == offset and not numbered:
                paragraph["role"], numbered = "page-number", True
            elif SOURCE_LINE.match(paragraph["text"]):
                paragraph["role"] = "note" if paragraph["text"].lower().startswith("note") else "source"
    return sorted(repeated, key=lambda text: (-counts[text], text))


def looks_like_cover(slide: dict) -> bool:
    """Slide 1 is taken as the cover when it has a title and either sits on a
    title layout (not a content, comparison or title-only one) or carries a
    subtitle and nothing else."""
    if not slide["title"]:
        return False
    layout = (slide["layout"] or "").lower()
    title_layout = "title" in layout and not re.search(r"content|(?<![a-z])(?:and|only)(?![a-z])", layout)
    bare = slide["subtitle"] is not None and not (slide["paragraphs"] or slide["tables"] or slide["charts"])
    return title_layout or bare


def draft_of(slide: dict) -> dict:
    """The slide's old copy for the author to revise; empty parts are left out."""
    draft = {
        "subtitle": slide["subtitle"],
        "text": ["  " * p["level"] + p["text"] for p in slide["paragraphs"] if p.get("role") not in ("furniture", "page-number")],
        "tables": [{"rows": t["rows"], "cols": t["cols"]} for t in slide["tables"]],
        "charts": [{"type": c["type"], "categories": len(c["categories"]), "series": [s["name"] for s in c["series"]]} for c in slide["charts"]],
        "pictures": [p["file"] for p in slide["pictures"] if p["file"]],
        "notes": slide["notes"],
    }
    return {key: value for key, value in draft.items() if value not in (None, [], "")}


def rules_version() -> int | None:
    """The rules version this runtime enforces (weight.json), which the
    starter records: a rule introduced after the import is an advisory to this
    revision, as it is to any deck that predates it."""
    try:
        return json.loads((Path(__file__).resolve().parent / "weight.json").read_text(encoding="utf-8"))["rulesVersion"]
    except (OSError, KeyError, ValueError):
        return None


def starter(identifier: str, inventory_name: str, slides: list, furniture: list, carry: bool = False) -> dict:
    """The deck head and a page per slide. A hidden slide's page carries
    `hidden: true` beside its draft, not in it: mapping the page replaces the
    draft, and the flag must outlive it so the rebuilt slide stays out of the
    slide show.

    `carry` is the point-change starter: every slide - the title slide too -
    is a page marked `carry: true`, which the build copies from the source deck
    unchanged, so the deck is the user's own until the author changes a page.

    Otherwise the starter is a rebuild's: a cover when slide 1 is a shown
    title slide (a cover cannot be hidden, so a hidden title slide stays a
    page), the deck's footer when one line repeats on the slides (several are
    a footer among a tracker's labels, and the author says which), and every
    other slide a page to map."""
    version = rules_version()
    deck = {"schema": DECK_SCHEMA, "id": identifier, "workflow": "existing_deck_revision", "inventory": inventory_name,
            **({"rulesVersion": version} if version is not None else {}), "request": ""}
    page = lambda s: {"id": s["id"], "sourceSlide": s["index"], "title": s["title"] or "", **({"hidden": True} if s["hidden"] else {})}
    if carry:
        return {"deck": deck, "pages": [{**page(s), "carry": True, "draft": draft_of(s)} for s in slides]}
    if len(furniture) == 1:
        deck["footer"] = furniture[0]
    body = slides
    if slides and looks_like_cover(slides[0]) and not slides[0]["hidden"]:
        first = slides[0]
        deck["cover"] = {"title": first["title"], **({"subtitle": first["subtitle"]} if first["subtitle"] else {})}
        body = slides[1:]
    return {"deck": deck, "pages": [{**page(s), "draft": draft_of(s)} for s in body]}


# The canvas the runtime composes a page on (core.mjs SLIDE), in px at 96 an inch.
COMPOSED = {"width": 1280, "height": 720}
UNHELD_NAMES = {"drawn": ("shape drawn without text", "shapes drawn without text"), "diagrams": ("diagram", "diagrams"),
                "objects": ("embedded object", "embedded objects"), "media": ("film or sound", "films or sounds")}


def decisions(inventory: dict) -> list:
    """What the author has to decide before anything is compiled, in
    sentences: whether a page the runtime composes can be set into this deck,
    and which slides draw what the inventory does not hold - so a slide that
    cannot be recomposed from the inventory is known at the import, not three
    steps later."""
    said = []
    size = inventory["slideSize"]
    if (size["width"], size["height"]) != (COMPOSED["width"], COMPOSED["height"]):
        said.append(f"The deck's slides are {size['width']} x {size['height']} px; the runtime composes a page at {COMPOSED['width']} x {COMPOSED['height']}. "
                    "A carried slide and a text edit on one (`replace`) keep the deck's own size; a page given a `type` cannot be set beside carried slides of another size, "
                    "so a change that needs a composed page means rebuilding the deck: import without --carry and map every slide.")
    partial = [(s["id"], s["unheld"]) for s in inventory["slides"] if s.get("unheld")]
    unreadable = [s["id"] for s in inventory["slides"] if any(c.get("unreadable") for c in s["charts"]) or any(p.get("unreadable") for p in s["pictures"])]
    if partial:
        listed = "; ".join(f"{sid} ({', '.join(f'{n} {UNHELD_NAMES[kind][n != 1]}' for kind, n in sorted(found.items()))})" for sid, found in partial[:12])
        said.append(f"{len(partial)} slide{'' if len(partial) == 1 else 's'} draw{'s' if len(partial) == 1 else ''} what the inventory does not hold: {listed}{'; ...' if len(partial) > 12 else ''}. "
                    "Keep such a slide as it is (`carry: true`), or redraw it as a page type from its words and decide what replaces the drawing: the inventory cannot rebuild it.")
    if unreadable:
        said.append(f"A chart or a picture on {', '.join(unreadable)} could not be read (`unreadable` in the inventory): carry the slide, or supply the data when you map it.")
    return said


def open_deck(path: Path):
    if not path.is_file():
        raise Refusal(f"No such deck: {path}")
    try:
        data = path.read_bytes()
    except OSError as error:
        raise Refusal(f"Cannot read {path}: {error.strerror or error}")
    try:
        return Presentation(io.BytesIO(data)), hashlib.sha256(data).hexdigest()
    except Exception as error:
        raise Refusal(f"{path.name} is not a PowerPoint deck (.pptx): {reason(error)}")


def read_inventory(path: Path, identifier: str, assets: Path | None) -> dict:
    """The deck as the inventory records it; its pictures are written to `assets` when one is given."""
    prs, digest = open_deck(path)
    width = 3 if len(prs.slides) > 99 else 2
    slides = [read_slide(slide, i, f"s{i:0{width}d}", assets) for i, slide in enumerate(prs.slides, 1)]
    furniture = mark_furniture(slides)
    return {
        "schema": INVENTORY_SCHEMA, "id": identifier,
        "source": {"file": path.name, "sha256": digest},
        # The rules version the import stamps on the starter is recorded here too: a pages file that later records a lower one says why (revision.mjs).
        **({"rulesVersion": rules_version()} if rules_version() is not None else {}),
        "slideSize": {"width": px(prs.slide_width or 0), "height": px(prs.slide_height or 0)},
        **({"furniture": furniture} if furniture else {}),
        "slides": slides,
    }


def import_deck(path: Path, out: Path, identifier: str, force: bool = False, carry: bool = False) -> dict:
    inventory_path, pages_path, copy_path = out / f"{identifier}.inventory.json", out / f"{identifier}.pages.json", out / f"{identifier}.source.pptx"
    if not path.is_file():
        raise Refusal(f"No such deck: {path}")
    if pages_path.exists() and not force:
        raise Refusal(f"{pages_path} already exists and may hold revisions; pass --force to overwrite it with a fresh starter")
    open_deck(path)  # refused before anything is written
    out.mkdir(parents=True, exist_ok=True)
    inventory = read_inventory(path, identifier, out / "assets" / identifier)
    slides = inventory["slides"]
    # The deck itself is kept beside its inventory: a carried slide is copied from it when the revision is built, and
    # the user's file can then move or change without the revision losing the deck it was imported from.
    if path.resolve() != copy_path.resolve():
        copy_path.write_bytes(path.read_bytes())
    inventory["source"]["copy"] = copy_path.name
    inventory_path.write_text(json.dumps(inventory, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    pages_path.write_text(json.dumps(starter(identifier, inventory_path.name, slides, inventory.get("furniture", []), carry), indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    return {
        "inventory": str(inventory_path), "pages": str(pages_path), "source": str(copy_path), "slides": len(slides),
        "pictures": sum(len(s["pictures"]) for s in slides), "charts": sum(len(s["charts"]) for s in slides),
        "tables": sum(len(s["tables"]) for s in slides),
        "starter": "every slide is carried over as it is (`carry: true`); change only the pages the request names" if carry
                   else "every slide waits to be mapped to a page type, or marked `carry: true` to keep it as it is",
        "decide": decisions(inventory),
    }


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("deck", type=Path, help="the .pptx to read")
    parser.add_argument("out", type=Path, nargs="?", help="the folder for the inventory, the starter pages and assets/")
    parser.add_argument("--id", type=deck_id, help="the deck id (default: a slug of the file name)")
    parser.add_argument("--force", action="store_true", help="overwrite an existing starter pages file")
    parser.add_argument("--carry", action="store_true", help="the point-change starter: every slide carried over as it is")
    parser.add_argument("--inventory-only", type=Path, metavar="FILE", help="write the inventory to FILE and nothing else: no starter, no pictures")
    args = parser.parse_args(argv)
    if (args.out is None) == (args.inventory_only is None):
        parser.error("give the output folder, or --inventory-only <file.json>")
    if args.inventory_only and args.carry:
        parser.error("--carry chooses the starter, and --inventory-only writes none")
    try:
        if args.inventory_only:
            inventory = read_inventory(args.deck, args.id or slug(args.deck.stem), None)
            args.inventory_only.parent.mkdir(parents=True, exist_ok=True)
            args.inventory_only.write_text(json.dumps(inventory, ensure_ascii=False) + "\n", encoding="utf-8")
            print(json.dumps({"inventory": str(args.inventory_only), "slides": len(inventory["slides"])}))
            return 0
        summary = import_deck(args.deck, args.out, args.id or slug(args.deck.stem), args.force, args.carry)
    except Refusal as refusal:
        print(refusal, file=sys.stderr)
        return 2
    print(json.dumps(summary, indent=1))
    return 0


if __name__ == "__main__":
    sys.exit(main())
