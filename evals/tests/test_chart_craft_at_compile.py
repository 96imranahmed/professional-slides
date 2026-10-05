"""Chart craft the compile can hold and draw itself.

A rendered-deck review of a real run found three things the runtime had the
information to prevent. The build bars - tables treated, charts annotated -
appeared at the compile only as craft floors lower down, and the deck was
refused for them at delivery; thresholds and assumed constants were drawn as
data series, with the markers and weight of recorded values; and eight charts
marked nothing although their titles named the value to mark. These hold the
repair: every compile prints each delivery bar with the pages short of it and
refuses a miss delivery would refuse (a revision that carries slides on the
pages it composed, as delivery does); a series that is assumed or a reference is drawn as one, in the scene
and in the exported file; and a bare chart's finding names the exact mark its
title asks for, proven by composing the page with it.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import tempfile
import unittest
from pathlib import Path

from node_probe import NODE, ROOT, RUNTIME, requires_python_package, run_node

FIXTURES = ROOT / "evals" / "quality" / "fixtures" / "evidence"

BARS = '''
import { craftFindings } from './skills/professional-slides/runtime/gates/craft_gates.mjs';
import { barOutcome, BUILD_BARS } from './skills/professional-slides/runtime/build-bars.mjs';
import { standingLine } from './skills/professional-slides/runtime/gates/gate_classes.mjs';
import { RULES } from './skills/professional-slides/runtime/weight.mjs';
// One analytical page: an action title, twenty drawn marks and an exhibit of `kind`. A chart page marks its finding where
// `marked`, a table carries a treatment where `treated`, and `frame` leaves a picture frame empty.
const page = (id, kind, { marked = false, treated = false, frame = false } = {}) => ({ id, nodes: [{ role: 'action-title', type: 'text' }, ...Array(20).fill({ role: 'm', type: 'rect' }),
  ...(kind.startsWith('chart.') ? [{ role: 'chart-mark', type: 'rect', data: { highlighted: marked } }, { role: 'chart-mark', type: 'rect' }] : []),
  ...(kind === 'table' ? [{ role: 'table-cell', type: 'text' }, ...(treated ? [{ role: 'table-harvey', type: 'ellipse' }] : [])] : []),
  ...(frame ? [{ role: 'image-frame', type: 'rect', frame: { x: 0, y: 0, width: 400, height: 300 } }] : [])], componentInstances: [{ component: 'slide-chrome' }, { component: kind }] });
const OTHERS = ['cards', 'fact-grid', 'timeline', 'matrix', 'steps', 'tree', 'funnel', 'quote', 'map', 'gantt'];
// Twenty analytical pages behind a cover: eight chart pages (`marked` of them marking their finding), two tables, ten other exhibits.
const deck = ({ marked = 8, treated = 2, changed = {} } = {}) => ({ slides: [{ id: 'cover', nodes: [{ role: 'cover-title', type: 'text' }], componentInstances: [] },
  ...['chart.line', 'chart.bar', 'chart.column', 'chart.scatter', 'chart.line', 'chart.bar', 'chart.column', 'chart.area'].map((kind, i) => page(`c${i + 1}`, kind, { marked: i < marked, ...(changed[`c${i + 1}`] ?? {}) })),
  ...[0, 1].map((i) => page(`t${i + 1}`, 'table', { treated: i < treated })), ...OTHERS.map((kind, i) => page(`o${i + 1}`, kind, changed[`o${i + 1}`] ?? {}))] });
const spec = (more = {}) => ({ schema: 'professional-slides.deck/v3', id: 'probe', workflow: 'new_deck', slides: [], ...more });
const run = (scene, more = {}) => { const standings = [], found = craftFindings(spec(more), scene, { standings });
  return { bars: found.filter((f) => f.code.startsWith('BAR_')).map((f) => ({ code: f.code, severity: f.severity, rule: f.rule ?? null, pages: f.pages ?? null, repair: f.repair, waived: f.waived ?? null })),
    lines: Object.fromEntries(standings.filter((st) => st.code.startsWith('BAR_')).map((st) => [st.code, standingLine(st)])), blocks: Object.fromEntries(standings.filter((st) => st.code.startsWith('BAR_')).map((st) => [st.code, st.blocks])),
    floors: found.filter((f) => f.code.startsWith('CRAFT_') && f.severity === 'blocker').map((f) => f.code) }; };
'''


class BuildBarsAtCompileTests(unittest.TestCase):
    def test_every_compile_prints_each_delivery_bar_and_a_miss_is_refused_with_the_pages_and_the_waiver(self):
        result = run_node(BARS + '''
const reason = 'The charts are reference plots a reader looks values up in, and no single mark is the finding.';
console.log(JSON.stringify({ floors: [BUILD_BARS.tablesTreated.min, BUILD_BARS.chartsAnnotated.min], met: run(deck({ marked: 6 })), missed: run(deck({ marked: 4 })),
  waived: run(deck({ marked: 4 }), { waivers: [{ code: 'BAR_CHARTS_ANNOTATED', reason }] }), catalogue: run(deck({ marked: 4 }), { purpose: 'catalogue' }),
  // What delivery reads of the same scene: the compile refuses exactly the misses delivery would.
  spare: run(deck({ marked: 8 })).lines.BAR_CHARTS_ANNOTATED,
  delivery: { missed: barOutcome(deck({ marked: 4 })).unwaived.map((f) => f.code), met: barOutcome(deck({ marked: 6 })).unwaived.map((f) => f.code) } }));
''')
        # No threshold moved.
        self.assertEqual(result["floors"], [0.75, 0.65])
        met, missed = result["met"], result["missed"]
        self.assertEqual(met["bars"], [])
        self.assertEqual(sorted(met["lines"]), ["BAR_CHARTS_ANNOTATED", "BAR_DRAWINGS_PER_PAGE", "BAR_EXHIBIT_VARIETY", "BAR_TABLES_TREATED", "BAR_UNSOURCED_PICTURES"])
        # The standing is the delivery bar itself: the value, the bar and the pages to spare or short.
        self.assertEqual(met["lines"]["BAR_CHARTS_ANNOTATED"], "BAR_CHARTS_ANNOTATED: chart pages marking something on the plot 6 of 8, 75%; delivery floor 65%: at the delivery floor, 1 fewer blocks")
        self.assertEqual(result["spare"], "BAR_CHARTS_ANNOTATED: chart pages marking something on the plot 8 of 8, 100%; delivery floor 65%: 2 chart pages to spare")
        self.assertEqual(missed["lines"]["BAR_CHARTS_ANNOTATED"], "BAR_CHARTS_ANNOTATED: chart pages marking something on the plot 4 of 8, 50%; delivery floor 65%: short by 2 chart pages - BLOCKS")
        self.assertEqual(missed["lines"]["BAR_TABLES_TREATED"], "BAR_TABLES_TREATED: tables carrying a treatment 2 of 2, 100%; delivery floor 75%: at the delivery floor, 1 fewer blocks")
        self.assertEqual([[f["code"], f["severity"], f["rule"], f["pages"]] for f in missed["bars"]], [["BAR_CHARTS_ANNOTATED", "blocker", None, ["c5", "c6", "c7", "c8"]]])
        repair = missed["bars"][0]["repair"]
        self.assertIn("4 of 8 chart pages (50%) against the delivery floor of 65%: 2 more reach it (holding it down: c5, c6, c7, c8)", repair)
        self.assertIn("Delivery refuses the deck for this bar", repair)
        # A missed bar is waivable, and the finding says how, as delivery's refusal does.
        self.assertIn('record `waivers: [{ "code": "BAR_CHARTS_ANNOTATED", "reason": "<a sentence saying why>" }]` on `deck`', repair)
        self.assertEqual(result["delivery"], {"missed": ["BAR_CHARTS_ANNOTATED"], "met": []})
        waived = result["waived"]
        self.assertEqual([[f["code"], f["severity"]] for f in waived["bars"]], [["BAR_CHARTS_ANNOTATED", "advisory"]])
        self.assertIn("The deck waives this bar", waived["bars"][0]["repair"])
        self.assertIn("delivery shows the waiver to the reviewer, who must confirm it", waived["bars"][0]["repair"])
        self.assertFalse(waived["blocks"]["BAR_CHARTS_ANNOTATED"])
        self.assertIn("waived on the deck", waived["lines"]["BAR_CHARTS_ANNOTATED"])
        # A catalogue is held to the ceilings only, here as at delivery.
        self.assertEqual(result["catalogue"]["bars"], [])

    def test_a_rate_under_its_craft_floor_is_one_finding_and_the_floor_says_what_the_bar_asks(self):
        result = run_node(BARS + '''
// The craft floors are read from a deck's content pages, so the spec names twenty.
const found = craftFindings(spec({ slides: Array.from({ length: 20 }, (_, i) => ({ id: `p${i}`, title: `A finding on page ${i}` })) }), deck({ marked: 1 }), {});
console.log(JSON.stringify(found.filter((f) => /CHARTS/.test(f.code)).map((f) => [f.code, f.severity, f.repair])));
''')
        self.assertEqual([[code, severity] for code, severity, _ in result], [["CRAFT_CHARTS_BARE", "blocker"]])
        self.assertIn("The build bar BAR_CHARTS_ANNOTATED asks more than this floor: 1 of 8 chart pages (13%) against the delivery floor of 65%: 5 more reach it", result[0][2])

    def test_a_revision_is_held_to_the_bars_on_the_pages_it_composed_at_the_compile_as_at_delivery(self):
        result = run_node(BARS + '''
const revision = { workflow: RULES.revisionWorkflow };
const carried = (n) => Array.from({ length: n }, (_, i) => ({ id: `s${i + 1}`, sourceSlide: i + 1, title: '', after: null }));
// The scene of a revision that carries slides is the pages it composed (build-deck.mjs composes no carried slide).
const only = (scene, ids) => ({ slides: scene.slides.filter((slide) => ids.includes(slide.id)) });
const MOST = ['c1', 'c2', 'c3', 'c4', 'c5', 'c6', 'c7', 'c8', 't1', 't2', 'o1', 'o2'];
const scenes = {
  // One bare chart page composed beside nineteen carried slides: a floor is read from as many analytical pages as any deck's.
  point: [only(deck({ marked: 0 }), ['c1']), { ...revision, carried: carried(19) }],
  // The composed page leaves a picture frame empty: a cap is read from the first page.
  own: [only(deck({ changed: { o1: { frame: true } } }), ['o1']), { ...revision, carried: carried(19) }],
  // A revision that composed most of the deck answers for the rate on those pages.
  wide: [only(deck({ marked: 4 }), MOST), { ...revision, carried: carried(8) }],
  // The same pages under a waiver, and under the rules recorded before the compile held the bars.
  waived: [only(deck({ marked: 4 }), MOST), { ...revision, carried: carried(8), waivers: [{ code: 'BAR_CHARTS_ANNOTATED', reason: 'The charts are reference plots a reader looks values up in, and no single mark is the finding.' }] }],
  older: [only(deck({ marked: 4 }), MOST), { ...revision, carried: carried(8), rulesVersion: 5 }],
  // A rebuild carries no slide: every page is the runtime's, and the bars are held in full.
  rebuilt: [deck({ marked: 4 }), revision],
};
console.log(JSON.stringify(Object.fromEntries(Object.entries(scenes).map(([name, [scene, more]]) => [name, { ...run(scene, more),
  // What delivery reads: the same scene, the same waivers (deliver-deck.mjs).
  delivery: barOutcome(scene, more.waivers || [], {}).unwaived.map((f) => f.code) }]))));
''')
        blocked = lambda case: sorted(f["code"] for f in result[case]["bars"] if f["severity"] == "blocker")
        # The compile and delivery read one scope and refuse the same bars, case by case.
        for case in ("point", "own", "wide", "waived", "rebuilt"):
            self.assertEqual(blocked(case), sorted(result[case]["delivery"]), case)
        # One composed page is too few to read a floor from, so a point change is refused for no rate...
        self.assertEqual(result["point"]["bars"], [])
        # ...and every line says where the bar is held and where it is not.
        for line in result["point"]["lines"].values():
            self.assertIn("read over the 1 page this revision composed; not held on the 19 slides it carries from the source deck, which are the user's own", line)
        # A cap is held from the first composed page.
        self.assertEqual([[f["code"], f["severity"], f["pages"]] for f in result["own"]["bars"]], [["BAR_UNSOURCED_PICTURES", "blocker", ["o1"]]])
        # Twelve composed pages are a deck's worth: the rate is theirs, and the refusal is not waived as a measure of the imported deck.
        self.assertEqual([[f["code"], f["severity"], f["waived"]] for f in result["wide"]["bars"]], [["BAR_CHARTS_ANNOTATED", "blocker", None]])
        said = result["wide"]["bars"][0]["repair"]
        self.assertTrue(said.startswith("4 of 8 chart pages (50%) against the delivery floor of 65%"), said)
        self.assertIn("read over the 12 pages this revision composed; not held on the 8 slides it carries from the source deck", said)
        self.assertIn("the compile and the build hold it by the same number over the same pages", said)
        self.assertIn('record `waivers: [{ "code": "BAR_CHARTS_ANNOTATED"', said)
        self.assertEqual([[f["code"], f["severity"]] for f in result["waived"]["bars"]], [["BAR_CHARTS_ANNOTATED", "advisory"]])
        # A rebuild is held to the bars over the whole deck, as a new deck is.
        self.assertEqual([[f["code"], f["severity"]] for f in result["rebuilt"]["bars"]], [["BAR_CHARTS_ANNOTATED", "blocker"]])
        self.assertNotIn("not held", result["rebuilt"]["bars"][0]["repair"])
        # A revision recorded under an older version is refused at the compile as it is at delivery: delivery has refused a
        # missed bar of every deck whatever version it records, so the bars are no version's rule and the stages hold one thing.
        self.assertEqual([[f["code"], f["severity"], f.get("rule"), f["waived"]] for f in result["older"]["bars"]], [["BAR_CHARTS_ANNOTATED", "blocker", None, None]])
        self.assertEqual(result["older"]["delivery"], ["BAR_CHARTS_ANNOTATED"])


SERIES = '''
import { authorDeck, workedExamples } from './skills/professional-slides/runtime/author-deck.mjs';
import { bindDeck } from './skills/professional-slides/runtime/bind.mjs';
const worked = workedExamples();
const FY = ['FY21', 'FY22', 'FY23', 'FY24', 'FY25', 'FY26', 'FY27', 'FY28'];
const graded = { finding: 'a finding', calculation: 'a calculation', soWhat: 'It bears on the decision at hand', strength: 'strong', sources: ['sources/a.csv'] };
const items = [{ ...graded, id: 'i1', shape: 'series', measures: {
    journeys: { unit: 'million', population: 'the operator', periods: FY.slice(0, 6), values: [14.9, 31.6, 42.8, 46.1, 47.5, 48.3] },
    plan: { unit: 'million', population: 'the operator', periods: FY.slice(6), values: [49.0, 49.6], assumed: true, rationale: 'the plan the board approved in March' },
    held: { unit: 'million', population: 'the operator', periods: FY, values: FY.map(() => 45), assumed: true, rationale: 'the level the team assumes the rival holds' },
    rival: { unit: 'million', population: 'the rival', periods: FY, values: [20, 25, 30, 36, 41, 45, 50, 55], assumed: true, rationale: 'an estimate made by the team' },
    threshold: { unit: 'million', population: 'the plan', period: 'FY26', value: 45, standard: true } } }];
const insights = new Map(items.map((i) => [i.id, i]));
const base = { type: 'trend', form: 'line', commentary: 'so-what-bar', why: 'The movement over time is the claim, so the series is drawn',
  bar: 'The plan needs the journeys to keep rising at the pace of the last three years, which the assumed path does not deliver' };
const bound = (id, series, more = {}) => ({ ...base, id, evidence: ['i1'], settles: { kind: 'rate', what: 'journeys a year', measures: ['i1/journeys'] }, title: `Journeys against the plan, case ${id}`, exhibit: { heading: 'Passenger journeys', series, ...more } });
const typed = (id, series, more = {}) => ({ ...base, id, settles: { kind: 'rate', what: 'journeys a year' }, title: `Journeys against the plan, case ${id}`, exhibit: { heading: 'Passenger journeys', unit: 'million', categories: FY, series, ...more } });
const REC = [14.9, 31.6, 42.8, 46.1, 47.5, 48.3, 49.0, 49.6], RIVAL = [20, 25, 30, 36, 41, 45, 50, 55];
// What a composed page draws of its chart: the strokes, the markers, the reference rule and the names in its key.
const drawn = (deck, id) => { const slide = deck.slides.find((s) => (s.sourceSlideId ?? s.id) === id), nodes = slide?.nodes ?? [], count = (test) => nodes.filter(test).length;
  return { lines: count((n) => n.role === 'chart-line'), dashed: count((n) => n.role === 'chart-line' && n.style?.dash === 'dash'), markers: count((n) => n.role === 'chart-marker'), open: count((n) => n.role === 'chart-marker' && n.data?.state === 'assumed'),
    reference: nodes.filter((n) => n.role === 'chart-reference-label').map((n) => n.text), rules: nodes.filter((n) => n.role === 'chart-reference-line').map((n) => n.style?.dash),
    named: nodes.filter((n) => /end-label|legend-label|data-label/.test(n.role) && /[A-Za-z]/.test(n.text ?? '')).map((n) => n.text), periods: nodes.filter((n) => n.role === 'chart-period-label').map((n) => n.text),
    bars: count((n) => n.role === 'chart-mark'), lighter: count((n) => n.role === 'chart-mark' && n.data?.state === 'assumed'),
    native: (slide?.componentInstances ?? []).filter((c) => String(c.component).startsWith('chart.')).map((c) => Boolean(c.nativeChart)) }; };
const TYPED = [
  typed('t1', [{ name: 'Journeys', values: REC, assumedFrom: 'FY27' }, { name: 'Plan threshold', values: FY.map(() => 45), state: 'reference' }]),
  typed('t2', [{ name: 'Journeys', values: REC }, { name: 'Rival path', values: RIVAL, state: 'assumed' }]),
  { ...typed('t3', [{ name: 'Journeys', values: REC, assumedFrom: 'FY27' }, { name: 'Capacity limit', values: FY.map(() => 52), state: 'reference' }]), form: 'column' },
  typed('t4', [{ name: 'Journeys', values: REC }, { name: 'Rival path', values: RIVAL }])];
'''


class SeriesStateTests(unittest.TestCase):
    def test_the_binding_tells_the_chart_which_series_are_assumed_or_a_reference_and_they_are_drawn_so(self):
        result = run_node(SERIES + '''
const pages = [
  bound('b1', [{ measure: ['i1/journeys', 'i1/plan'], name: 'Journeys' }, { measure: 'i1/threshold', name: 'Plan threshold' }]),
  bound('b2', [{ measure: 'i1/journeys', name: 'Journeys' }, { measure: 'i1/rival', name: 'Rival path' }], { select: { from: 'FY21', to: 'FY26' } }),
  bound('b3', [{ measure: ['i1/journeys', 'i1/plan'], name: 'Journeys' }, { measure: 'i1/held', name: 'Rival, held flat' }]),
  bound('b4', [{ measure: ['i1/journeys', 'i1/plan'], name: 'Journeys', state: 'recorded' }])];
const doc = { deck: { ...worked.deck, id: 'refs' }, pages };
const written = bindDeck(doc, insights).doc.pages;
const result = await authorDeck(doc, { baseDir: '.', insights, fit: false, fill: false });
console.log(JSON.stringify(Object.fromEntries(pages.map((p) => [p.id, { series: written.find((w) => w.id === p.id).exhibit.series.map((s) => [s.name, s.state ?? null, s.assumedFrom ?? null]), note: written.find((w) => w.id === p.id).note ?? null, ...drawn(result.deck, p.id) }]))));
''')
        b1, b2, b3, b4 = (result[key] for key in ("b1", "b2", "b3", "b4"))
        # A recorded series joined to a scenario turns at the first assumed period; a measure of one value is a reference.
        self.assertEqual(b1["series"], [["Journeys", None, "FY27"], ["Plan threshold", "reference", None]])
        # The record keeps its stroke and filled markers; the assumed run is dashed with open markers, and the turn is bracketed.
        self.assertEqual([b1["lines"], b1["dashed"], b1["markers"], b1["open"]], [7, 2, 8, 2])
        self.assertEqual(b1["periods"], ["Recorded", "Assumed"])
        # The threshold is a dashed rule across the plot, named and valued at its end, with no markers and no place among the series.
        self.assertEqual([b1["reference"], b1["rules"]], [["Plan threshold 45.0"], ["dash"]])
        self.assertEqual(b1["named"], ["Journeys 49.6"])
        # A series that is an assumption throughout is dashed end to end and says so in its key; the record beside it is not bracketed as assumed.
        self.assertEqual(b2["series"], [["Journeys", None, None], ["Rival path", "assumed", None]])
        self.assertEqual([b2["lines"], b2["dashed"], b2["markers"], b2["open"]], [10, 5, 12, 6])
        self.assertEqual(b2["named"], ["Journeys 48.3", "Rival path (assumed) 45.0"])
        self.assertEqual(b2["periods"], [])
        self.assertEqual(b2["note"], "Assumed, not recorded: Rival path.")
        # An assumption held at one level is a reference line, not a recorded series; the bracket marks where the record ends, and the page says the level is assumed.
        self.assertEqual(b3["series"][1], ["Rival, held flat", "reference", None])
        self.assertEqual([b3["reference"], b3["markers"], b3["periods"]], [["Rival, held flat 45.0"], 8, ["Recorded", "Assumed"]])
        self.assertEqual(b3["note"], "Assumed, not recorded: Rival, held flat.")
        # A series that says its own state keeps it.
        self.assertEqual(b4["series"], [["Journeys", "recorded", None]])
        self.assertEqual(b4["dashed"], 0)
        # A chart that draws a series by its state is assembled as shapes, so the exported file carries the same strokes.
        for case in (b1, b2, b3):
            self.assertEqual(case["native"], [False])

    def test_a_typed_exhibit_declares_the_same_per_series_on_a_deck_with_no_insight_log(self):
        result = run_node(SERIES + '''
// A revision's pages are typed, and it has no log: the series say what they are themselves.
const doc = { deck: { ...worked.deck, id: 'refs', workflow: 'existing_deck_revision', request: undefined }, pages: TYPED };
const result = await authorDeck(doc, { baseDir: '.', fit: false, fill: false });
const wrong = await authorDeck({ ...doc, pages: [typed('w1', [{ name: 'Journeys', values: REC, state: 'forecast' }, { name: 'Rival path', values: RIVAL }])] }, { baseDir: '.', fit: false, fill: false });
const late = await authorDeck({ ...doc, pages: [typed('w2', [{ name: 'Journeys', values: REC, assumedFrom: 'FY31' }, { name: 'Rival path', values: RIVAL }])] }, { baseDir: '.', fit: false, fill: false });
const refused = (r) => r.blocking.filter((f) => f.code === 'PAGE_DOES_NOT_COMPOSE').map((f) => f.repair);
console.log(JSON.stringify({ composes: refused(result), wrong: refused(wrong), late: refused(late), ...Object.fromEntries(TYPED.map((p) => [p.id, drawn(result.deck, p.id)])) }));
''')
        self.assertEqual(result["composes"], [])
        t1, t2, t3, t4 = (result[key] for key in ("t1", "t2", "t3", "t4"))
        self.assertEqual([t1["lines"], t1["dashed"], t1["markers"], t1["open"]], [7, 2, 8, 2])
        self.assertEqual([t1["reference"], t1["rules"], t1["periods"]], [["Plan threshold 45.0"], ["dash"], ["Recorded", "Assumed"]])
        self.assertEqual([t2["dashed"], t2["open"], t2["named"]], [7, 8, ["Journeys 49.6", "Rival path (assumed) 55.0"]])
        # Columns: the assumed years take a lighter fill, the limit is a rule across them, and the turn is bracketed.
        self.assertEqual([t3["bars"], t3["lighter"], t3["reference"], t3["periods"]], [8, 2, ["Capacity limit 52.0"], ["Recorded", "Assumed"]])
        # Series that say nothing are drawn as recorded data, as before.
        self.assertEqual([t4["dashed"], t4["open"], t4["reference"], t4["periods"]], [0, 0, [], []])
        self.assertEqual(len(result["wrong"]), 1)
        self.assertIn('`state` is one of "recorded", "assumed", "reference"', result["wrong"][0])
        self.assertEqual(len(result["late"]), 1)
        self.assertIn("`assumedFrom` names the category its assumed run starts at (FY21, FY22", result["late"][0])

    @requires_python_package("pptx")
    def test_the_exported_file_carries_the_dashed_strokes_and_no_native_series_for_them(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = run_node(SERIES + f'''
import fs from 'node:fs';
const doc = {{ deck: {{ ...worked.deck, id: 'refs' }}, pages: [worked.pages.find((p) => p.type === 'summary'), ...TYPED] }};
const result = await authorDeck(doc, {{ baseDir: '.', fit: false, fill: false }});
fs.writeFileSync({json.dumps(str(Path(tmp) / "refs.deck.json"))}, JSON.stringify(result.spec));
console.log(JSON.stringify({{ scene: Object.fromEntries(TYPED.map((p) => [p.id, result.deck.slides.find((s) => (s.sourceSlideId ?? s.id) === p.id).nodes.filter((n) => n.type === 'line' && n.style?.dash === 'dash').length])) }}));
''')
            built = subprocess.run([NODE, str(RUNTIME / "build-deck.mjs"), str(Path(tmp) / "refs.deck.json"), str(Path(tmp) / "out"), "--no-render"], capture_output=True, text=True, cwd=ROOT, timeout=600)
            pptx = Path(tmp) / "out" / "refs.pptx"
            self.assertTrue(pptx.exists(), built.stderr[-1500:])
            from pptx import Presentation

            def read(slide):
                found = {"dashed": 0, "charts": 0}

                def walk(shapes):
                    for shape in shapes:
                        if shape.shape_type == 6:
                            walk(shape.shapes)
                            continue
                        if getattr(shape, "has_chart", False) and shape.has_chart:
                            found["charts"] += 1
                        if "Connector" in type(shape).__name__ and shape.line.dash_style is not None:
                            found["dashed"] += 1
                walk(slide.shapes)
                return found
            # The cover and the summary lead the file; the four chart pages follow in order.
            slides = [read(slide) for slide in Presentation(str(pptx)).slides][-4:]
        # Every dashed stroke of the scene is a dashed connector in the file - the assumed run, the reference rule, the bracket's divider.
        self.assertEqual([slide["dashed"] for slide in slides], [out["scene"][key] for key in ("t1", "t2", "t3", "t4")])
        self.assertEqual([slide["dashed"] > 0 for slide in slides], [True, True, True, False])
        # A chart that states a series' state stays drawn; the plain one is the native chart it always was.
        self.assertEqual([slide["charts"] for slide in slides], [0, 0, 0, 1])


MARKS = '''
import fs from 'node:fs';
import path from 'node:path';
import { authorDeck, readInsights, scaffoldReport } from './skills/professional-slides/runtime/author-deck.mjs';
import { claimMarks } from './skills/professional-slides/runtime/gates/dependency_gates.mjs';
import { pageChartAnnotated } from './skills/professional-slides/runtime/build-bars.mjs';
const dir = DIR;
const load = () => JSON.parse(fs.readFileSync(path.join(dir, 'finance.pages.json'), 'utf8'));
const insights = await readInsights(dir, 'finance', {});
const page = (doc, id) => doc.pages.find((p) => p.id === id);
'''


class ClaimMarkTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls._tmp = tempfile.TemporaryDirectory()
        cls.dir = Path(cls._tmp.name)
        for file in FIXTURES.glob("finance.*"):
            shutil.copy(file, cls.dir / file.name)
        cls.probe = MARKS.replace("DIR", json.dumps(str(cls.dir)))

    @classmethod
    def tearDownClass(cls):
        cls._tmp.cleanup()

    def test_a_bare_chart_is_told_the_mark_its_title_names_and_the_mark_is_proven_on_the_page(self):
        result = run_node(self.probe + '''
const doc = load();
// A chart under a strip of figures, its mark taken off: the title states the measure at FY26 by a token.
delete page(doc, 'f2').exhibit.highlights; page(doc, 'f2').title = 'Operating cash flow fell to {{i-cash/ocf@FY26}} million while profit rose again';
// A panel's chart, its mark taken off: the title types a number that is one value of the measure the panel plots.
delete page(doc, 'f4').exhibits[0].highlights; page(doc, 'f4').title = 'Costs fell to 62 percent of income without closing branches';
const proposed = claimMarks(doc, insights);
const before = JSON.stringify(doc);
const result = await authorDeck(doc, { baseDir: dir, insights, fit: false, fill: false });
const told = [...result.blocking, ...result.advisories].filter((f) => f.code === 'UNANNOTATED').map((f) => ({ id: f.id, marks: f.marks ?? null, repair: f.repair }));
// The mark written as the finding says it: the page then marks its chart.
const marked = structuredClone(doc); page(marked, 'f2').exhibit.highlights = proposed.get('f2')[0].mark.highlights; page(marked, 'f4').exhibits[0].highlights = proposed.get('f4')[0].mark.highlights;
const after = await authorDeck(marked, { baseDir: dir, insights, fit: false, fill: false });
const annotated = (r, id) => pageChartAnnotated(r.deck.slides.find((s) => (s.sourceSlideId ?? s.id) === id));
console.log(JSON.stringify({ proposed: Object.fromEntries(proposed), untouched: JSON.stringify(doc) === before, told, bare: [annotated(result, 'f2'), annotated(result, 'f4')], marked: [annotated(after, 'f2'), annotated(after, 'f4')],
  after: [...after.blocking, ...after.advisories].filter((f) => f.marks).map((f) => f.id) }));
''')
        self.assertEqual(result["proposed"]["f2"], [{"exhibit": '"Operating cash flow"', "at": 0, "category": ["FY26"], "mark": {"highlights": [{"category": "FY26"}]}, "because": "the title's 66 is i-cash/ocf at FY26"}])
        self.assertEqual(result["proposed"]["f4"][0]["because"], "the title's 62 percent is i-efficiency/cost-income at FY26")
        # Nothing is written into the page: the author is told, and the copy stays theirs.
        self.assertTrue(result["untouched"])
        self.assertEqual(result["bare"], [False, False])
        told = {item["id"]: item for item in result["told"]}
        self.assertIn('The title already names what to mark: `"highlights": [{"category":"FY26"}]` on "Operating cash flow" (the title\'s 66 is i-cash/ocf at FY26)', told["f2"]["repair"])
        self.assertIn('`"highlights": [{"category":"FY26"}]` on "Cost-to-income ratio"', told["f4"]["repair"])
        self.assertEqual(told["f2"]["marks"][0]["mark"], {"highlights": [{"category": "FY26"}]})
        # A page whose title names no value of what its chart plots is told nothing it cannot use.
        self.assertTrue(all(item["marks"] is None for key, item in told.items() if key not in ("f2", "f4")))
        self.assertEqual(result["marked"], [True, True])
        self.assertEqual(result["after"], [])

    def test_a_chart_alternative_the_fit_search_offers_carries_the_mark_the_bare_chart_is_asked_for(self):
        result = run_node(self.probe + '''
import { fitLines } from './skills/professional-slides/runtime/fit-search.mjs';
const doc = load();
// A panel's chart with its mark taken off, under a title that names the value to mark; and headings too long for two panels
// side by side, so the page does not fit its row and the fit search offers the forms that do.
delete page(doc, 'f4').exhibits[0].highlights; page(doc, 'f4').title = 'Costs fell to 62 percent of income without closing branches';
page(doc, 'f4').exhibits[0].heading = 'Cost-to-income ratio of the bank across the whole of its retail and commercial network at each financial year end';
page(doc, 'f4').exhibits[1].heading = 'Branches open at the end of each financial year across the retail network in every region the bank serves';
const result = await authorDeck(doc, { baseDir: dir, insights, fill: false, fitPages: new Set(['f4']) });
const fit = result.fits.get('f4');
const told = [...result.blocking, ...result.advisories].find((f) => f.code === 'UNANNOTATED' && f.id === 'f4');
// The same page with the mark written: the chart is not bare, and the alternatives have nothing to carry.
const marked = structuredClone(doc); page(marked, 'f4').exhibits[0].highlights = [{ category: 'FY26' }];
const after = (await authorDeck(marked, { baseDir: dir, insights, fill: false, fitPages: new Set(['f4']) })).fits.get('f4');
console.log(JSON.stringify({ misfit: result.blocking.filter((f) => f.id === 'f4').map((f) => f.code), told: told?.marks ?? null, pass: fit.pass, partial: fit.partial, lines: fitLines(fit).lines,
  after: [...after.pass, ...after.partial].map((alt) => [alt.form, alt.commentary, alt.marks ?? null]) }));
''')
        self.assertIn("HEADING_WRAPS", result["misfit"])
        mark = {"highlights": [{"category": "FY26"}]}
        self.assertEqual(result["told"][0]["mark"], mark)
        offered = result["pass"] + result["partial"]
        self.assertIn(["stack", "captions"], [[alt["form"], alt["commentary"]] for alt in result["pass"]])
        # Every alternative that draws the chart is offered with the exact mark the page's own finding names, proven drawn on it.
        for alt in offered:
            self.assertEqual([item["mark"] for item in alt["marks"]], [mark], alt)
            self.assertEqual(alt["marks"], result["told"])
        said = next(line for line in result["lines"] if 'form "stack", commentary "captions"' in line)
        self.assertIn('composed with the mark the page\'s bare chart is asked for - `"highlights": [{"category":"FY26"}]` on "Cost-to-income ratio', said)
        self.assertIn("which this form draws: write it with the form", said)
        # A page that already marks its chart is offered its alternatives as it always was.
        self.assertTrue(result["after"])
        self.assertTrue(all(marks is None for _, _, marks in result["after"]))

    def test_the_bar_a_deck_misses_says_which_of_its_bare_pages_already_name_their_mark(self):
        result = run_node(self.probe + '''
const doc = load();
// Twelve chart pages with nothing marked, each title stating its chart's last value; four more whose titles name no value of theirs.
const bare = (id, title) => { const copy = structuredClone(page(load(), 'f2')); delete copy.exhibit.highlights; return { ...copy, id, title }; };
doc.pages = [page(doc, 'f0'), ...Array.from({ length: 12 }, (_, i) => bare(`g${i + 1}`, `Operating cash flow fell to {{i-cash/ocf@FY26}} million, as page ${i + 1} shows`)),
  ...Array.from({ length: 4 }, (_, i) => bare(`h${i + 1}`, `Lending absorbed the cash that earnings produced, as page ${i + 13} shows`))];
const result = await authorDeck(doc, { baseDir: dir, insights, fit: false, fill: false });
const rate = result.blocking.filter((f) => ['BAR_CHARTS_ANNOTATED', 'CRAFT_CHARTS_BARE'].includes(f.code)).map((f) => ({ code: f.code, marked: Object.keys(f.marks ?? {}), repair: f.repair }));
console.log(JSON.stringify({ rate, standing: result.standings.filter((st) => st.code === 'BAR_CHARTS_ANNOTATED').map((st) => st.line) }));
''')
        # One finding for the rate, and it names the pages whose own titles say what to mark - not the ones that do not.
        self.assertEqual(len(result["rate"]), 1, result["rate"])
        self.assertEqual(result["rate"][0]["marked"], [f"g{n}" for n in range(1, 13)])
        self.assertIn('On 12 of the bare chart pages the title already names what to mark - g1: `"highlights": [{"category":"FY26"}]` on "Operating cash flow" (the title\'s 66 is i-cash/ocf at FY26)', result["rate"][0]["repair"])
        self.assertIn("delivery floor 65%: short by", result["standing"][0])

    def test_a_scaffold_bound_to_a_series_marks_its_latest_period_where_the_form_would_draw_it_bare(self):
        result = run_node(self.probe + '''
import { markedChart } from './skills/professional-slides/runtime/page-types.mjs';
const series = [...insights.values()].find((i) => i.id === 'i-liquidity');
const report = scaffoldReport('panels', { form: 'row', insight: series });
const doc = { ...load(), pages: [page(load(), 'f0'), { ...report.page, id: 'sc1', title: 'Liquid assets and the liabilities due within a year, side by side' }] };
const result = await authorDeck(doc, { baseDir: dir, insights, fit: false, fill: false });
const slide = result.deck.slides.find((s) => (s.sourceSlideId ?? s.id) === 'sc1');
console.log(JSON.stringify({ bound: report.bound, exhibits: report.page.exhibits.map((ex) => [ex.type, ex.highlights ?? null, markedChart(ex)]), composes: result.blocking.filter((f) => f.id === 'sc1' && /COMPILE|COMPOSE/.test(f.code)).map((f) => f.repair), annotated: slide ? pageChartAnnotated(slide) : null }));
''')
        self.assertTrue(result["bound"])
        self.assertTrue(result["exhibits"])
        for kind, highlights, marked in result["exhibits"]:
            self.assertEqual(highlights, [{"category": "FY26"}], kind)
            self.assertTrue(marked)
        self.assertEqual(result["composes"], [])
        self.assertTrue(result["annotated"])


if __name__ == "__main__":
    unittest.main()
