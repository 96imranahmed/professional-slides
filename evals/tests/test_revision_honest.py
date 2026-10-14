"""A revision is honest about what it changed and what it is held to.

An independent review of the point-change path found that one carried slide
turned every deck rule into an advisory, that a text edit rewrote longer
numbers it happened to stand inside, that a changed figure left standing in
other words, at another scale, in a chart or in the speaker's notes was not
found, that a title rewritten through `replace` staged no critique, that a
cut slide left its chart and workbook in the file, and that the compile and
the build disagreed about what a revision is held to. These hold the repairs,
each on a deck made here: a small deck written with python-pptx that this
skill did not build, imported with `--carry`, and - for the deck rules - the
shipped worked example set beside one carried slide.
"""
from __future__ import annotations

import copy
import json
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from node_probe import HAS_PPTX, NODE, ROOT, RUNTIME, RUNTIME_PYTHON, run_node

PYTHON = RUNTIME_PYTHON or sys.executable
IMPORT = RUNTIME / "import-deck.py"
ASSEMBLE = RUNTIME / "emit" / "assemble_pptx.py"
AUTHOR = RUNTIME / "author-deck.mjs"
BUILD = RUNTIME / "build-deck.mjs"
EXAMPLE = ROOT / "skills" / "professional-slides" / "examples"


def sh(*command, cwd=None):
    return subprocess.run([str(part) for part in command], cwd=cwd, capture_output=True, text=True)


PROBE = """
import fs from 'node:fs';
import path from 'node:path';
import { authorDeck } from './skills/professional-slides/runtime/author-deck.mjs';
import { loadJudgements, withJudgements } from './skills/professional-slides/runtime/judgements.mjs';
const dir = DIR, doc = JSON.parse(fs.readFileSync(path.join(dir, NAME + '.pages.json'), 'utf8'));
// The answers recorded beside the pages file are read, as the compile reads them.
const result = await withJudgements(await loadJudgements(dir, doc.deck?.id ?? NAME), () => authorDeck(doc, { baseDir: dir, insights: null, fit: false, fill: false }));
const brief = (f) => ({ code: f.code, id: f.id ?? null, severity: f.severity, pages: f.pages ?? null, measured: f.measured ?? null, waived: f.waived ?? null, repair: f.repair ?? f.reason ?? '' });
console.log(JSON.stringify({ blocking: result.blocking.map(brief), advisories: result.advisories.map(brief) }));
"""


def judged(work: Path, name: str = "deck"):
    """What the compile says of the pages file in `work`, finding by finding: `{ blocking, advisories }`."""
    return run_node(PROBE.replace("DIR", json.dumps(str(work))).replace("NAME", json.dumps(name)))


def last_json(text: str):
    at = text.rfind("\n{")
    return json.loads(text[at + 1:] if at >= 0 else text)


# A deck a user might own: a figure stated in a sentence, in other words, at another scale, in a table, in two charts and in
# the speaker's notes; numbers that hold its digits inside longer ones; a slide to cut, with a chart of its own and a
# picture another slide shows too.
SLIDES = [
    {"cover": True, "title": "Orchard Bakeries: annual review", "subtitle": "Partners' pre-read"},
    {"title": "Summary", "lines": ["Sales reached £31.7m in the year, up 9% on the year before", "Cost per loaf was £131.75 against a plan of £131.7", "Shop margin held at 31.7% across the estate"]},
    {"title": "Sales in other words", "lines": ["Sales for the year were £31.7 million", "Sales of GBP 31.70m in the year", "Sales of £31,700k in the year", "Sales: 31.7"]},
    {"title": "Regions", "table": [["Region", "Sales (£m)"], ["Coast", "9.4"], ["Hills", "10.2"], ["Valley", "12.1"], ["Total", "31.7"]]},
    {"title": "Sales by year", "chart": {"categories": ["Last year", "This year"], "series": {"Sales": [29.1, 31.7]}}},
    {"title": "Sales by year, to the penny", "chart": {"categories": ["Last year", "This year"], "series": {"Sales": [29.08, 31.73]}}},
    {"title": "Margin by shop", "lines": ["Margin was 31.7% at the harbour shop"], "notes": "Remind them that sales were £31.7m"},
    {"title": "The mill we may buy", "chart": {"categories": ["Last year", "This year"], "series": {"Mill profit": [4.2, 5.9]}}, "picture": True, "notes": "Not for circulation: the offer is nine times profit"},
    {"title": "Close", "lines": ["The year closed on sales of £31.7m", "The year closed on sales of £31.7m again, said twice"], "picture": True},
]


def build_deck(path: Path, slides=SLIDES) -> Path:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    from PIL import Image

    path.parent.mkdir(parents=True, exist_ok=True)
    picture = path.parent / "mark.png"
    Image.new("RGB", (64, 64), (40, 90, 140)).save(picture)
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    for spec in slides:
        if spec.get("cover"):
            slide = deck.slides.add_slide(deck.slide_layouts[0])
            slide.shapes.title.text, slide.placeholders[1].text = spec["title"], spec.get("subtitle", "")
            continue
        slide = deck.slides.add_slide(deck.slide_layouts[1] if spec.get("lines") else deck.slide_layouts[5])
        slide.shapes.title.text = spec["title"]
        if spec.get("lines"):
            frame = slide.placeholders[1].text_frame
            for at, line in enumerate(spec["lines"]):
                (frame.paragraphs[0] if at == 0 else frame.add_paragraph()).text = line
        if spec.get("table"):
            rows = spec["table"]
            table = slide.shapes.add_table(len(rows), len(rows[0]), Inches(1), Inches(2), Inches(10), Inches(3)).table
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    table.cell(r, c).text = cell
        if spec.get("chart"):
            data = CategoryChartData()
            data.categories = spec["chart"]["categories"]
            for name, values in spec["chart"]["series"].items():
                data.add_series(name, values)
            slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(8), Inches(4.5), data)
        if spec.get("picture"):
            slide.shapes.add_picture(str(picture), Inches(11.5), Inches(0.4), Inches(1), Inches(1))
        if spec.get("notes"):
            slide.notes_slide.notes_text_frame.text = spec["notes"]
    deck.save(path)
    return path


@unittest.skipUnless(NODE and HAS_PPTX, "needs Node.js and python-pptx")
class RevisedDeck(unittest.TestCase):
    """The deck, built and imported once; each test revises a copy of it."""

    @classmethod
    def setUpClass(cls):
        cls.home = Path(tempfile.mkdtemp(prefix="revision-honest-"))
        cls.source = build_deck(cls.home / "source" / "orchard.pptx")
        cls.bed = cls.home / "bed"
        done = sh(PYTHON, IMPORT, cls.source, cls.bed, "--carry", "--id", "deck")
        assert done.returncode == 0, done.stderr
        cls.starter = json.loads((cls.bed / "deck.pages.json").read_text())
        cls.inventory = json.loads((cls.bed / "deck.inventory.json").read_text())

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.home, ignore_errors=True)

    def revise(self, change=None):
        work = Path(tempfile.mkdtemp(prefix="rev-", dir=self.home))
        shutil.copytree(self.bed, work, dirs_exist_ok=True)
        doc = copy.deepcopy(self.starter)
        doc["deck"]["request"] = "Sales for the year have been restated from £31.7m to £31.9m."
        if change:
            change(doc)
        (work / "deck.pages.json").write_text(json.dumps(doc, indent=1, ensure_ascii=False))
        return work

    @staticmethod
    def page(doc, key):
        return next(page for page in doc["pages"] if page["id"] == key)

    def edit(self, key, *edits):
        return self.revise(lambda doc: self.page(doc, key).update(replace=list(edits)))

    def author(self, work, *flags):
        return sh(NODE, AUTHOR, work / "deck.pages.json", *flags)

    def log(self, work):
        return judged(work)["blocking"]

    def refused(self, work):
        return [(f["code"], f.get("id")) for f in self.log(work)]

    def assembled(self, work):
        """The revision compiled and built without a render; returns the build result and the written deck as a zip."""
        done = self.author(work)
        self.assertEqual(done.returncode, 0, done.stderr[:1500])
        built = sh(NODE, BUILD, work / "deck.deck.json", work / "out", "--no-render")
        self.assertEqual(built.returncode, 0, (built.stdout + built.stderr)[-1500:])
        result = json.loads((work / "out" / "build-result.json").read_text())
        return result, zipfile.ZipFile(result["pptxPath"])

    def slide_text(self, archive, number):
        from pptx import Presentation
        slide = Presentation(archive.filename).slides[number - 1]
        texts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
        texts += [cell.text for shape in slide.shapes if shape.has_table for row in shape.table.rows for cell in row.cells]
        return texts, (slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else "")


class WholeWordTests(RevisedDeck):
    """An edit rewrites whole words, and the stale check reads the slide the way the edit does."""

    CASES = [("Sales reached £31.7m in the year", "31.7"), ("Sales reached £31.7m in the year", "£31.7m"), ("a margin of 31.7% held", "31.7"), ("a margin of 31.7% held", "31.7%"),
             ("Cost was £131.75 against £131.7", "31.7"), ("Sales: 31.7", "31.7"), ("(31.7) in brackets, and 31.7, twice", "31.7"), ("The shop's lease ends", "shop"),
             ("The shop, then the shop.", "shop"), ("Total 31.7.5 is a version", "31.7"), ("year-on-year growth", "year"), ("growth year on year—and more", "year"),
             ("“Sales rose” she said", "Sales rose"), ("Sales rose", "Sales  rose"), ("", "x"), ("31.7", "31.7")]

    def test_the_edit_and_the_stale_check_share_one_reading_of_what_a_slide_prints(self):
        script = "import sys, json, importlib.util\nspec = importlib.util.spec_from_file_location('assemble', sys.argv[1])\nmodule = importlib.util.module_from_spec(spec)\nspec.loader.exec_module(module)\n" \
                 "print(json.dumps([[len(module.whole_words(text, words)[0]), module.whole_words(text, words)[1]] for text, words in json.loads(sys.argv[2])]))"
        python = json.loads(sh(PYTHON, "-c", script, ASSEMBLE, json.dumps(self.CASES)).stdout)
        node = run_node(f"""
import {{ wholeWords, printsWords }} from './skills/professional-slides/runtime/revision.mjs';
console.log(JSON.stringify({json.dumps(self.CASES)}.map(([text, words]) => {{ const found = wholeWords(text, words); return [found.whole.length, found.inside.map((item) => item.token), printsWords(text, words)]; }})));
""")
        self.assertEqual([row[:2] for row in node], python)
        whole = {(text, words): row[0] for (text, words), row in zip(self.CASES, node)}
        # A number is not printed by a longer number, by a currency amount or by a percentage that holds its digits.
        self.assertEqual(whole[("Sales reached £31.7m in the year", "31.7")], 0)
        self.assertEqual(whole[("a margin of 31.7% held", "31.7")], 0)
        self.assertEqual(whole[("Cost was £131.75 against £131.7", "31.7")], 0)
        self.assertEqual(whole[("Total 31.7.5 is a version", "31.7")], 0)
        # Brackets, quotes and the punctuation that ends a clause are not part of the word.
        self.assertEqual(whole[("(31.7) in brackets, and 31.7, twice", "31.7")], 2)
        self.assertEqual(whole[("The shop, then the shop.", "shop")], 2)
        self.assertEqual(whole[("“Sales rose” she said", "Sales rose")], 1)
        # A possessive and a hyphenated compound are words of their own; an em dash set without spaces parts two.
        self.assertEqual(whole[("The shop's lease ends", "shop")], 0)
        self.assertEqual(whole[("year-on-year growth", "year")], 0)
        self.assertEqual(whole[("growth year on year—and more", "year")], 2)
        self.assertEqual([row[2] for row in node], [row[0] > 0 for row in node])

    def test_an_edit_that_stands_only_inside_longer_numbers_is_refused_with_them_named(self):
        # The bare figure: on 59e7d1a this rewrote the cost per loaf to £131.95 and the margin to 31.9%, and compiled.
        work = self.edit("s02", {"old": "31.7", "new": "31.9", "only": ["s03", "s04", "s07", "s09"]})
        done = self.author(work, "--check")
        self.assertEqual(done.returncode, 2)
        self.assertEqual(self.refused(work), [("REVISION_CARRY_INVALID", "s02")])
        said = self.log(work)[0]["repair"]
        self.assertIn('slide 2 prints "31.7" only inside longer words or numbers - "£31.7m", "£131.75", "£131.7", "31.7%"', said)
        self.assertIn('name the one you mean as the slide prints it (`"old": "£31.7m"`)', said)

    def test_an_edit_that_stands_more_than_once_is_refused_until_it_is_narrowed_or_says_all(self):
        twice = {"old": "£31.7m", "new": "£31.9m"}
        others = ["s02", "s03", "s04", "s07"]
        work = self.edit("s09", {**twice, "only": others})
        self.assertEqual(self.author(work, "--check").returncode, 2)
        said = self.log(work)[0]["repair"]
        self.assertEqual(self.refused(work), [("REVISION_CARRY_INVALID", "s09")])
        self.assertIn('slide 9 prints "£31.7m" 2 times - "The year closed on sales of £31.7m"; "The year closed on sales of £31.7m again, said twice"', said)
        self.assertIn('`"all": true`', said)
        # Narrowed to the one sentence, the other is left standing - and is then the edited slide's own stale copy.
        work = self.edit("s09", {"old": "sales of £31.7m again", "new": "sales of £31.9m again", "only": others})
        self.assertEqual(self.author(work, "--check").returncode, 2)
        self.assertEqual(self.refused(work), [("NUMBER_STALE", "s09")])
        self.assertIn('s09 itself still prints "£31.7m"', self.log(work)[0]["repair"])
        # `all` says every one is meant.
        work = self.edit("s09", {**twice, "all": True, "only": others})
        result, archive = self.assembled(work)
        texts, _ = self.slide_text(archive, 9)
        self.assertEqual([text for text in texts if "sales of" in text], ["The year closed on sales of £31.9m\nThe year closed on sales of £31.9m again, said twice"])
        self.assertEqual([edit["changes"] for edit in result["revision"]["edits"]], [[{"old": "£31.7m", "new": "£31.9m", "count": 2}]])

    def test_a_whole_word_edit_leaves_the_longer_numbers_beside_it_as_they_were(self):
        work = self.edit("s02", {"old": "£31.7m", "new": "£31.9m", "only": ["s03", "s04", "s07", "s09"]})
        _, archive = self.assembled(work)
        texts, _ = self.slide_text(archive, 2)
        body = next(text for text in texts if "Sales reached" in text)
        self.assertEqual(body, "Sales reached £31.9m in the year, up 9% on the year before\nCost per loaf was £131.75 against a plan of £131.7\nShop margin held at 31.7% across the estate")


# What the orchard deck's passages state, as a reader reads them.
ORCHARD = [["margin", "the shop margin"], ["last year", "last year's sales"], ["sales|total", "this year's sales"]]


class CutSlideTests(RevisedDeck):
    def test_a_cut_slide_takes_its_chart_and_its_data_with_it_and_a_shared_picture_stays(self):
        source = zipfile.ZipFile(self.bed / "deck.source.pptx")
        charts = sorted(name for name in source.namelist() if name.startswith("ppt/charts/chart") and name.endswith(".xml"))
        self.assertEqual(len(charts), 3)
        mill = next(name for name in charts if b"Mill profit" in source.read(name))
        book = next(name for name in source.namelist() if name.startswith("ppt/embeddings/") and name in source.read(f"ppt/charts/_rels/{Path(mill).name}.rels").decode().replace("../", "ppt/"))
        media = [name for name in source.namelist() if name.startswith("ppt/media/")]
        self.assertEqual(len(media), 1)
        # Slide 8 is cut: the mill's economics must not circulate.
        work = self.revise(lambda doc: doc["pages"].remove(self.page(doc, "s08")))
        result, archive = self.assembled(work)
        names = archive.namelist()
        self.assertEqual(result["revision"]["dropped"], [8])
        # Its chart, the chart's relationships and its embedded workbook leave the file, and the build says so.
        self.assertNotIn(mill, names)
        self.assertNotIn(book, names)
        self.assertFalse(any(b"Mill profit" in archive.read(name) for name in names if name.endswith(".xml")))
        self.assertFalse(any(b"nine times profit" in archive.read(name) for name in names if name.endswith(".xml")))
        self.assertIn(mill, result["revision"]["removed"])
        self.assertIn(book, result["revision"]["removed"])
        # The picture slide 9 also shows is drawn on by a kept slide: it stays, byte for byte, and so does every other chart.
        self.assertIn(media[0], names)
        self.assertEqual(archive.read(media[0]), source.read(media[0]))
        self.assertNotIn(media[0], result["revision"]["removed"])
        self.assertEqual(sorted(name for name in names if name.startswith("ppt/charts/chart") and name.endswith(".xml")), sorted(set(charts) - {mill}))
        # Nothing in the package still points at what was removed, and the file opens.
        for name in names:
            if name.endswith(".rels"):
                self.assertNotIn(Path(mill).name, archive.read(name).decode())
        self.assertNotIn(Path(mill).name, archive.read("[Content_Types].xml").decode())
        from pptx import Presentation
        self.assertEqual(len(Presentation(archive.filename).slides), 8)
        self.assertEqual(result["revision"]["preserved"]["drifted"], [])


class OneDecisionTests(RevisedDeck):
    def test_an_evaluations_length_is_decided_once_for_the_compile_and_the_build(self):
        # A point change to a deck marked as an evaluation: the compile said the rule was not held, and the build refused it.
        work = self.revise(lambda doc: (doc["deck"].update(purpose="evaluation"), self.page(doc, "s07").update(replace=[{"old": "harbour shop", "new": "quay shop"}])))
        done = self.author(work)
        self.assertEqual(done.returncode, 0, done.stderr[:1500])
        built = sh(NODE, BUILD, work / "deck.deck.json", work / "out", "--no-render")
        self.assertEqual(built.returncode, 0, (built.stdout + built.stderr)[-800:])
        self.assertNotIn("EVALUATION_TOO_SHORT", built.stdout + built.stderr)


@unittest.skipUnless(NODE, "needs Node.js")
class ComposedPageChangesTests(unittest.TestCase):
    PROBE = '''
import { revisionChanges } from './skills/professional-slides/runtime/revision.mjs';
import { stampKept } from './skills/professional-slides/runtime/revision.mjs';
const slide = (index, title, more = {}) => ({ index, id: `s0${index}`, title, hidden: false, subtitle: null, paragraphs: [], tables: [], charts: [], pictures: [], notes: null, ...more });
const inventory = { source: { file: 'deck.pptx' }, slideSize: { width: 1280, height: 720 }, slides: [slide(1, 'Cover'), slide(2, 'Summary', { paragraphs: [{ text: 'Sales reached 31.7 in the year' }] }),
  slide(3, 'Sales by channel', { paragraphs: [{ text: 'Shops are 62 of every 100 sold' }], charts: [{ type: 'PIE', title: null, categories: ['Shops', 'Wholesale'], series: [{ name: 'Share', values: [62, 38] }] }] }), slide(4, 'Close')] };
const carry = (n, more = {}) => ({ id: `s0${n}`, sourceSlide: n, title: inventory.slides[n - 1].title, after: null, ...more });
const typed = (id, title, sourceSlide, more = {}) => ({ id, title, pageType: { type: 'composition', sourceSlide }, ...more });
const deck = (carried, slides = []) => ({ workflow: 'existing_deck_revision', inventory: 'deck.inventory.json', slides, carried });
'''

    def test_a_composed_page_is_the_slides_argument_only_where_it_says_what_the_slide_said(self):
        result = run_node(self.PROBE + '''
const redrawn = { exhibit: { type: 'chart.donut', labels: ['Shops', 'Wholesale'], values: [62, 38] }, points: ['Shops are 62 of every 100 sold'] };
const rest = [carry(1), carry(2), { ...carry(4), after: 's03' }];
console.log(JSON.stringify({
  // The pie redrawn as a donut under the slide's own title, with the slide's own words: nothing of the argument moved.
  same: revisionChanges(deck(rest, [typed('s03', 'Sales by channel', 3, redrawn)]), inventory),
  // Under the same title, an argument the slide never made: other numbers, other words.
  other: revisionChanges(deck(rest, [typed('s03', 'Sales by channel', 3, { exhibit: { type: 'chart.donut', labels: ['Shops', 'Wholesale', 'Online'], values: [48, 31, 21] }, points: ['Online is now a fifth of everything sold, and the only channel growing'] })]), inventory),
  // A source of title-only slides, each composed page naming one: the titles in common keep nothing from being read.
  shell: revisionChanges(deck([carry(1)], [typed('s02', 'Summary', 2, { points: ['Sales reached 31.7 in the year', 'Three shops opened and none closed, on a margin that held'] }), typed('s04', 'Close', 4, { points: ['Open four more shops next year'] })]),
    { ...inventory, slides: [inventory.slides[0], slide(2, 'Summary'), inventory.slides[2], slide(4, 'Close')] }) }));''')
        self.assertEqual((result["same"]["spine"], result["same"]["spineChanged"]), ([], False))
        self.assertEqual(result["same"]["drawn"], ["s03"])
        self.assertEqual(result["other"]["spine"], ["s03"])
        self.assertTrue(result["other"]["spineChanged"])
        self.assertEqual(result["shell"]["spine"], ["s02", "s04"])

    def test_what_a_composed_page_keeps_of_its_slide_is_checked_against_the_inventory_and_only_that_is_excused(self):
        result = run_node(self.PROBE + '''
import { applyRulesVersion } from './skills/professional-slides/runtime/weight.mjs';
const judge = (title) => { const spec = stampKept(deck([carry(1), carry(2), { ...carry(4), after: 's03' }], [typed('s03', title, 3)]), inventory);
  return [spec.slides[0].pageType.kept ?? null, ...['CONTENT_NO_CLAIM', 'SOURCE_UNCITED', 'TITLE_WORDS', 'PAGE_DOES_NOT_COMPOSE'].map((code) => applyRulesVersion([{ code, id: 's03', severity: 'blocker' }], spec)[0].severity)]; };
const added = stampKept(deck([carry(1)], [{ id: 'n01', title: 'Sales by channel', pageType: { type: 'composition' } }]), inventory);
console.log(JSON.stringify({ own: judge('Sales by channel'), spaced: judge('  Sales   by channel '), retitled: judge('Channel mix'), added: applyRulesVersion([{ code: 'CONTENT_NO_CLAIM', id: 'n01', severity: 'blocker' }], added)[0].severity,
  // A new deck is not a revision: nothing is kept.
  fresh: applyRulesVersion([{ code: 'CONTENT_NO_CLAIM', id: 's03', severity: 'blocker' }], { workflow: 'new_deck', slides: [{ id: 's03', pageType: { kept: { title: true } } }] })[0].severity }));''')
        # The slide's own title, word for word: the rule that a title is a claim advises, and every other rule on the page holds.
        self.assertEqual(result["own"], [{"title": True}, "advisory", "blocker", "blocker", "blocker"])
        self.assertEqual(result["spaced"][:2], [{"title": True}, "advisory"])
        # A title the revision wrote - a label of its own - is the revision's, and is refused; so is a page that stands for no slide.
        self.assertEqual(result["retitled"], [None, "blocker", "blocker", "blocker", "blocker"])
        self.assertEqual(result["added"], "blocker")
        self.assertEqual(result["fresh"], "blocker")

class ComposedPageOnARevisionTests(RevisedDeck):
    def test_a_breadth_recorded_under_another_shapes_key_is_told_the_key_not_sent_to_research(self):
        result = run_node("""
import { breadthProblem } from './skills/professional-slides/runtime/evidence.mjs';
const mix = (breadth) => breadthProblem({ id: 'i-mix', shape: 'mix', breadth });
console.log(JSON.stringify({ wrongKey: mix({ members: 4 }), narrow: mix({ parts: 2 }), right: mix({ parts: 4 }), series: breadthProblem({ id: 'i-s', shape: 'series', breadth: { members: 9 } }) }));""")
        # Four parts recorded as `members`: the data is wide enough, and the message says which key a mix is counted in.
        self.assertIn("`breadth` records `members`, and a mix is counted in `parts`", result["wrongKey"])
        self.assertIn("Write `breadth: { parts: n }`", result["wrongKey"])
        self.assertNotIn("research task", result["wrongKey"])
        # Data that really is too narrow is still a research task.
        self.assertIn("the data has 2 parts", result["narrow"])
        self.assertIn("research task", result["narrow"])
        self.assertIsNone(result["right"])
        self.assertIn("counted in `periods` and `series`", result["series"])

    def test_the_delivery_record_lists_every_change_the_revision_made(self):
        result = run_node("""
import { changesMade } from './skills/professional-slides/runtime/revision.mjs';
const revision = { order: [{ id: 's01', carried: 1 }, { id: 's02', carried: 2, edited: true }, { id: 's03', composed: 1 }, { id: 'n01', composed: 2 }, { id: 's05', carried: 5, edited: true }],
  edits: [{ slide: 2, changes: [{ old: '£31.7m', new: '£31.9m', count: 2 }, { title: 'Sales rose' }] }, { slide: 5, notes: true, changes: [{ old: '£31.7m', new: '£31.9m', count: 1 }, { hidden: true }] }],
  dropped: [4], removed: ['ppt/charts/chart2.xml', 'ppt/embeddings/Microsoft_Excel_Sheet2.xlsx'] };
const spec = { slides: [{ id: 's03', pageType: { sourceSlide: 3 } }, { id: 'n01', pageType: {} }] };
console.log(JSON.stringify({ made: changesMade(revision, spec).map((line) => [line.kind, line.text]), none: changesMade(null, spec) }));""")
        self.assertEqual(result["none"], [])
        self.assertEqual(result["made"], [
            ["replaced", 'slide 2 (s02): "£31.7m" replaced by "£31.9m", in 2 places'],
            ["retitled", 'slide 2 (s02): title set to "Sales rose"'],
            ["replaced", 'slide 5 (s05): "£31.7m" replaced by "£31.9m" (its speaker notes among them, where they printed it)'],
            ["hidden", "slide 5 (s05): hidden from the slide show"],
            ["redrawn", "slide 3 (s03): redrawn - composed by the runtime in the slide's place"],
            ["added", "page n01: added, composed by the runtime"],
            ["removed", "slide 4: removed"],
            ["parts", "2 parts only the removed slide drew on left the file with it: ppt/charts/chart2.xml, ppt/embeddings/Microsoft_Excel_Sheet2.xlsx"]])


@unittest.skipUnless(NODE, "needs Node.js")
class ClaimCitationTests(unittest.TestCase):
    def test_a_claim_about_a_measure_the_page_does_not_show_still_cites_its_source(self):
        result = run_node('''
import { dependencyFindings, requiredCitations } from './skills/professional-slides/runtime/gates/dependency_gates.mjs';
import { measureRegistry } from './skills/professional-slides/runtime/measures.mjs';
const YEARS = ['Y1', 'Y2', 'Y3', 'Y4', 'Y5', 'Y6', 'Y7', 'Y8'];
const graded = { finding: 'a finding', calculation: 'a calculation', soWhat: 'It bears on the decision at hand', strength: 'strong', sources: ['sources/a.csv'] };
const items = [{ ...graded, id: 'i1', shape: 'series', cite: ['accounts', 'tills'], measures: {
  sales: { unit: 'GBP m', population: 'the bakery', periods: YEARS, values: [2.1, 2.2, 2.3, 2.4, 2.5, 2.6, 2.7, 2.8], cite: ['accounts'] },
  loaves: { unit: 'thousand loaves', population: 'the bakery', periods: YEARS, values: [310, 318, 331, 340, 346, 359, 371, 390], cite: ['tills'] },
  staff: { unit: 'employees', population: 'the bakery', periods: YEARS, values: [20, 21, 21, 22, 24, 25, 25, 27] } } }];
const insights = new Map(items.map((i) => [i.id, i])), registry = measureRegistry(insights);
const SOURCES = { accounts: { name: 'Orchard Bakeries annual accounts' }, tills: { name: 'Till records supplied with the brief' } };
const page = (claimed, shown, source) => ({ id: 'p1', type: 'trend', form: 'line', commentary: 'so-what-bar', why: 'The movement over time is the claim, so the series is drawn', evidence: ['i1'],
  title: 'Loaves sold grew faster than sales in every year of the eight', bar: 'Volume is outrunning price, which is what the plan assumes it will keep doing for three more years',
  settles: { kind: 'rate', what: 'growth a year', measures: claimed.map((name) => `i1/${name}`) }, exhibit: { heading: 'Growth', series: shown.map((name) => ({ measure: `i1/${name}`, name })), highlights: [{ category: 'Y8' }] }, ...(source ? { source } : {}) });
const uncited = (p) => dependencyFindings({ deck: { schema: 'professional-slides.deck/v3', id: 'probe', workflow: 'new_deck', request: 'How has the bakery grown?' }, sources: SOURCES, pages: [p] }, insights).filter((f) => f.code === 'SOURCE_UNCITED').map((f) => f.measured);
console.log(JSON.stringify({ needs: requiredCitations(page(['sales', 'loaves'], ['sales']), registry), shownOnly: requiredCitations(page(['sales'], ['sales']), registry),
  // The claim is about sales and loaves, the chart shows sales alone, and the citation names the accounts alone.
  wider: uncited(page(['sales', 'loaves'], ['sales'], ['accounts'])), both: uncited(page(['sales', 'loaves'], ['sales'], ['accounts', 'tills'])),
  // What the page neither shows nor claims - the rest of the insight - is still not asked.
  exact: uncited(page(['sales'], ['sales'], ['accounts'])) }));''')
        self.assertEqual(result["needs"], ["accounts", "tills"])
        self.assertEqual(result["shownOnly"], ["accounts"])
        self.assertEqual(result["wider"], [["tills"]])
        self.assertEqual(result["both"], [])
        self.assertEqual(result["exact"], [])


if __name__ == "__main__":
    unittest.main()
