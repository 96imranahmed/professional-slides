"""Renderer selection and export boundaries, including OSes unavailable to CI."""
import contextlib
import json
import os
import subprocess
import sys
import tempfile
import types
import unittest
from pathlib import Path
from unittest import mock

from node_probe import NODE, RUNTIME, ROOT, requires_python_package, run_node

sys.path.insert(0, str(RUNTIME / "emit"))
import renderer_backends as backends
import render_pptx


class RendererDiscoveryTests(unittest.TestCase):
    def discovery(self, platform, requested="auto", ppt="PowerPoint", lo="soffice"):
        with mock.patch.object(backends, "powerpoint", return_value=ppt), mock.patch.object(backends, "soffice", return_value=lo), \
             mock.patch.object(backends.shutil, "which", side_effect=lambda name: "/bin/" + name):
            return backends.discover(requested, platform)

    def test_installed_powerpoint_is_preferred_on_both_desktop_platforms(self):
        for platform in ("darwin", "win32"):
            report = self.discovery(platform)
            self.assertEqual(report["selected"], "powerpoint")
            self.assertEqual([c["id"] for c in report["available"]], ["powerpoint", "libreoffice"])
            self.assertFalse(report["missing"])

    def test_linux_and_desktops_without_powerpoint_use_libreoffice(self):
        for platform in ("linux", "darwin", "win32"):
            self.assertEqual(self.discovery(platform, ppt=None)["selected"], "libreoffice")
        self.assertIsNone(backends.powerpoint("linux"))

    def test_pinned_renderer_never_selects_another_installed_app(self):
        self.assertEqual(self.discovery("darwin", "libreoffice")["selected"], "libreoffice")
        report = self.discovery("linux", "powerpoint", ppt=None)
        self.assertIsNone(report["selected"])
        self.assertEqual(report["available"], [])
        self.assertIn("powerpoint", report["missing"])

    def test_environment_pin_and_invalid_values(self):
        with mock.patch.dict(os.environ, {backends.RENDERER_ENV: "libreoffice"}):
            self.assertEqual(backends.requested_renderer(), "libreoffice")
            self.assertEqual(backends.requested_renderer("auto"), "auto")
        with mock.patch.dict(os.environ, {backends.RENDERER_ENV: "typo"}):
            with self.assertRaisesRegex(ValueError, "must be one of"):
                backends.discover()

    def test_powerpoint_only_installation_does_not_request_libreoffice(self):
        report = self.discovery("darwin", lo=None)
        self.assertEqual(report["selected"], "powerpoint")
        self.assertNotIn("soffice", report["missing"])
        with mock.patch.object(backends, "powerpoint", return_value="PowerPoint"), mock.patch.object(backends, "soffice", return_value=None), \
             mock.patch.object(backends.shutil, "which", return_value=None):
            self.assertEqual(set(backends.discover("auto", "win32")["missing"]), {"pdftoppm", "pdftotext"})

    def test_windows_com_registration_is_probed_without_launching_office(self):
        registry = types.SimpleNamespace(HKEY_CLASSES_ROOT="HKCR", OpenKey=mock.Mock(return_value=contextlib.nullcontext()))
        with mock.patch.dict(sys.modules, {"winreg": registry}), mock.patch.object(backends, "powershell", return_value="powershell.exe"):
            self.assertEqual(backends.powerpoint("win32"), "powershell.exe")
            registry.OpenKey.assert_called_with("HKCR", r"PowerPoint.Application\CLSID")
            registry.OpenKey.side_effect = OSError("not registered")
            self.assertIsNone(backends.powerpoint("win32"))

    def test_windows_libreoffice_is_found_off_path(self):
        with tempfile.TemporaryDirectory() as tmp, mock.patch.dict(os.environ, {"ProgramFiles": tmp}), \
             mock.patch.object(backends.shutil, "which", return_value=None):
            binary = Path(tmp) / "LibreOffice/program/soffice.exe"
            binary.parent.mkdir(parents=True); binary.touch()
            self.assertEqual(backends.soffice("win32"), str(binary))


class PowerPointCommandTests(unittest.TestCase):
    def setUp(self):
        patch = mock.patch.object(backends, "powerpoint_lock", side_effect=contextlib.nullcontext)
        patch.start(); self.addCleanup(patch.stop)

    def test_paths_are_arguments_and_export_opens_a_copy_on_each_os(self):
        for platform in ("darwin", "win32"):
            with self.subTest(platform=platform), tempfile.TemporaryDirectory() as tmp:
                source = Path(tmp) / "café ' source.pptx"
                source.write_bytes(b"original")
                directory = Path(tmp) / "attempt"
                container = Path(tmp) / "Library/Containers/com.microsoft.Powerpoint/Data"
                container.mkdir(parents=True)
                def export(cmd, **kwargs):
                    self.assertFalse(kwargs.get("shell", False))
                    self.assertNotIn(str(source), cmd)
                    copied = Path(cmd[2]) if platform == "darwin" else Path(cmd[cmd.index("-InputPath") + 1])
                    self.assertEqual(copied.read_bytes(), b"original")
                    pdf = Path(cmd[3]) if platform == "darwin" else Path(cmd[cmd.index("-OutputPath") + 1])
                    if platform == "darwin":
                        self.assertIn(container.resolve(), copied.parents)
                        self.assertIn(container.resolve(), pdf.parents)
                    pdf.write_bytes(b"pdf")
                    return subprocess.CompletedProcess(cmd, 0)
                with mock.patch.object(backends.subprocess, "run", side_effect=export) as run, mock.patch.object(backends.Path, "home", return_value=Path(tmp)):
                    pdf = backends.powerpoint_pdf(source, directory, {"path": "PowerPoint" if platform == "darwin" else "powershell.exe"}, platform)
                self.assertTrue(pdf.is_file())
                self.assertEqual(source.read_bytes(), b"original")
                cmd = run.call_args.args[0]
                self.assertEqual(cmd[0], "osascript" if platform == "darwin" else "powershell.exe")
                if platform == "darwin":
                    self.assertEqual(list((container / "tmp/ProfessionalSlides").iterdir()), [], "scratch exports must be removed")

    def test_failed_export_reports_activation_or_permission_error(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "source.pptx"; source.write_bytes(b"pptx")
            with mock.patch.object(backends.subprocess, "run", side_effect=subprocess.CalledProcessError(1, ["osascript"], stderr=b"denied (-1743)")):
                with self.assertRaisesRegex(RuntimeError, r"denied.*-1743"):
                    backends._powerpoint_pdf(source, Path(tmp) / "out", {"path": "PowerPoint"}, "darwin")


class PowerPointLockTests(unittest.TestCase):
    def test_concurrent_export_fails_promptly_and_the_lock_is_released(self):
        with tempfile.TemporaryDirectory() as tmp, mock.patch.object(backends.tempfile, "gettempdir", return_value=tmp):
            with backends.powerpoint_lock():
                with self.assertRaisesRegex(RuntimeError, "busy with another renderer"):
                    with backends.powerpoint_lock():
                        self.fail("a peer exporter acquired the same lock")
            with backends.powerpoint_lock():
                pass


@requires_python_package("pptx", "pypdf")
class ExportFallbackTests(unittest.TestCase):
    def setUp(self):
        from pptx import Presentation
        self.tmp = tempfile.TemporaryDirectory(); self.addCleanup(self.tmp.cleanup)
        self.directory = Path(self.tmp.name)
        self.source = self.directory / "source.pptx"
        deck = Presentation()
        for _ in range(2):
            deck.slides.add_slide(deck.slide_layouts[6])
        deck.save(str(self.source))
        self.discovery = {"requested": "auto", "available": [{"id": "powerpoint", "path": "PowerPoint"}, {"id": "libreoffice", "path": "soffice"}]}

    def pdf(self, directory, pages=2):
        from pypdf import PdfWriter
        target = directory / "source.pdf"
        writer = PdfWriter()
        for _ in range(pages):
            writer.add_blank_page(width=960, height=540)
        with target.open("wb") as out:
            writer.write(out)
        return target

    def test_auto_fallback_records_the_failed_exporter(self):
        with mock.patch.object(render_pptx, "powerpoint_pdf", side_effect=RuntimeError("not activated")), \
             mock.patch.object(render_pptx, "to_pdf", side_effect=lambda source, td: self.pdf(td)):
            pdf, used, warnings = render_pptx.export_pdf(self.source, self.directory / "exports", self.discovery)
        self.assertTrue(pdf.is_file())
        self.assertEqual(used, "libreoffice")
        self.assertEqual(warnings, [{"renderer": "powerpoint", "message": "not activated"}])

    def test_missing_pages_cannot_pass_as_a_successful_powerpoint_export(self):
        with mock.patch.object(render_pptx, "powerpoint_pdf", side_effect=lambda source, td, backend: self.pdf(td, 1)), \
             mock.patch.object(render_pptx, "to_pdf", side_effect=lambda source, td: self.pdf(td)):
            _, used, warnings = render_pptx.export_pdf(self.source, self.directory / "exports", self.discovery)
        self.assertEqual(used, "libreoffice")
        self.assertIn("1 pages for 2 slides", warnings[0]["message"])

    def test_corrupt_pdf_is_rejected_before_rasterising(self):
        def corrupt(source, td, backend):
            target = td / "bad.pdf"; target.write_bytes(b"not a PDF"); return target
        with mock.patch.object(render_pptx, "powerpoint_pdf", side_effect=corrupt), \
             mock.patch.object(render_pptx, "to_pdf", side_effect=lambda source, td: self.pdf(td)):
            _, used, warnings = render_pptx.export_pdf(self.source, self.directory / "exports", self.discovery)
        self.assertEqual(used, "libreoffice")
        self.assertEqual(warnings[0]["renderer"], "powerpoint")

    def test_explicit_failure_does_not_fall_back(self):
        self.discovery["requested"] = "powerpoint"
        with mock.patch.object(render_pptx, "powerpoint_pdf", side_effect=RuntimeError("denied")), mock.patch.object(render_pptx, "to_pdf") as lo:
            with self.assertRaisesRegex(RuntimeError, "denied"):
                render_pptx.export_pdf(self.source, self.directory / "exports", self.discovery)
            lo.assert_not_called()


@unittest.skipUnless(NODE, "needs Node.js")
class DoctorRendererTests(unittest.TestCase):
    def test_powerpoint_satisfies_the_doctor_without_libreoffice(self):
        report = run_node('''
import { installLines } from './skills/professional-slides/runtime/doctor.mjs';
const yes = { found: true, path: '/bin/tool' }, no = { found: false, path: null };
const base = { platform: 'darwin', node: { ok: true }, python: { chosen: 'python3', candidates: [] }, binaries: { soffice: no, pdftoppm: yes, pdftotext: yes } };
console.log(JSON.stringify(installLines({ ...base, renderer: { selected: 'powerpoint', missing: {} } })));
''')
        self.assertEqual(report, [])

    def test_doctor_and_renderer_share_the_same_probe(self):
        report = run_node('''
import { rendererProbe } from './skills/professional-slides/runtime/doctor.mjs';
console.log(JSON.stringify(rendererProbe({ ...process.env, PROFESSIONAL_SLIDES_RENDERER: 'libreoffice' })));
''')
        result = subprocess.run([os.environ.get("RUNTIME_PYTHON") or sys.executable, str(RUNTIME / "emit/renderer_backends.py"), "--renderer", "libreoffice"], capture_output=True, text=True, check=True)
        self.assertEqual(report, json.loads(result.stdout))
