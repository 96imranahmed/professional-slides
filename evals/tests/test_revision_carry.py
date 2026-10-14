"""A point change to an existing deck is judged on what it changes.

A revision carries the slides it does not change from the source deck
(`carry: true`, runtime/revision.mjs): the build copies them part for part
(emit/assemble_pptx.py) and composes, gates, critiques and reviews only the
pages the revision gave a type or edited in place. These tests run that path
on a deck the skill did not build (evals/point-change/fixtures.py): the
import's two starters, the compile's scope - every page rule on a composed
page, no deck measure while slides are carried, every rule version by
version - the stale-text check, the critique's and the review's scope, the
assembler, and what the built file proves it kept.
"""
import copy
import hashlib
import importlib.util
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from node_probe import HAS_PILLOW, HAS_PPTX, HAS_RENDERER, NODE, ROOT, RUNTIME, RUNTIME_PYTHON, run_node

PYTHON = RUNTIME_PYTHON or sys.executable
IMPORT = RUNTIME / "import-deck.py"
ASSEMBLE = RUNTIME / "emit" / "assemble_pptx.py"
AUTHOR = RUNTIME / "author-deck.mjs"
BENCH = ROOT / "evals" / "point-change"
WEIGHT = json.loads((RUNTIME / "weight.json").read_text())


def load(name: str, file: Path):
    spec = importlib.util.spec_from_file_location(name, file)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def sh(*command, cwd=None):
    return subprocess.run([str(part) for part in command], cwd=cwd, capture_output=True, text=True)


def last_json(text: str):
    at = text.rfind("\n{")
    return json.loads(text[at + 1:] if at >= 0 else text)


# A composed page that passes every page rule in the user's deck: the pie of slide 7 redrawn as a donut.
DONUT = {"id": "s07", "sourceSlide": 7, "type": "composition", "form": "donut", "commentary": "below",
         "title": "Wholesale is now 22% of sales, more than twice delivery",
         "why": "Four channels of one year's sales are parts of a whole; the donut keeps the slide's reading and sets the shares side by side.",
         "settles": {"kind": "share", "what": "Share of FY26 sales by channel"},
         "adds": "The points say which channels carry the stores and why wholesale matters to the plan.",
         "exhibit": {"heading": "Share of FY26 sales, %", "labels": ["In store", "Takeaway", "Delivery", "Wholesale"], "values": [41, 27, 10, 22]},
         "points": ["Wholesale is 22% of sales, up from 16% in FY22, and it is the one channel that grows without a new lease",
                    "In store and takeaway together are 68% of sales, so the estate still carries the business",
                    "Delivery is the smallest channel at 10% and the only one that pays a commission to a platform"]}


@unittest.skipUnless(NODE and HAS_PPTX and HAS_PILLOW, "needs Node.js, python-pptx and Pillow")
class CarriedDeck(unittest.TestCase):
    """The user's deck, built and imported once; each test revises a copy of it."""

    @classmethod
    def setUpClass(cls):
        cls.home = Path(tempfile.mkdtemp(prefix="revision-carry-"))
        cls.source = load("point_change_fixtures", BENCH / "fixtures.py").build_user(cls.home / "source")
        cls.bed = cls.home / "bed"
        done = sh(PYTHON, IMPORT, cls.source, cls.bed, "--carry", "--id", "deck")
        assert done.returncode == 0, done.stderr
        cls.summary = json.loads(done.stdout)
        cls.starter = json.loads((cls.bed / "deck.pages.json").read_text())
        cls.inventory = json.loads((cls.bed / "deck.inventory.json").read_text())

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.home, ignore_errors=True)

    def revise(self, change=None, name=None):
        """A copy of the imported bed with `change(doc)` made to its pages file; returns the folder."""
        work = Path(tempfile.mkdtemp(prefix="rev-", dir=self.home)) if name is None else self.home / name
        shutil.copytree(self.bed, work, dirs_exist_ok=True)
        doc = copy.deepcopy(self.starter)
        if change:
            change(doc)
        (work / "deck.pages.json").write_text(json.dumps(doc, indent=1, ensure_ascii=False))
        return work

    @staticmethod
    def page(doc, key):
        return next(page for page in doc["pages"] if page["id"] == key)

    def author(self, work, *flags):
        self.last_run = sh(NODE, AUTHOR, work / "deck.pages.json", *flags)
        return self.last_run

    def refused(self, work):
        """The findings the last author run refused the deck for, as (code, page id): each printed "  CODE [id]" over its repair."""
        refused = self.last_run.stderr.partition("nothing was written.")[2]
        found = re.findall(r"^ {2}([A-Z][A-Z_]+)(?: \[([^\]]+)\])?\n {4}", refused, re.M)
        return [(code, where.split(" in ")[0].split(", ")[0] if where and "," not in where else None) for code, where in found]


class ImportTests(CarriedDeck):
    def test_the_carry_starter_keeps_every_slide_and_records_what_it_was_made_under(self):
        deck, pages = self.starter["deck"], self.starter["pages"]
        self.assertEqual(deck, {"schema": "professional-slides.deck/v3", "id": "deck", "workflow": "existing_deck_revision", "inventory": "deck.inventory.json",
                                "rulesVersion": WEIGHT["rulesVersion"], "request": ""})
        # Every slide is a page, the title slide among them: nothing of the user's deck is the runtime's to redraw until the author says so.
        self.assertEqual([(p["id"], p["sourceSlide"], p["carry"]) for p in pages], [(f"s0{n}", n, True) for n in range(1, 9)])
        self.assertTrue(all("draft" in p and "type" not in p for p in pages))
        self.assertEqual(pages[0]["title"], "Harbour Coffee: FY26 review")

    def test_the_source_deck_is_kept_beside_the_inventory_byte_for_byte(self):
        copy_path = self.bed / "deck.source.pptx"
        self.assertEqual(copy_path.read_bytes(), self.source.read_bytes())
        self.assertEqual(self.inventory["source"], {"file": self.source.name, "sha256": hashlib.sha256(self.source.read_bytes()).hexdigest(), "copy": "deck.source.pptx"})
        self.assertEqual(self.summary["source"], str(copy_path))

    def test_the_import_says_what_the_author_must_decide(self):
        # This deck is the size the runtime composes at and draws nothing the inventory cannot hold.
        self.assertEqual(self.summary["decide"], [])
        self.assertIn("carry: true", self.summary["starter"])
        # A deck of another size says so at the import: a composed page cannot be set beside its slides.
        other = self.home / "four-three"
        other.mkdir()
        made = sh(PYTHON, "-c", "import sys\nfrom pptx import Presentation\nfrom pptx.util import Inches\nprs = Presentation()\ns = prs.slides.add_slide(prs.slide_layouts[5])\ns.shapes.title.text = 'A slide'\n"
                  "s.shapes.add_shape(1, Inches(1), Inches(2), Inches(2), Inches(1))\ns.shapes.add_connector(1, Inches(1), Inches(4), Inches(3), Inches(4))\nprs.save(sys.argv[1])", other / "old.pptx")
        self.assertEqual(made.returncode, 0, made.stderr)
        summary = json.loads(sh(PYTHON, IMPORT, other / "old.pptx", other / "out", "--carry").stdout)
        self.assertEqual(len(summary["decide"]), 2)
        self.assertIn("960 x 720", summary["decide"][0])
        self.assertIn("1280 x 720", summary["decide"][0])
        # ...and so does a slide that draws what the inventory does not hold: kept as it is, or redrawn by decision.
        self.assertIn("s01 (2 shapes drawn without text)", summary["decide"][1])
        inventory = json.loads((other / "out" / "old.inventory.json").read_text())
        self.assertEqual(inventory["slides"][0]["unheld"], {"drawn": 2})

    def test_the_inventory_marks_page_furniture_and_the_draft_leaves_it_out(self):
        roles = {(p["text"], p.get("role")) for slide in self.inventory["slides"] for p in slide["paragraphs"]}
        self.assertIn(("Harbour Coffee | Board pre-read | Confidential", "furniture"), roles)
        self.assertIn(("4", "page-number"), roles)
        self.assertIn(("Source: management accounts, FY22 to FY26", "source"), roles)
        self.assertEqual(self.inventory["furniture"], ["Harbour Coffee | Board pre-read | Confidential"])
        draft = self.page(self.starter, "s04")["draft"]["text"]
        self.assertEqual(draft, ["London has the fastest growth and the thinnest margin: rents took 31% of its revenue in FY26"])

    def test_the_rebuild_starter_still_maps_every_slide_and_takes_the_cover(self):
        out = self.home / "rebuild"
        self.assertEqual(sh(PYTHON, IMPORT, self.source, out, "--id", "deck").returncode, 0)
        doc = json.loads((out / "deck.pages.json").read_text())
        self.assertEqual(doc["deck"]["cover"], {"title": "Harbour Coffee: FY26 review", "subtitle": "Board pre-read, March 2026"})
        self.assertEqual(doc["deck"]["footer"], "Harbour Coffee | Board pre-read | Confidential")
        self.assertEqual(doc["deck"]["rulesVersion"], WEIGHT["rulesVersion"])
        self.assertEqual([p["id"] for p in doc["pages"]], [f"s0{n}" for n in range(2, 9)])
        self.assertTrue(all("carry" not in p for p in doc["pages"]))
        # An unmapped slide is still refused - and the refusal says it can be carried instead.
        full = sh(NODE, AUTHOR, out / "deck.pages.json")
        self.assertEqual(full.returncode, 2)
        self.assertIn("REVISION_UNMAPPED", full.stderr)
        self.assertIn("`carry: true`", full.stderr)


class CompileScopeTests(CarriedDeck):
    def test_the_starter_compiles_as_written_in_every_mode(self):
        work = self.revise()
        for flags in (["--draft"], ["--check"], []):
            done = self.author(work, *flags)
            self.assertEqual(done.returncode, 0, f"{flags}: {done.stderr[:800]}")
        spec = json.loads((work / "deck.deck.json").read_text())
        self.assertEqual(spec["slides"], [])
        self.assertEqual([(c["id"], c["sourceSlide"], c["after"]) for c in spec["carried"]], [(f"s0{n}", n, None) for n in range(1, 9)])
        self.assertEqual(spec["rulesVersion"], WEIGHT["rulesVersion"])
        summary = last_json(done.stdout)
        self.assertEqual(summary["revision"]["carried"], 8)
        self.assertEqual(summary["revision"]["composed"], [])
        self.assertIn("8 carried as they are", done.stderr)

    def test_a_carried_page_is_the_slide_and_at_most_a_text_edit(self):
        cases = {
            "points": lambda p: p.update(points=["A point"]),
            "type": lambda p: p.update(type="summary"),
            "replace": lambda p: p.update(replace=[{"old": "", "new": "x"}]),
            "same": lambda p: p.update(replace=[{"old": "Summary", "new": "Summary"}]),
            "slide": lambda p: p.pop("sourceSlide"),
        }
        for name, change in cases.items():
            work = self.revise(lambda doc: change(self.page(doc, "s02")))
            done = self.author(work, "--check")
            self.assertEqual(done.returncode, 2, name)
            self.assertEqual(self.refused(work), [("COMPILE", "s02")], name)
        # Outside a revision there is no slide to carry.
        work = self.revise(lambda doc: doc["deck"].update(workflow="new_deck", request="Make the deck."))
        self.assertIn("only a revision", self.author(work, "--check").stderr)

    def test_an_edit_the_slide_cannot_take_is_refused_at_the_compile(self):
        def edits(doc):
            # 9.7 is a value of slide 3's chart, not words the slide prints.
            self.page(doc, "s03")["replace"] = [{"old": "9.7", "new": "9.9"}]
            self.page(doc, "s05")["replace"] = [{"old": "It took £41k a week", "new": "It took £44k a week"}, {"old": "no such words", "new": "x"}]
        work = self.revise(edits)
        done = self.author(work, "--check")
        self.assertEqual(done.returncode, 2)
        self.assertEqual(self.refused(work), [("REVISION_CARRY_INVALID", "s03"), ("REVISION_CARRY_INVALID", "s05")])
        self.assertIn('slide 3 does not print "9.7"', done.stderr)
        self.assertIn("giving the page a `type`", done.stderr)

    def test_a_slide_is_carried_once_and_not_also_redrawn(self):
        def twice(doc):
            doc["pages"].append({**self.page(doc, "s02"), "id": "s02b"})
            doc["pages"].append({**copy.deepcopy(DONUT), "id": "n01"})
            doc["deck"]["density"] = "live-pitch"
        work = self.revise(twice)
        self.assertEqual(self.author(work, "--check").returncode, 2)
        found = self.refused(work)
        self.assertIn(("REVISION_CARRY_INVALID", "s02b"), found)
        # Slide 7 is carried as s07 and redrawn as n01.
        self.assertIn(("REVISION_CARRY_INVALID", "s07"), found)

    def test_the_source_deck_must_be_the_one_imported(self):
        work = self.revise()
        (work / "deck.source.pptx").unlink()
        done = self.author(work, "--check")
        self.assertEqual(done.returncode, 2)
        self.assertEqual(self.refused(work), [("REVISION_SOURCE_MISSING", None)])
        (work / "deck.source.pptx").write_bytes(b"not the deck")
        self.assertEqual(self.author(work, "--check").returncode, 2)
        self.assertIn("its hash differs", self.author(work, "--check").stderr)

    def test_a_composed_page_needs_slides_of_the_size_it_is_composed_at(self):
        def swap(doc):
            doc["deck"]["density"] = "live-pitch"
            doc["pages"][doc["pages"].index(self.page(doc, "s07"))] = copy.deepcopy(DONUT)
        work = self.revise(swap)
        inventory = json.loads((work / "deck.inventory.json").read_text())
        inventory["slideSize"] = {"width": 960, "height": 720}
        (work / "deck.inventory.json").write_text(json.dumps(inventory))
        done = self.author(work, "--check")
        self.assertEqual(done.returncode, 2)
        self.assertIn(("REVISION_SLIDE_SIZE", None), self.refused(work))
        # Text edits on carried slides need no composed page, so they keep any size.
        work = self.revise(lambda doc: self.page(doc, "s04").update(title="London grows fastest on the thinnest margin"))
        (work / "deck.inventory.json").write_text(json.dumps(inventory))
        self.assertEqual(self.author(work, "--check").returncode, 0)


class StaleTextTests(CarriedDeck):
    def test_replaced_words_left_on_another_slide_block_and_the_bare_number_advises(self):
        work = self.revise(lambda doc: self.page(doc, "s02").update(replace=[{"old": "£12.4m", "new": "£12.6m"}]))
        done = self.author(work, "--check")
        self.assertEqual(done.returncode, 2)
        self.assertEqual(self.refused(work), [("NUMBER_STALE", "s02")])
        # The closing slide prints the same figure; the table's total shows its number.
        self.assertIn('s08 still prints "£12.4m"', done.stderr)

        def everywhere(doc):
            for key in ("s02", "s08"):
                self.page(doc, key)["replace"] = [{"old": "£12.4m", "new": "£12.6m"}]
        # The table's total is the same quantity in another notation - 12.4 under "Revenue (£m)" - and blocks as the words do.
        work = self.revise(everywhere)
        done = self.author(work, "--check")
        self.assertEqual(done.returncode, 2, done.stderr[:800])
        self.assertEqual(set(self.refused(work)), {("NUMBER_STALE", "s02"), ("NUMBER_STALE", "s08")})
        self.assertIn('s04 still states 12.4 ("Total" under "Revenue (£m)")', done.stderr)

        def all_three(doc):
            everywhere(doc)
            self.page(doc, "s04")["replace"] = [{"old": "12.4", "new": "12.6"}]
        work = self.revise(all_three)
        done = self.author(work, "--check")
        self.assertEqual(done.returncode, 0)
        self.assertFalse(any("STALE" in line for line in last_json(done.stdout)["advisories"]))

    def test_an_edit_names_the_pages_whose_same_words_are_another_thing_and_each_is_listed(self):
        edit = lambda only: (lambda doc: self.page(doc, "s02").update(replace=[{"old": "£12.4m", "new": "£12.6m", "only": only}]))
        # All-or-nothing is not a statement anyone can check: `only` names pages.
        done = self.author(self.revise(edit(True)), "--check")
        self.assertEqual(done.returncode, 2)
        self.assertIn('`"only": ["s05", ...]` naming the other pages', done.stderr)
        # A page it does not name is still read.
        work = self.revise(edit(["s08"]))
        done = self.author(work, "--check")
        self.assertEqual(done.returncode, 2)
        self.assertEqual(self.refused(work), [("NUMBER_STALE", "s02")])
        self.assertIn("s04 still states 12.4", done.stderr)
        self.assertNotIn("s08 still prints", done.stderr)
        # A page that is not one of the deck's excuses nothing, and is said.
        work = self.revise(edit(["s04", "s08", "s99"]))
        self.assertEqual(self.author(work, "--check").returncode, 2)
        self.assertEqual(self.refused(work), [("REVISION_CARRY_INVALID", "s02")])
        # Every page named, the edit stands - and the run and the compiled deck's statement list each for the reviewer.
        work = self.revise(edit(["s04", "s08"]))
        done = self.author(work, "--check")
        self.assertEqual(done.returncode, 0, done.stderr[:800])
        self.assertFalse(any("STALE" in line for line in last_json(done.stdout)["advisories"]))
        self.assertEqual(last_json(done.stdout)["revision"]["excused"], [{"page": "s02", "old": "£12.4m", "only": ["s04", "s08"]}])
        self.assertIn('Left standing on purpose (`only`), for the reviewer to check: s02 "£12.4m" not looked for on s04, s08', done.stderr)

    def test_a_number_is_not_found_inside_a_longer_one(self):
        # 2.7 is wholesale revenue on slide 3 and London's on slide 4; 12.7 and 2.75 would not be it.
        result = run_node('''
import { consistencyFindings } from "./skills/professional-slides/runtime/gates/consistency_gates.mjs";
const inventory = { source: {}, slideSize: { width: 1280, height: 720 }, slides: [
  { index: 1, title: "One", paragraphs: [{ text: "Margin was 2.7% in the year" }], tables: [], charts: [] },
  { index: 2, title: "Two", paragraphs: [{ text: "It rose to 12.7% and then 2.75%" }], tables: [], charts: [] },
  { index: 3, title: "Three", paragraphs: [{ text: "A margin of 2.7% again" }], tables: [], charts: [{ series: [{ values: [2.7] }] }] },
  { index: 4, title: "Four", paragraphs: [], tables: [], charts: [{ series: [{ values: [2.7, 3] }] }] }] };
const spec = { workflow: "existing_deck_revision", inventory: "none.inventory.json", slides: [], carried: [
  { id: "s1", sourceSlide: 1, title: "One", after: null, replace: [{ old: "2.7%", new: "2.9%" }] }, { id: "s2", sourceSlide: 2, title: "Two", after: null },
  { id: "s3", sourceSlide: 3, title: "Three", after: null }, { id: "s4", sourceSlide: 4, title: "Four", after: null }] };
const doc = { deck: { workflow: spec.workflow }, pages: spec.carried.map((entry) => ({ id: entry.id, sourceSlide: entry.sourceSlide, title: entry.title, carry: true, ...(entry.replace ? { replace: entry.replace } : {}) })) };
console.log(JSON.stringify(consistencyFindings(doc, null, { spec, inventory, changed: ["s1"] }).filter((f) => f.code === "NUMBER_STALE").map((f) => ({ severity: f.severity, pages: f.pages }))));
''')
        self.assertEqual(result, [{"severity": "blocker", "pages": ["s1", "s3"]}, {"severity": "advisory", "pages": ["s1", "s4"]}])

    def test_no_recorded_version_excuses_what_a_revision_leaves_behind_and_a_lowered_one_says_why(self):
        def older(reason=None):
            def change(doc):
                doc["deck"]["rulesVersion"] = 5
                if reason:
                    doc["deck"]["rulesVersionReason"] = reason
                self.page(doc, "s02")["replace"] = [{"old": "£12.4m", "new": "£12.6m"}]
            return change
        # The import stamped the current version: a pages file that records a lower one is asking to be held to less, and says why.
        self.assertEqual(self.inventory["rulesVersion"], WEIGHT["rulesVersion"])
        work = self.revise(older())
        done = self.author(work, "--check")
        self.assertEqual(done.returncode, 2)
        self.assertEqual(set(self.refused(work)), {("NUMBER_STALE", "s02"), ("REVISION_RULES_VERSION", None)})
        self.assertIn('`"rulesVersionReason": "..."`', done.stderr)
        # With its reason the version stands - and the stale rule is no version's: it judges only what this revision changed.
        work = self.revise(older("The deck was authored and accepted under version five last quarter"))
        done = self.author(work, "--check")
        self.assertEqual(done.returncode, 2)
        self.assertEqual(self.refused(work), [("NUMBER_STALE", "s02")])
        self.assertIn(f"Held to rules version 5, below the {WEIGHT['rulesVersion']} its import stamped", done.stderr)

        def everywhere(doc):
            older("The deck was authored and accepted under version five last quarter")(doc)
            self.page(doc, "s08")["replace"] = [{"old": "£12.4m", "new": "£12.6m"}]
            self.page(doc, "s04")["replace"] = [{"old": "12.4", "new": "12.6"}]
        work = self.revise(everywhere)
        self.assertEqual(self.author(work).returncode, 0)
        # The version the deck records is kept: the compile does not stamp today's over it.
        self.assertEqual(json.loads((work / "deck.deck.json").read_text())["rulesVersion"], 5)


@unittest.skipUnless(NODE, "needs Node.js")
class RevisionChangesTests(unittest.TestCase):
    PROBE = '''
import { revisionChanges } from "./skills/professional-slides/runtime/revision.mjs";
const slide = (index, title, text) => ({ index, id: `s0${index}`, title, subtitle: null, hidden: false, paragraphs: [{ text, level: 0 }], tables: [], charts: [], pictures: [] });
const inventory = { slides: [slide(1, "Cover", "A deck"), slide(2, "Summary", "Revenue reached 12.4"), slide(3, "Revenue by year", "It doubled"), slide(4, "Regions", "London is thin"), slide(5, "Next steps", "Approve the plan")] };
const carry = (n, more = {}) => ({ id: `s0${n}`, sourceSlide: n, title: inventory.slides[n - 1].title, after: null, ...more });
const typed = (id, title, sourceSlide) => ({ id, title, points: ["It doubled"], pageType: { type: "summary", form: "takeaways", ...(sourceSlide ? { sourceSlide } : {}), content: { claim: title } } });
const deck = (carried, slides = []) => ({ workflow: "existing_deck_revision", inventory: "x", slides, carried });
'''

    def changes(self, body):
        return run_node(self.PROBE + body)

    def test_a_deck_carried_whole_changed_nothing(self):
        result = self.changes('console.log(JSON.stringify(revisionChanges(deck([1, 2, 3, 4, 5].map((n) => carry(n))), inventory)));')
        self.assertEqual(result["spine"], [])
        self.assertEqual(result["content"], [])
        self.assertFalse(result["spineChanged"])
        self.assertFalse(result["restyle"])
        self.assertEqual(result["carried"], ["s01", "s02", "s03", "s04", "s05"])
        self.assertEqual(result["pages"], ["s01", "s02", "s03", "s04", "s05"])

    def test_a_new_title_changes_the_spine_and_replaced_words_the_copy_alone(self):
        result = self.changes('''
const retitled = revisionChanges(deck([carry(1), carry(2), carry(3), carry(4, { title: "London grows fastest on the thinnest margin" }), carry(5)]), inventory);
const reworded = revisionChanges(deck([carry(1), carry(2, { replace: [{ old: "12.4", new: "12.6" }] }), carry(3), carry(4), carry(5)]), inventory);
console.log(JSON.stringify({ retitled, reworded }));''')
        self.assertEqual(result["retitled"]["spine"], ["s04"])
        self.assertEqual(result["retitled"]["content"], ["s04"])
        self.assertEqual(result["reworded"]["spine"], [])
        self.assertFalse(result["reworded"]["spineChanged"])
        self.assertEqual(result["reworded"]["content"], ["s02"])
        self.assertEqual(result["reworded"]["carried"], ["s01", "s03", "s04", "s05"])

    def test_a_composed_page_is_drawn_anew_beside_the_users_slides_and_always_read(self):
        result = self.changes('''
// Slide 3 redrawn under its own title, and a page added after it.
const spec = deck([carry(1), carry(2), { ...carry(4), after: "n01" }, { ...carry(5), after: "n01" }], [typed("s03", "Revenue by year", 3), typed("n01", "Wholesale carries the growth")]);
console.log(JSON.stringify(revisionChanges(spec, inventory)));''')
        self.assertEqual(result["pages"], ["s01", "s02", "s03", "n01", "s04", "s05"])
        # The redrawn page keeps its title, so the spine holds there; the added page changes it.
        self.assertEqual(result["spine"], ["n01"])
        self.assertEqual(result["drawn"], ["s03", "n01"])
        self.assertEqual(result["content"], ["s03", "n01"])
        self.assertEqual(result["carried"], ["s01", "s02", "s04", "s05"])

    def test_a_dropped_slide_and_a_moved_one_change_the_spine_where_they_were(self):
        result = self.changes('''
const dropped = revisionChanges(deck([carry(1), carry(2), carry(4), carry(5)]), inventory);
const moved = revisionChanges(deck([carry(1), carry(3), carry(2), carry(4), carry(5)]), inventory);
console.log(JSON.stringify({ dropped, moved }));''')
        self.assertEqual(result["dropped"]["dropped"], [3])
        self.assertEqual(result["dropped"]["spine"], ["s02", "s04"])
        self.assertTrue(result["dropped"]["spineChanged"])
        self.assertEqual(result["moved"]["spine"], ["s02"])

    def test_a_slide_hidden_or_shown_again_is_the_revisions_to_review(self):
        # The assembler hides or shows a carried slide (`hidden` in its plan) and the emitter a composed one; each is read
        # as its slide by its title, words and drawing. Read as unchanged, the review was told nothing changed on it and
        # refused any finding there.
        result = self.changes('''
import { stampKept } from "./skills/professional-slides/runtime/revision.mjs";
const backup = { slides: inventory.slides.map((s) => (s.index === 4 ? { ...s, hidden: true } : s)) };
const said = (n, more = {}) => ({ id: `s0${n}`, title: inventory.slides[n - 1].title, points: [inventory.slides[n - 1].paragraphs[0].text],
  pageType: { type: "summary", form: "takeaways", sourceSlide: n, content: { claim: inventory.slides[n - 1].title } }, ...more });
const rebuilt = (four) => stampKept(deck([], [1, 2, 3, 4, 5].map((n) => (n === 4 ? said(4, four) : said(n)))), backup);
const of = (spec) => { const c = revisionChanges(spec, backup); return { content: c.content, shown: c.shown, carried: c.carried ?? null, spineChanged: c.spineChanged }; };
console.log(JSON.stringify({
  hid: of(deck([carry(1), carry(2), carry(3, { hidden: true }), carry(4), carry(5)])),
  kept: of(deck([carry(1), carry(2), carry(3, { hidden: false }), carry(4, { hidden: true }), carry(5)])),
  shownAgain: of(deck([carry(1), carry(2), carry(3), carry(4, { hidden: false }), carry(5)])),
  // A page composed where the hidden slide stood: left without the flag it stays hidden, as the carried slide would.
  composedKept: { hidden: rebuilt({}).slides[3].hidden ?? null, ...of(rebuilt({})) },
  composedShown: { hidden: rebuilt({ hidden: false }).slides[3].hidden ?? null, ...of(rebuilt({ hidden: false })) },
  composedHid: of(stampKept(deck([], [1, 2, 3, 4, 5].map((n) => said(n, n === 2 || n === 4 ? { hidden: true } : {}))), backup)) }));''')
        self.assertEqual(result["hid"], {"content": ["s03"], "shown": ["s03"], "carried": ["s01", "s02", "s04", "s05"], "spineChanged": False})
        self.assertEqual(result["kept"], {"content": [], "shown": [], "carried": ["s01", "s02", "s03", "s04", "s05"], "spineChanged": False})
        self.assertEqual(result["shownAgain"], {"content": ["s04"], "shown": ["s04"], "carried": ["s01", "s02", "s03", "s05"], "spineChanged": False})
        self.assertEqual(result["composedKept"], {"hidden": True, "content": [], "shown": [], "carried": None, "spineChanged": False})
        self.assertEqual(result["composedShown"], {"hidden": False, "content": ["s04"], "shown": ["s04"], "carried": None, "spineChanged": False})
        self.assertEqual(result["composedHid"]["shown"], ["s02"])
        self.assertEqual(result["composedHid"]["content"], ["s02"])

    def test_a_picture_is_its_slides_by_its_path_not_its_name(self):
        # Every imported deck names its pictures s01-1.png, s01-2.png, ...: another deck's picture of the same name is
        # another picture, and the slide's own, written from a folder above the inventory's, is the same one.
        result = self.changes('''
const pictured = { slides: inventory.slides.map((s) => (s.index === 3 ? { ...s, pictures: [{ file: "assets/board/s03-1.png" }] } : s)) };
const drawn = (file) => revisionChanges(deck([], [{ ...typed("s03", "Revenue by year", 3), pictures: [{ path: file, alt: "The chart" }] }]), pictured).drawn;
console.log(JSON.stringify({ own: drawn("assets/board/s03-1.png"), dotted: drawn("./assets/board/s03-1.png"), above: drawn("../work/assets/board/s03-1.png"),
  windows: drawn("work\\\\assets\\\\board\\\\s03-1.png"), otherDeck: drawn("assets/budget/s03-1.png"), other: drawn("assets/board/s03-2.png") }));''')
        self.assertEqual({name: drawn for name, drawn in result.items()},
                         {"own": [], "dotted": [], "above": [], "windows": [], "otherDeck": ["s03"], "other": ["s03"]})

class SwappedDeck(CarriedDeck):
    """The user's deck with slide 7 redrawn as a composed page, compiled and built without a render."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        # One composed page, emitted as the build emits it, to set among the user's slides.
        cls.swapped = cls.home / "swapped"
        shutil.copytree(cls.bed, cls.swapped)
        doc = copy.deepcopy(cls.starter)
        doc["deck"]["density"] = "live-pitch"
        doc["pages"][6] = copy.deepcopy(DONUT)
        (cls.swapped / "deck.pages.json").write_text(json.dumps(doc, ensure_ascii=False))
        authored = sh(NODE, AUTHOR, cls.swapped / "deck.pages.json")
        # The exit code is said with the output: a compile killed from outside (a negative code) prints nothing of its own.
        assert authored.returncode == 0, f"exit {authored.returncode}\n" + authored.stdout[-3000:] + authored.stderr[-3000:]
        built = sh(NODE, RUNTIME / "build-deck.mjs", cls.swapped / "deck.deck.json", cls.swapped / "out", "--no-render")
        assert built.returncode == 0, built.stderr[:2000] + built.stdout[:2000]
        cls.build = json.loads((cls.swapped / "out" / "build-result.json").read_text())


class AssemblerTests(SwappedDeck):
    def assemble(self, slides, composed=None, check=False):
        work = Path(tempfile.mkdtemp(prefix="asm-", dir=self.home))
        (work / "plan.json").write_text(json.dumps({"source": str(self.bed / "deck.source.pptx"), "sha256": self.inventory["source"]["sha256"], "composed": composed, "slides": slides}))
        done = sh(PYTHON, ASSEMBLE, work / "plan.json", *(["--check"] if check else [work / "out.pptx"]))
        return done, work / "out.pptx"

    @staticmethod
    def parts(path):
        with zipfile.ZipFile(path) as archive:
            return {name: archive.read(name) for name in archive.namelist()}

    def test_a_deck_carried_whole_is_the_source_deck_part_for_part(self):
        done, out = self.assemble([{"carry": n} for n in range(1, 9)])
        self.assertEqual(done.returncode, 0, done.stderr)
        report = json.loads(done.stdout)
        self.assertEqual((report["carried"], report["edited"], report["composed"], report["dropped"]), (8, 0, 0, []))
        self.assertEqual(report["preserved"]["slides"], 8)
        self.assertEqual(report["preserved"]["drifted"], [])
        before, after = self.parts(self.source), self.parts(out)
        # Nothing changed, so nothing is written again: every part is the user's bytes, the list of slides among them.
        self.assertEqual(after, before)
        self.assertEqual(report["preserved"]["sourcePartsIdentical"], report["preserved"]["sourceParts"])
        # The same plan gives the same file.
        again, out2 = self.assemble([{"carry": n} for n in range(1, 9)])
        self.assertEqual(out.read_bytes(), out2.read_bytes())

    def test_text_is_replaced_in_place_and_nothing_else_of_the_slide_moves(self):
        from pptx import Presentation
        done, out = self.assemble([{"carry": 1}, {"carry": 2, "edits": [{"old": "£12.4m", "new": "£12.6m"}]}, {"carry": 3},
                                   {"carry": 4, "title": "London grows fastest on the thinnest margin", "edits": [{"old": "12.4", "new": "12.6"}]},
                                   {"carry": 5, "hidden": True}, {"carry": 6}, {"carry": 7}, {"carry": 8}])
        self.assertEqual(done.returncode, 0, done.stderr)
        report = json.loads(done.stdout)
        self.assertEqual(report["edits"][0], {"slide": 2, "changes": [{"old": "£12.4m", "new": "£12.6m", "count": 1}]})
        self.assertEqual((report["carried"], report["edited"]), (5, 3))
        before, after = Presentation(str(self.source)), Presentation(str(out))
        frames = lambda slide: [(s.shape_type, s.left, s.top, s.width, s.height) for s in slide.shapes]
        for a, b in zip(before.slides, after.slides):
            self.assertEqual(frames(a), frames(b))
        self.assertEqual(after.slides[1].placeholders[1].text_frame.paragraphs[0].text, "Revenue reached £12.6m in FY26, up 14% on FY25")
        self.assertEqual(after.slides[3].shapes.title.text, "London grows fastest on the thinnest margin")
        table = next(s for s in after.slides[3].shapes if s.has_table).table
        self.assertEqual([cell.text for cell in table.rows[5].cells], ["Total", "42", "12.6", "7.9%", "13.4%"])
        self.assertEqual(after.slides[4]._element.get("show"), "0")
        # The edited run keeps its own formatting: the size the deck set on it.
        self.assertEqual(after.slides[1].placeholders[1].text_frame.paragraphs[0].runs[0].font.size, before.slides[1].placeholders[1].text_frame.paragraphs[0].runs[0].font.size)

    def test_words_that_run_across_runs_and_a_soft_break_are_replaced_as_the_reader_reads_them(self):
        module = load("assemble_pptx", ASSEMBLE)
        from lxml import etree
        a = "http://schemas.openxmlformats.org/drawingml/2006/main"
        paragraph = etree.fromstring(f'<a:p xmlns:a="{a}"><a:r><a:rPr b="1"/><a:t>Revenue reached £12</a:t></a:r><a:r><a:t>.4m in</a:t></a:r><a:br/><a:r><a:t>FY26 and £12.4m again</a:t></a:r></a:p>')
        self.assertEqual(module.replace_in_paragraph(paragraph, "£12.4m", "£12.6m"), 2)
        texts = [t.text for t in paragraph.iter(f"{{{a}}}t")]
        self.assertEqual(texts, ["Revenue reached £12.6m", " in", "FY26 and £12.6m again"])
        # Words either side of a soft break read with a space between them, and the break goes with the words replaced.
        self.assertEqual(module.replace_in_paragraph(paragraph, "in FY26", "for FY26"), 1)
        self.assertEqual("".join(t.text for t in paragraph.iter(f"{{{a}}}t")), "Revenue reached £12.6m for FY26 and £12.6m again")
        self.assertEqual(len(paragraph.findall(f"{{{a}}}br")), 0)

    def test_every_edit_that_cannot_be_made_is_listed_before_anything_is_written(self):
        done, out = self.assemble([{"carry": 3, "edits": [{"old": "9.7", "new": "9.9"}, {"old": "Revenue has doubled", "new": "Revenue doubled"}, {"old": "absent", "new": "x"}]}, {"carry": 12}], check=True)
        self.assertEqual(done.returncode, 2)
        refusals = json.loads(done.stdout)["refusals"]
        self.assertEqual([r["slide"] for r in refusals], [3, 3, 12])
        self.assertIn('does not print "9.7"', refusals[0]["message"])
        done, out = self.assemble([{"carry": 3, "edits": [{"old": "9.7", "new": "9.9"}]}])
        self.assertEqual(done.returncode, 2)
        self.assertFalse(out.exists())
        # A source deck that is not the one imported is refused, by its hash.
        work = Path(tempfile.mkdtemp(prefix="asm-", dir=self.home))
        (work / "plan.json").write_text(json.dumps({"source": str(self.source), "sha256": "0" * 64, "composed": None, "slides": [{"carry": 1}]}))
        self.assertIn("its hash differs", sh(PYTHON, ASSEMBLE, work / "plan.json", work / "out.pptx").stderr)

    def test_a_dropped_slide_leaves_with_its_notes_and_a_composed_one_brings_its_own_master(self):
        from pptx import Presentation
        composed = self.build["revision"]["composedPptx"]
        done, out = self.assemble([{"carry": 1}, {"carry": 3}, {"composed": 1}, {"carry": 8}], composed=composed)
        self.assertEqual(done.returncode, 0, done.stderr)
        report = json.loads(done.stdout)
        self.assertEqual(report["dropped"], [2, 4, 5, 6, 7])
        before, after = self.parts(self.source), self.parts(out)
        # Of the parts both files hold, only the three that list the slides were written again.
        self.assertEqual({name for name in before if name in after and after[name] != before[name]}, {"ppt/presentation.xml", "ppt/_rels/presentation.xml.rels", "[Content_Types].xml"})
        # Slide 2 and its notes page are gone; the notes master the deck keeps is still there.
        self.assertNotIn("ppt/slides/slide2.xml", after)
        self.assertNotIn("ppt/notesSlides/notesSlide1.xml", after)
        self.assertIn("ppt/notesMasters/notesMaster1.xml", after)
        deck = Presentation(str(out))
        self.assertEqual([slide.shapes.title.text for slide in deck.slides], ["Harbour Coffee: FY26 review", "Revenue by year", DONUT["title"], "Next steps"])
        # The composed slide is drawn on its own master and theme, added beside the user's: none of theirs is redefined.
        self.assertEqual(len(deck.slide_masters), 2)
        self.assertEqual(after["ppt/slideMasters/slideMaster1.xml"], before["ppt/slideMasters/slideMaster1.xml"])
        self.assertEqual(after["ppt/theme/theme1.xml"], before["ppt/theme/theme1.xml"])
        from lxml import etree
        ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
        presentation = etree.fromstring(after["ppt/presentation.xml"])
        master_ids = [int(el.get("id")) for el in presentation.findall("p:sldMasterIdLst/p:sldMasterId", ns)]
        layout_ids = [int(el.get("id")) for name, data in after.items() if name.startswith("ppt/slideMasters/") and name.endswith(".xml") for el in etree.fromstring(data).iter("{%s}sldLayoutId" % ns["p"])]
        self.assertEqual(len(set(master_ids + layout_ids)), len(master_ids + layout_ids))
        slide_ids = [int(el.get("id")) for el in presentation.findall("p:sldIdLst/p:sldId", ns)]
        self.assertEqual(len(set(slide_ids)), 4)

    P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
    PML = "http://schemas.openxmlformats.org/presentationml/2006/main"

    def sectioned(self, slides):
        """The source deck filed in two sections as PowerPoint writes them - "Opening" (slides 1-3) and "Body" (the rest) -
        assembled by `slides`: the code, and the slide ids of the written deck in order and section by section."""
        from lxml import etree
        parts = self.parts(self.source)
        root = etree.fromstring(parts["ppt/presentation.xml"])
        ids = [el.get("id") for el in root.find(f"{{{self.PML}}}sldIdLst")]
        sections = "".join(f'<p14:section name="{name}" id="{{0000000{n}-0000-0000-0000-000000000000}}"><p14:sldIdLst>{"".join(f"""<p14:sldId id="{i}"/>""" for i in group)}</p14:sldIdLst></p14:section>'
                           for n, (name, group) in enumerate((("Opening", ids[:3]), ("Body", ids[3:]))))
        root.append(etree.fromstring(f'<p:extLst xmlns:p="{self.PML}" xmlns:p14="{self.P14}"><p:ext uri="{{521415D9-36F7-43E2-AB2F-B90AF26B5E84}}"><p14:sectionLst>{sections}</p14:sectionLst></p:ext></p:extLst>'))
        parts["ppt/presentation.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        work = Path(tempfile.mkdtemp(prefix="asm-", dir=self.home))
        with zipfile.ZipFile(work / "sectioned.pptx", "w", zipfile.ZIP_DEFLATED) as archive:
            for name, data in parts.items():
                archive.writestr(name, data)
        (work / "plan.json").write_text(json.dumps({"source": str(work / "sectioned.pptx"), "composed": self.build["revision"]["composedPptx"], "slides": slides}))
        done = sh(PYTHON, ASSEMBLE, work / "plan.json", work / "out.pptx")
        if done.returncode != 0:
            return done, None, None
        written = etree.fromstring(self.parts(work / "out.pptx")["ppt/presentation.xml"])
        listed = [el.get("id") for el in written.find(f"{{{self.PML}}}sldIdLst")]
        return done, listed, [[el.get("id") for el in section.iter(f"{{{self.P14}}}sldId")] for section in written.iter(f"{{{self.P14}}}section")]

    def test_a_deck_with_sections_keeps_every_slide_in_one(self):
        done, listed, grouped = self.sectioned([{"carry": 1}, {"carry": 2}, {"carry": 4}, {"composed": 1}, {"carry": 5}])
        self.assertEqual(done.returncode, 0, done.stderr)
        # Slide 3 left the opening; the composed slide, which stands for no slide, joined the body after the slide it follows.
        self.assertEqual(grouped, [listed[:2], listed[2:]])
        self.assertEqual(len(listed), 5)

    def test_a_slide_drawn_in_place_of_the_first_of_a_section_opens_that_section(self):
        # Slide 4 opens the body. Placed by the slide before it, its replacement joined the opening and the body lost its
        # first page; it takes slide 4's place instead, and the opening keeps its three.
        done, listed, grouped = self.sectioned([{"carry": 1}, {"carry": 2}, {"carry": 3}, {"composed": 1, "sourceSlide": 4}, {"carry": 5}])
        self.assertEqual(done.returncode, 0, done.stderr)
        self.assertEqual(grouped, [listed[:3], listed[3:]])
        self.assertEqual(len(listed), 5)
        # A slide that is not the source deck's is no place to take.
        done, _, _ = self.sectioned([{"carry": 1}, {"composed": 1, "sourceSlide": 99}])
        self.assertEqual(done.returncode, 2)
        self.assertIn("in place of source slide 99", done.stderr)


class BuildTests(SwappedDeck):
    def test_the_build_composes_the_changed_page_and_assembles_the_deck_around_it(self):
        revision = self.build["revision"]
        self.assertEqual((revision["carried"], revision["edited"], revision["composed"], revision["dropped"]), (7, 0, 1, [7]))
        self.assertEqual(revision["preserved"]["slides"], 7)
        self.assertEqual(revision["preserved"]["drifted"], [])
        self.assertEqual([item.get("carried", "composed") for item in revision["order"]], [1, 2, 3, 4, 5, 6, "composed", 8])
        self.assertEqual(self.build["status"], "built-unrendered")
        self.assertEqual(self.build["blockers"], [])
        out = self.swapped / "out"
        # The plan names the slide the composed page is drawn in place of, whose place in the deck's sections it takes.
        self.assertEqual(json.loads((out / "assembly.json").read_text())["slides"][6], {"composed": 1, "id": "s07", "sourceSlide": 7})
        self.assertEqual(Path(self.build["pptxPath"]).resolve(), (out / "deck.pptx").resolve())
        self.assertTrue((out / "deck.composed.pptx").exists())
        # The composed pages' own scene holds the one page; the assembled deck's holds all eight, in order.
        scene = json.loads((out / "scene.json").read_text())
        self.assertEqual([s["id"] for s in scene["slides"]], ["s07"])
        deck = json.loads((out / "deck-scene.json").read_text())
        self.assertEqual([(s["id"], "carried" in s) for s in deck["slides"]], [(f"s0{n}", n != 7) for n in range(1, 9)])
        # The composed page prints its place in the deck, not its place among the composed pages.
        number = next(node["text"] for node in scene["slides"][0]["nodes"] if node.get("role") == "page-number")
        self.assertEqual(int(number), 7)

    def test_the_measure_of_what_was_kept_agrees_with_the_build(self):
        report = self.swapped / "preservation.json"
        done = sh(PYTHON, BENCH / "preservation.py", self.bed / "deck.source.pptx", self.build["pptxPath"], "--json", report)
        self.assertEqual(done.returncode, 0, done.stderr)
        slides = json.loads(report.read_text())["slides"]
        self.assertEqual([s["parts"] for s in slides], [True] * 6 + [False, True])
        self.assertEqual([s["geometry"] for s in slides if s["parts"]], [0.0] * 7)

    def test_a_revision_of_text_edits_composes_nothing(self):
        def number(doc):
            for key in ("s02", "s08"):
                self.page(doc, key)["replace"] = [{"old": "£12.4m", "new": "£12.6m"}]
            self.page(doc, "s04")["replace"] = [{"old": "12.4", "new": "12.6"}]
        work = self.revise(number)
        self.assertEqual(self.author(work).returncode, 0)
        built = sh(NODE, RUNTIME / "build-deck.mjs", work / "deck.deck.json", work / "out", "--no-render")
        self.assertEqual(built.returncode, 0, built.stderr[:1500] + built.stdout[:1500])
        result = json.loads((work / "out" / "build-result.json").read_text())
        self.assertEqual((result["revision"]["carried"], result["revision"]["edited"], result["revision"]["composed"]), (5, 3, 0))
        self.assertFalse((work / "out" / "deck.composed.pptx").exists())


class EvidenceOnARevisionTests(CarriedDeck):
    def test_a_revision_needs_no_insight_log_and_one_page_may_bring_its_own(self):
        def swap(doc):
            doc["deck"]["density"] = "live-pitch"
            doc["pages"][doc["pages"].index(self.page(doc, "s07"))] = copy.deepcopy(DONUT)
        work = self.revise(swap)
        done = self.author(work, "--check")
        self.assertEqual(done.returncode, 0, done.stderr[:800])
        advisories = last_json(done.stdout)["advisories"]
        # No log: nothing to trace the page's typed numbers against, and nothing asked of the deck's evidence.
        self.assertFalse(any(code in line for line in advisories for code in ("NUMBER_UNTRACED", "MEASURES_MISSING", "ANALYSIS_REQUIRED", "CLAIM_MEASURES_MISSING")))

        # A log for the changed page alone: the page states its figure by reference, and the rest of the deck is asked for nothing.
        def bound(doc):
            swap(doc)
            page = self.page(doc, "s07")
            page["evidence"] = ["I1"]
            page["exhibit"]["basis"] = {"measures": ["I1/channel_share"], "role": "proof"}
            page["settles"] = {"kind": "share", "what": "Share of FY26 sales by channel", "measures": ["I1/channel_share"]}
            page["points"][0] = "Wholesale is {{I1/channel_share@Wholesale | 0}}% of sales, up from 16% in FY22, and it is the one channel that grows without a new lease"
        work = self.revise(bound)
        (work / "deck.insights.json").write_text(json.dumps({"schema": "professional-slides.insights/v1", "insights": [{
            "id": "I1", "finding": "Wholesale is 22% of FY26 sales, more than twice delivery", "soWhat": "Wholesale is the channel that grows without a lease", "shape": "mix", "strength": "strong",
            "calculation": "channel sales over total sales", "sources": ["Harbour Coffee management accounts FY26"], "breadth": {"parts": 4},
            "measures": {"channel_share": {"unit": "%", "population": "FY26 sales", "period": "FY26", "members": ["In store", "Takeaway", "Delivery", "Wholesale"], "values": [41, 27, 10, 22]}}}]}))
        done = self.author(work, "--check")
        self.assertEqual(done.returncode, 0, done.stderr[:1500])
        summary = last_json(done.stdout)
        self.assertFalse(any(code in line for line in summary["advisories"] for code in ("MEASURES_MISSING", "ANALYSIS_REQUIRED", "PILLAR_UNSUPPORTED")))
        # The runtime wrote the page's values from the measure it names.
        self.assertEqual(self.author(work).returncode, 0)
        slide = json.loads((work / "deck.deck.json").read_text())["slides"][0]
        self.assertIn("Wholesale is 22% of sales", json.dumps(slide))
        self.assertEqual(slide["pageType"]["content"]["evidence"], ["I1"])


@unittest.skipUnless(os.environ.get("PS_RUN_SLOW") == "1", "opt-in: set PS_RUN_SLOW=1 or pass --slow")
@unittest.skipUnless(NODE and HAS_PPTX and HAS_PILLOW and HAS_RENDERER, "needs Node.js, python-pptx, Pillow, LibreOffice (soffice) and pdftoppm")
class PointChangeBenchmarkTests(unittest.TestCase):
    def test_the_scripted_author_meets_every_check_of_every_task(self):
        with tempfile.TemporaryDirectory() as work:
            report = Path(work) / "report.json"
            done = sh(NODE, BENCH / "run.mjs", "--work", work, "--json", report)
            self.assertEqual(done.returncode, 0, done.stdout + done.stderr[-2000:])
            results = json.loads(report.read_text())
            self.assertEqual([r["task"] for r in results], ["number-propagates", "retitle", "exhibit-swap", "notations", "bare-number", "swap-own-title", "retitle-contradicts"])
            by = {r["task"]: r for r in results}
            for r in results:
                self.assertTrue(r["accepted"], (r["task"], {name: check for name, check in r.get("checks", {}).items() if not check.get("ok")} or r.get("error")))
                self.assertEqual(r["checks"]["scoped"]["refusalsOnUntouched"], 0)
                self.assertEqual(r["checks"]["preserved"]["identical"], r["checks"]["preserved"]["untouched"])
                # Nothing was changed that the request did not ask for.
                self.assertEqual(r["checks"]["asked"]["unasked"], [], r["task"])
            number = by["number-propagates"]
            # The first compile is refused for the figure left on the closing slide and in the table's total - the same
            # quantity under "Revenue (£m)".
            self.assertEqual(number["cost"]["refusedRuns"], 1)
            self.assertEqual(sorted(number["checks"]["stale"]["named"]), ["s04", "s08"])
            self.assertEqual(by["exhibit-swap"]["cost"]["composed"], 1)
            # The swap's `made` asks for a title that is no longer the label, in whatever words the author chose.
            self.assertTrue(by["exhibit-swap"]["checks"]["made"]["retitled"][0]["ok"])
            # One figure said five ways: the first move is refused naming the table, "£12.4 million", the close and the speaker's
            # notes, and the edits reach all of them.
            notations = by["notations"]
            self.assertEqual(sorted(notations["checks"]["stale"]["named"]), ["s04", "s05", "s06", "s08"])
            # The figure typed bare stands only inside longer numbers: refused with them named, and they are untouched at the end.
            bare = by["bare-number"]
            self.assertEqual(bare["checks"]["refusedFor"], {"ok": True, "wanted": ["REVISION_CARRY_INVALID"], "got": ["NUMBER_STALE", "REVISION_CARRY_INVALID"]})
            self.assertTrue(all(p["ok"] for p in bare["checks"]["made"]["prints"]))
            # The pie redrawn under the slide's own label title: compiled at the first attempt.
            swap = by["swap-own-title"]
            self.assertEqual(swap["cost"]["sequence"], ["import", "check", "compile", "build"])
            # The retitle that contradicts a slide nobody asked to change: that slide is byte for byte the source's.
            self.assertEqual(by["retitle-contradicts"]["checks"]["preserved"]["identical"], 7)

    def test_an_edit_nobody_asked_for_fails_the_benchmark(self):
        # The restated figure, and beside it a "correction" the request did not ask for: the growth rate the figure implies.
        tasks = json.loads((BENCH / "tasks.json").read_text())
        task = copy.deepcopy(next(t for t in tasks["tasks"] if t["id"] == "number-propagates"))
        task["moves"] = [{"edits": [{"page": "s02", "replace": [{"old": "£12.4m", "new": "£12.6m"}, {"old": "14%", "new": "16%"}]},
                                    {"page": "s08", "replace": [{"old": "£12.4m", "new": "£12.6m"}]}, {"page": "s04", "replace": [{"old": "12.4", "new": "12.6"}]}]}]
        task["accept"]["stale"] = None
        with tempfile.TemporaryDirectory() as work:
            (Path(work) / "tasks.json").write_text(json.dumps({"tasks": [task]}))
            report = Path(work) / "report.json"
            done = sh(NODE, BENCH / "run.mjs", "--tasks", Path(work) / "tasks.json", "--work", Path(work) / "run", "--json", report)
            self.assertEqual(done.returncode, 2, done.stdout + done.stderr[-2000:])
            result = json.loads(report.read_text())[0]
            self.assertFalse(result["accepted"])
            self.assertEqual({name for name, check in result["checks"].items() if not check["ok"]}, {"asked"})
            self.assertEqual([item["slide"] for item in result["checks"]["asked"]["unasked"]], [2])
            self.assertIn("up 16% on FY25", result["checks"]["asked"]["unasked"][0]["what"])

    def test_a_retitle_is_read_for_what_it_says_and_not_for_the_scripted_sentence(self):
        # An author asked for a title writes its own sentence. The benchmark accepts one that says what was asked, refuses one
        # that does not, and still refuses a line of the same slide rewritten beside it.
        tasks = json.loads((BENCH / "tasks.json").read_text())
        base = next(t for t in tasks["tasks"] if t["id"] == "retitle")
        self.assertEqual(base["accept"]["retitled"], [{"slide": 4, "says": ["London", "fastest", "margin"]}])
        variants = {"own": [{"page": "s04", "title": "London grows fastest, at 11.0%, on the thinnest margin"}],
                    "other": [{"page": "s04", "title": "The South has the widest margin of the four"}],
                    "beside": [{"page": "s04", "title": "London grows fastest on the thinnest margin"}, {"page": "s04", "replace": [{"old": "31%", "new": "33%"}]}]}
        made = []
        for name, edits in variants.items():
            task = copy.deepcopy(base)
            task.update(id=name, moves=[{"note": "the author's own title", "edits": edits}])
            made.append(task)
        with tempfile.TemporaryDirectory() as work:
            (Path(work) / "tasks.json").write_text(json.dumps({"tasks": made}))
            report = Path(work) / "report.json"
            done = sh(NODE, BENCH / "run.mjs", "--tasks", Path(work) / "tasks.json", "--work", Path(work) / "run", "--json", report)
            self.assertEqual(done.returncode, 2, done.stdout + done.stderr[-2000:])
            by = {r["task"]: r for r in json.loads(report.read_text())}
            failed = {name: {check for check, read in r["checks"].items() if not read["ok"]} for name, r in by.items()}
            self.assertEqual(failed, {"own": set(), "other": {"made"}, "beside": {"asked"}})
            self.assertTrue(by["own"]["accepted"])
            self.assertIn("33%", by["beside"]["checks"]["asked"]["unasked"][0]["what"])

if __name__ == "__main__":
    unittest.main()
