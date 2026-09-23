"""Three export defects a fifty-page evaluation deck surfaced.

A standalone donut named none of its categories; a highlighted phrase just
after a line wrap split a takeaway into two PowerPoint paragraphs; and a 100%
stack printed 35.8 in its scene and 36 in PowerPoint.
"""
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from node_probe import run_node  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
RUNTIME = ROOT / "skills/professional-slides/runtime"


class DonutLegendTests(unittest.TestCase):
    def test_a_composed_donut_names_its_categories_and_its_native_chart_keeps_the_legend(self):
        result = run_node("""
import { toDeckPlan } from './skills/professional-slides/runtime/compose.mjs';
import { planDeck } from './skills/professional-slides/runtime/planner.mjs';
const spec={schema:'professional-slides.deck/v3',id:'d',cover:{title:'x'},slides:[{title:'Two firms supply most of the judged pages',layout:'exhibit-full',
  exhibit:{type:'chart.donut',heading:'Pages by firm',unit:'pages',labels:['BCG','McKinsey','Bain'],values:[75,52,5]}}]};
const {deck}=planDeck(toDeckPlan(spec,'.'));
const page=deck.slides.at(-1);
const donut=page.componentInstances.find(c=>c.component==='chart.donut');
console.log(JSON.stringify({variant:donut.variant,legend:page.nodes.filter(n=>n.role==='legend-label').map(n=>n.text),native:donut.nativeChart?.legend}));
""")
        self.assertEqual(result["variant"], "legend-top-right")
        self.assertEqual(result["legend"], ["BCG", "McKinsey", "Bain"])
        self.assertTrue(result["native"])


class HighlightWrapTests(unittest.TestCase):
    def test_a_highlight_after_a_wrap_keeps_one_source_paragraph(self):
        result = run_node("""
import { REGISTRY } from './skills/professional-slides/runtime/registry.mjs';
const text='Hold pages to the strong benchmarks. The diagram floor still binds, because the floor is hard, but it is the first number to re-measure as more diagram pages are judged.';
const nodes=REGISTRY.get('insight').render({id:'i',frame:{x:0,y:0,width:900,height:120},props:{text,highlight:'diagram'}}).nodes;
const body=nodes.find(n=>n.role==='insight-body');
const src=body.data.textLayout.sourceRuns||[];
console.log(JSON.stringify({wrapped:body.text.includes('\\n'),runsHaveBreak:(body.runs||[]).some(r=>r.text.includes('\\n')),
  sourceRuns:src.length,sourceBreak:src.some(r=>r.text.includes('\\n')),joined:src.map(r=>r.text).join('')===text}));
""")
        self.assertTrue(result["wrapped"])
        self.assertGreater(result["sourceRuns"], 1)
        self.assertFalse(result["sourceBreak"])
        self.assertTrue(result["joined"])


class ImageDividerFooterTests(unittest.TestCase):
    def test_the_footer_stays_in_the_panel_not_on_the_photograph(self):
        # White footer type over a crowd in the photo was unreadable.
        result = run_node("""
import { REGISTRY } from './skills/professional-slides/runtime/registry.mjs';
const img = {dataUri:'data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg==',width:1,height:1,alt:'A photograph',authorization:'test'};
const nodes = REGISTRY.get('section-divider').render({id:'d',frame:{x:0,y:0,width:1280,height:720},props:{title:'The test',companyName:'A deck footer',image:img}}).nodes;
const photo = nodes.find(n=>n.role==='divider-image'), foot = nodes.find(n=>n.role==='footer-right');
console.log(JSON.stringify({photoLeft:photo.frame.x, footRight:foot ? foot.frame.x+foot.frame.width : null}));
""")
        self.assertIsNotNone(result["footRight"])
        self.assertLessEqual(result["footRight"], result["photoLeft"])


class DeclaredDecimalsTests(unittest.TestCase):
    def test_a_declared_format_rounds_half_away_from_zero_like_powerpoint(self):
        # 3.55 is stored as 3.5499...; toFixed printed 3.5 on the drawn chart
        # while PowerPoint printed 3.6 on the native one.
        result = run_node("""
import { formatValue } from './skills/professional-slides/runtime/value-format.mjs';
console.log(JSON.stringify([3.55, -1.25, 2.345].map(v => formatValue(v, {valueFormat: {decimals: 1}}))));
""")
        self.assertEqual(result, ["3.6", "-1.3", "2.3"])


class SeriesRoundingTests(unittest.TestCase):
    def test_the_native_label_format_matches_the_drawn_one(self):
        sys.path.insert(0, str(RUNTIME / "emit"))
        from pptx import Presentation  # noqa: E402
        scene = run_node("""
import { toDeckPlan } from './skills/professional-slides/runtime/compose.mjs';
import { planDeck } from './skills/professional-slides/runtime/planner.mjs';
const spec={schema:'professional-slides.deck/v3',id:'r',cover:{title:'x'},slides:[{title:'BCG leads with charts and McKinsey with text',layout:'exhibit-full',
  exhibit:{type:'chart.stacked-column',percent:true,heading:'Pages by family',categories:['BCG','McKinsey'],
    series:[{name:'Chart',values:[24,12]},{name:'Text',values:[13,15]},{name:'Table',values:[14,10]}]}}]};
console.log(JSON.stringify(planDeck(toDeckPlan(spec,'.')).deck));
""")
        labels = [n["text"] for n in scene["slides"][-1]["nodes"] if n["role"] == "data-label"]
        self.assertIn("47.1", labels)
        with tempfile.TemporaryDirectory() as tmp:
            scene_path, pptx_path = Path(tmp) / "scene.json", Path(tmp) / "r.pptx"
            scene_path.write_text(json.dumps(scene))
            subprocess.run([sys.executable, str(RUNTIME / "emit" / "emit_pptx.py"), str(scene_path), str(pptx_path)],
                           check=True, capture_output=True)
            chart = next(s.chart for s in Presentation(pptx_path).slides[-1].shapes if s.has_chart)
            # One decimal where the scene has one, none where it has none: the
            # drawn chart writes 47.1 beside 27, so the native chart must too.
            self.assertEqual(chart.plots[0].data_labels.number_format, "General")


if __name__ == "__main__":
    unittest.main()
